import logging
import re
from sqlalchemy.orm import Session
from sqlalchemy.exc import IntegrityError
from . import models, schemas
from sqlalchemy import func # Asegúrate de poner esto al INICIO de tu archivo crud.py si no lo tienes
from typing import List, Optional
from sqlalchemy import and_, or_
import uuid as uuid_lib
from datetime import date
from modulos.superadmin import models as superadmin_models
from modulos.superadmin import crud as superadmin_crud
import io
from openpyxl import Workbook
from datetime import datetime, timedelta, timezone
from modulos.instituciones import models as inst_models
import hashlib
import operator as _operator
from modulos.educativo import models
from core.algoritmo import calcular_resultados, invertir_valor_item, resolver_escala
from core.cola_tareas import app as celery_app
from core.concurrencia import commit_o_conflicto
from fastapi import HTTPException, status
from modulos.auth import models as auth_models

logger = logging.getLogger("aurorasoft")

# ==========================================
# CONSTANTES Y HELPERS DE ALERTAS
# ==========================================

def detectar_niveles_de_alerta(diagnosticos):
    """Devuelve los niveles de los diagnósticos marcados como alerta.

    Usa el indicador explícito `DiagnosticoAlumno.es_alerta` (copiado del
    baremo al interpretar). Antes se comparaba `nivel` contra un conjunto
    cerrado de cadenas, así que un baremo llamado "Alto riesgo" o "Severo"
    no disparaba nada y el informe de un alumno en crisis se liberaba solo.
    """
    return [
        (d.nivel or "Nivel sin nombre")
        for d in diagnosticos
        if getattr(d, "es_alerta", False)
    ]


def obtener_destinatarios_alerta(db: Session, evaluacion):
    """Cadena de respaldo para que una alerta clínica nunca se pierda (C-13).

    Antes el aviso salía únicamente hacia `evaluacion.creador_id`: si ese
    psicólogo estaba desactivado, dado de baja o sin correo, nadie se enteraba
    de un caso crítico.

    Orden de destinatarios, acumulativo y sin duplicados:
      1. el creador de la campaña, si sigue activo y tiene correo;
      2. los admins activos de ESA institución (respaldo institucional);
      3. la cuenta genérica "colegio" activa, si existe;
      4. como último recurso, el email de contacto de la institución.

    Todas las consultas se filtran por `evaluacion.institucion_id`, así que una
    alerta jamás puede cruzarse hacia otro colegio.
    """
    destinatarios = []
    motivos = []

    def agregar(correo, motivo):
        if correo and correo not in destinatarios:
            destinatarios.append(correo)
            motivos.append(motivo)

    if evaluacion is None:
        return [], []

    # 1. Creador de la campaña — solo si sigue habilitado
    creador = db.query(auth_models.Usuario).filter(
        auth_models.Usuario.id == evaluacion.creador_id
    ).first()
    if creador and creador.activo:
        agregar(creador.email, "creador de la campaña")
    elif creador:
        logger.warning(
            "El creador de la campaña %s (usuario %s) está inactivo; "
            "la alerta se deriva al respaldo institucional.",
            evaluacion.id, creador.id,
        )

    # 2 y 3. Respaldo institucional, acotado a la institución de la campaña
    respaldos = db.query(auth_models.Usuario).filter(
        auth_models.Usuario.institucion_id == evaluacion.institucion_id,
        auth_models.Usuario.rol.in_(["admin", "colegio"]),
        auth_models.Usuario.activo == True,  # noqa: E712
    ).order_by(auth_models.Usuario.rol, auth_models.Usuario.id).all()
    for usuario in respaldos:
        agregar(usuario.email, f"respaldo institucional ({usuario.rol})")

    # 4. Último recurso: el correo de contacto del colegio
    if not destinatarios:
        institucion = db.query(inst_models.Institucion).filter(
            inst_models.Institucion.id == evaluacion.institucion_id
        ).first()
        if institucion:
            agregar(institucion.email_contacto, "correo de contacto de la institución")

    if not destinatarios:
        logger.error(
            "ALERTA SIN DESTINATARIO: la campaña %s de la institución %s no tiene "
            "ningún correo al que avisar. Revisar usuarios activos del colegio.",
            evaluacion.id, evaluacion.institucion_id,
        )

    return destinatarios, motivos


def obtener_revisor_original_de_inscripcion(db: Session, inscripcion) -> Optional[auth_models.Usuario]:
    """El usuario detrás de `Inscripcion.revisado_por_usuario_id`, si sigue activo.

    Usado SOLO por `reprocesar_informe_de_inscripcion` para avisarle,
    puntualmente, a quien había liberado el informe que un reproceso
    acaba de retenerlo de nuevo (riesgo documentado en
    docs/gaps-pendientes.md: esa persona es la única con contexto real
    para saber si ya se le entregó el resultado al alumno/apoderado fuera
    de la plataforma). `revisado_por_usuario_id` no se toca durante el
    reproceso (ver reprocesar_informe_de_inscripcion), así que en este
    punto sigue describiendo quién tomó la última decisión de
    revisión/liberación ANTES de la corrección. `None` si nadie reviso
    nunca a mano (quedó liberado automáticamente, sin alerta) o si ese
    usuario ya no está activo — en ambos casos la alerta genérica de
    `obtener_destinatarios_alerta` sigue cubriendo el caso.
    """
    if not getattr(inscripcion, "revisado_por_usuario_id", None):
        return None
    usuario = db.query(auth_models.Usuario).filter(
        auth_models.Usuario.id == inscripcion.revisado_por_usuario_id,
    ).first()
    return usuario if usuario and usuario.activo else None


def consentimiento_vigente_para_inscripcion(db: Session, inscripcion) -> bool:
    """¿Esta inscripción sigue amparada por un consentimiento válido?

    🔒 B-8 — La comprobación se acota a la institución DUEÑA DE LA CAMPAÑA.
    Antes se preguntaba solo por `persona_id`, así que un consentimiento
    registrado en el colegio A habilitaba una inscripción del colegio B.

    Los links anónimos (sin `persona_id`) no dependen del padrón: su
    consentimiento es el que el propio alumno acepta al entrar — el flag de
    política de abajo no tiene sentido para ellos (no hay identidad a la que
    pedirle un consentimiento de estudiante), así que siguen exentos.

    🆕 Política institucional de consentimientos — el consentimiento del
    APODERADO sigue siendo obligatorio siempre (comportamiento sin cambios).
    Cuando `Institucion.exigir_asentimiento_estudiante` está activo, se
    exige EN VIVO, en cada llamada (igual que el apoderado), también el
    asentimiento vigente del propio estudiante — una revocación de
    cualquiera de los dos bloquea el reintento/reanudación del examen. Con
    el flag apagado (default de toda institución existente) esta función es
    idéntica, campo por campo, a como era antes de este cambio.
    """
    if not inscripcion.persona_id:
        return True

    evaluacion = db.query(models.Evaluacion).filter(
        models.Evaluacion.id == inscripcion.evaluacion_id
    ).first()
    institucion_id = evaluacion.institucion_id if evaluacion else None

    if not superadmin_crud.tiene_consentimiento_vigente(db, inscripcion.persona_id, institucion_id):
        return False

    institucion = db.query(inst_models.Institucion).filter(
        inst_models.Institucion.id == institucion_id
    ).first()
    if institucion and institucion.exigir_asentimiento_estudiante:
        return superadmin_crud.tiene_consentimiento_vigente(
            db, inscripcion.persona_id, institucion_id,
            tipo_consentimiento=superadmin_models.TIPO_CONSENTIMIENTO_ASENTIMIENTO,
        )
    return True


def campana_esta_expirada(evaluacion) -> bool:
    """🔒 True si esta `Evaluacion` ya pasó su `fecha_fin`.

    Único punto de verdad para el candado de expiración que gatea los 7
    puntos del flujo de examen en modulos/educativo/router.py (guardar,
    finalizar, comenzar-test, progreso, preguntas, campos-socio ×2) y
    `matricular_roster_en_campana` más abajo en este mismo archivo — antes
    de esto, la condición `evaluacion.fecha_fin and datetime.utcnow() >
    evaluacion.fecha_fin` estaba pegada a mano en cada uno de esos 8 sitios.
    Hoy las 8 copias coinciden byte a byte, pero nada obligaba a que
    siguieran así: una edición futura que tocara el candado en un sitio y se
    olvidara de los otros 7 los habría desincronizado en silencio (algunos
    endpoints seguirían aceptando un examen sobre una campaña ya cerrada).
    Cerrar/cancelar una campaña (ver `_cerrar_o_cancelar_campana`) reutiliza
    exactamente este mismo candado forzando `fecha_fin` a "ahora", así que
    centralizarlo aquí también protege esa reutilización.

    🆕 Cierre de gap (docs/gaps-pendientes.md) — deliberadamente sigue
    usando datetime.utcnow() naive, no datetime.now(timezone.utc):
    `Evaluacion.fecha_fin` es DateTime SIN zona horaria, y
    `EvaluacionCreate.fecha_fin_a_utc_naive` (schemas.py) normaliza
    explícitamente cualquier fecha entrante a naive-UTC antes de guardarla
    — comparar aquí contra un datetime aware lanzaría TypeError en los 8
    puntos del flujo de examen que dependen de esta función.
    """
    return bool(evaluacion and evaluacion.fecha_fin and datetime.utcnow() > evaluacion.fecha_fin)


def inscripcion_esta_activa(inscripcion) -> bool:
    """🔒 C-8 — Una inscripción eliminada o anonimizada deja de ser utilizable.

    Antes el link mágico seguía funcionando después de un borrado lógico o de
    una anonimización ARCO: el alumno podía seguir rindiendo y generando datos
    sobre un registro que ya había sido dado de baja.
    """
    return not (
        getattr(inscripcion, "is_deleted", False)
        or getattr(inscripcion, "is_anonymized", False)
    )


def bloquear_inscripcion_para_finalizar(db: Session, token: str):
    """Trae la `Inscripcion` de este token con FOR UPDATE, forzando estado fresco.

    Se usa desde `finalizar_examen_seguro`, que antes de esto ya consultó la
    misma `Inscripcion` sin bloqueo (Fase 1, para validar sin saturar la BD si
    la petición se va a rechazar de entrada). Esa consulta previa deja el
    objeto cargado en el identity map de esta Session; sin `populate_existing()`,
    SQLAlchemy devuelve por defecto ese MISMO objeto sin refrescar sus atributos
    con la fila que Postgres acaba de desbloquear — aunque el candado de fila
    en sí funcione perfectamente. El síntoma: un segundo envío concurrente
    esperaba correctamente a que el primero confirmara, pero al desbloquearse
    seguía viendo el `estado` obsoleto de su propia Fase 1 en vez de
    "Finalizado", así que el candado de idempotencia de abajo nunca se
    disparaba para él — recalculaba todo de nuevo y confirmaba una segunda vez
    (dos 200, alerta de riesgo duplicada, intento duplicado de generar el
    informe -> UniqueViolation en `uq_edu_informes_inscripcion`). Aislado en su
    propia función para que la prueba de concurrencia
    (`tests/verificar_concurrencia_finalizar_e2e.py`) pueda ejercitar
    exactamente esta consulta sin reimplementarla.
    """
    return db.query(models.Inscripcion).filter(
        models.Inscripcion.token == token
    ).with_for_update().populate_existing().first()


def contar_preguntas_sin_responder(db: Session, inscripcion_id: int, prueba_id: int) -> int:
    """🆕 Segunda capa de defensa contra un examen incompleto (hallazgo de
    auditoría, sesión 2026-10-16/17, ver docs/gaps-pendientes.md) —
    independiente del candado de frontend (`faltanPreguntas`, ya endurecido
    contra fallos de autoguardado). Hasta ahora, `finalizar_examen_seguro`
    confiaba por completo en que el frontend solo dejaba pasar exámenes
    completos: un bug futuro de frontend, una llamada directa a la API, o un
    cliente distinto del oficial podía finalizar un examen incompleto sin que
    el backend lo notara — `calcular_resultados_dimensiones` simplemente suma
    lo que de verdad llegó a guardarse, así que una pregunta sin responder no
    producía ningún error, solo un puntaje silenciosamente más bajo.

    Cuenta cuántas `Pregunta` de la prueba (incluidas las que no participan
    en ningún puntaje, ej. un ítem `incluir_en_informe` sin dimension_id —
    el alumno igual debe responderlas, el frontend las muestra igual) NO
    tienen todavía una `RespuestaAlumno` para esta inscripción. Se usa desde
    `finalizar_examen_seguro` DESPUÉS del candado de idempotencia y ANTES de
    calcular nada — un examen incompleto nunca llega a `calcular_resultados_
    dimensiones`, se rechaza con 409 y la inscripción queda exactamente como
    estaba (ni "Finalizado" ni con resultados parciales), igual de
    reintentable que antes de este cambio.
    """
    total_preguntas = db.query(func.count(models.Pregunta.id)).filter(
        models.Pregunta.prueba_id == prueba_id
    ).scalar() or 0

    preguntas_respondidas = db.query(
        func.count(func.distinct(models.RespuestaAlumno.pregunta_id))
    ).filter(
        models.RespuestaAlumno.inscripcion_id == inscripcion_id
    ).scalar() or 0

    return max(total_preguntas - preguntas_respondidas, 0)


def bloquear_inscripcion_para_reprocesar(db: Session, inscripcion_id: int):
    """Trae la `Inscripcion` de este id con FOR UPDATE, forzando estado fresco.

    Mismo motivo que `bloquear_inscripcion_para_finalizar` (ver su docstring):
    el endpoint de reproceso ya consultó esta misma `Inscripcion` sin bloqueo
    para las validaciones previas (existe, está "Finalizado", está activa),
    dejándola en el identity map de esta Session. Sin `populate_existing()`,
    un reproceso lanzado dos veces casi al mismo tiempo para la MISMA
    inscripción vería, tras desbloquearse, los datos obsoletos de su propia
    consulta previa en vez del estado que el primer reproceso ya escribió.

    🆕 Hallazgo de auditoría (sesión 2026-10-16): `liberar_informe`/
    `actualizar_revision_informe` (router.py) también la usan ahora — ambos
    escriben `informe_liberado`/`resultado_visible_alumno`, los MISMOS dos
    campos que `reprocesar_informe_de_inscripcion` recalcula. Sin este
    candado compartido, una liberación/revisión manual podía intercalarse
    con un reproceso concurrente sobre la misma inscripción de forma no
    determinista (el orden real de ejecución dependía del azar de
    scheduling, no de cuál request llegó primero). El nombre de la función
    quedó igual a propósito (evita romper los tests existentes que la
    referencian por nombre) aunque ya no sea exclusiva de reprocesar — el
    candado en sí es completamente genérico, nada de su lógica depende de
    quién lo use.
    """
    return db.query(models.Inscripcion).filter(
        models.Inscripcion.id == inscripcion_id
    ).with_for_update().populate_existing().first()


def limpiar_resultados_previos(db: Session, inscripcion_id: int):
    """Borra puntajes y diagnósticos anteriores de esta inscripción.

    Hace que recalcular sea idempotente: si una finalización se revirtió a
    medias, el reintento no acumula filas duplicadas. No hace commit — la
    finalización controla la transacción completa (ver B-9).
    """
    db.query(models.DiagnosticoAlumno).filter(
        models.DiagnosticoAlumno.inscripcion_id == inscripcion_id
    ).delete(synchronize_session=False)
    db.query(models.ResultadoDimension).filter(
        models.ResultadoDimension.inscripcion_id == inscripcion_id
    ).delete(synchronize_session=False)

# ==========================================
# CRUD PARA LA FICHA SOCIODEMOGRÁFICA
# ==========================================
def crear_campo_socio(db: Session, campo: schemas.CampoSociodemograficoCreate):
    # Toma los datos del esquema y los guarda en la base de datos
    nuevo_campo = models.CampoSociodemografico(**campo.model_dump())
    db.add(nuevo_campo)
    db.commit()
    db.refresh(nuevo_campo)
    return nuevo_campo

# ==========================================
# CRUD PARA EL MOTOR DE TESTS
# ==========================================
def crear_prueba(db: Session, prueba: schemas.PruebaCreate, creado_por_usuario_id: int = None):
    datos_prueba = prueba.model_dump() if hasattr(prueba, "model_dump") else prueba.dict()
    
    db_prueba = models.Prueba(
        **datos_prueba,
        creado_por_usuario_id=creado_por_usuario_id,
    )
    db.add(db_prueba)
    db.commit()
    db.refresh(db_prueba)
    return db_prueba

def _validar_dimension_no_compuesta_para_pregunta(db: Session, dimension_id: Optional[int]):
    """Una dimensión 'suma_de_dimensiones' (ver Dimension.tipo_calculo) no
    tiene preguntas propias — su puntaje viene siempre de otras dimensiones,
    nunca de respuestas directas. Guard compartido por crear_pregunta,
    actualizar_pregunta y cargar_test_masivo."""
    if dimension_id is None:
        return
    dimension = db.query(models.Dimension).filter(models.Dimension.id == dimension_id).first()
    if dimension and dimension.tipo_calculo == models.TIPO_CALCULO_SUMA_DE_DIMENSIONES:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                f"La dimensión '{dimension.nombre}' es 'suma_de_dimensiones' (calculada a partir "
                "de otras dimensiones): no puede tener preguntas propias."
            ),
        )


def crear_pregunta(db: Session, pregunta: schemas.PreguntaCreate):
    _validar_dimension_no_compuesta_para_pregunta(db, pregunta.dimension_id)
    nueva_pregunta = models.Pregunta(**pregunta.model_dump())
    db.add(nueva_pregunta)
    db.commit()
    db.refresh(nueva_pregunta)
    return nueva_pregunta

def crear_dimension(db: Session, dimension: schemas.DimensionCreate):
    """Guarda una nueva dimensión en la base de datos"""
    nueva_dimension = models.Dimension(**dimension.model_dump())
    db.add(nueva_dimension)
    db.commit()
    db.refresh(nueva_dimension)
    return nueva_dimension

def crear_opcion_respuesta(db: Session, opcion: schemas.OpcionRespuestaCreate):
    """Guarda una alternativa de respuesta y su valor en la base de datos"""
    nueva_opcion = models.OpcionRespuesta(**opcion.model_dump())
    db.add(nueva_opcion)
    db.commit()
    db.refresh(nueva_opcion)
    return nueva_opcion


# ==========================================
# MOTOR DE CALIFICACIÓN AUTOMÁTICA
# ==========================================
# ==========================================
# MOTOR DE CALIFICACIÓN AUTOMÁTICA
# ==========================================
def calcular_resultados_dimensiones(db: Session, inscripcion_id: int):
    """
    Motor Base: Agrupa las respuestas del alumno por Dimensión y calcula su Puntaje Directo.
    Delega la lógica de cálculo (Likert, Inversión, Correcto/Incorrecto) al motor centralizado 
    en algoritmo.py pasando la prueba completa.
    """
    # 1. Averiguamos qué prueba está resolviendo este alumno
    inscripcion = db.query(models.Inscripcion).filter(models.Inscripcion.id == inscripcion_id).first()
    if not inscripcion:
        return []

    evaluacion = db.query(models.Evaluacion).filter(models.Evaluacion.id == inscripcion.evaluacion_id).first()
    prueba = db.query(models.Prueba).filter(models.Prueba.id == evaluacion.prueba_id).first()

    # 2. Obtenemos los ítems de la prueba con sus atributos (ej: es_inverso)
    items_db = db.query(models.Pregunta).filter(models.Pregunta.prueba_id == prueba.id).all()

    # 3. Traemos todas las respuestas registradas para esta inscripción
    respuestas_db = db.query(
        models.RespuestaAlumno, 
        models.Pregunta
    ).join(
        models.Pregunta, models.RespuestaAlumno.pregunta_id == models.Pregunta.id
    ).filter(
        models.RespuestaAlumno.inscripcion_id == inscripcion_id
    ).all()

    # 🟢 FIX PUNTO 8: Definición de DTO ligero para evitar NameError
    class RespuestaDTO:
        def __init__(self, pregunta_id: int, dimension: int, valor_raw: float):
            self.pregunta_id = pregunta_id
            self.dimension = dimension
            self.valor_raw = valor_raw

    # 4. Adaptamos las respuestas al formato genérico necesario para algoritmo.py
    respuestas_adaptadas = [
        RespuestaDTO(
            pregunta_id=preg.id,
            dimension=preg.dimension_id,  # Agrupamos por ID de dimensión
            valor_raw=resp.puntaje_obtenido
        )
        for resp, preg in respuestas_db
    ]

    # 5. Ejecutamos el motor centralizado de cálculo pasando el objeto prueba completo
    #
    # `codigo_prueba` solo se usa para identificar la prueba en los logs. La
    # cadena anterior —getattr(p,"nombre_tecnico", getattr(p,"codigo", p.nombre))—
    # reventaba SIEMPRE con AttributeError: Python evalúa los argumentos por
    # adelantado, así que `prueba.nombre` se ejecutaba antes que ningún getattr
    # y `Prueba` no tiene esa columna (tiene nombre_tecnico y nombre_publico).
    codigo_prueba = (
        getattr(prueba, "nombre_tecnico", None)
        or getattr(prueba, "nombre_publico", None)
        or "prueba_desconocida"
    )

    # 🆕 Dimensiones 'conteo_sintomas' (ver Dimension.tipo_calculo/
    # conteo_umbral_item en models.py) — en vez de sumar los valores crudos de
    # sus ítems, cuentan cuántos alcanzan su propio umbral (ej. NICHQ
    # Vanderbilt: "6 de 9 ítems con puntuación ≥2"). El resto de dimensiones
    # (ausentes de este dict) se suman exactamente como siempre.
    dimensiones_prueba = db.query(models.Dimension).filter(models.Dimension.prueba_id == prueba.id).all()
    dimensiones_conteo = {
        d.id: d.conteo_umbral_item
        for d in dimensiones_prueba
        if d.tipo_calculo == models.TIPO_CALCULO_CONTEO_SINTOMAS and d.conteo_umbral_item is not None
    }

    resultado_calculo = calcular_resultados(
        codigo_prueba=codigo_prueba,
        respuestas=respuestas_adaptadas,
        items_db=items_db,
        prueba=prueba,  # Contiene escala_maxima, modo_calculo, etc.
        dimensiones_conteo=dimensiones_conteo,
    )

    puntajes_brutos = resultado_calculo.get("puntajes_brutos", {})
    nuevos_resultados = []

    # 6. Guardamos los Puntajes Directos calculados en la base de datos
    for dimension_id, puntaje in puntajes_brutos.items():
        if dimension_id is None:
            continue

        nuevo_resultado = models.ResultadoDimension(
            inscripcion_id=inscripcion_id,
            dimension_id=dimension_id,
            # 🔢 C-1 — 4 decimales, la misma precisión que la columna NUMERIC(12,4).
            # Antes se redondeaba a 2 y la columna entera truncaba el resto.
            puntaje_directo=round(float(puntaje), 4)
        )
        db.add(nuevo_resultado)
        nuevos_resultados.append(nuevo_resultado)

    # 🆕 Segunda pasada — dimensiones 'suma_de_dimensiones' (ver
    # Dimension.tipo_calculo/DimensionComponente en models.py). No tienen
    # preguntas propias, así que nunca aparecen en `puntajes_brutos`; su
    # puntaje_directo es la suma de los puntajes YA calculados en la primera
    # pasada de sus componentes (que, por construcción — ver la validación en
    # reemplazar_componentes_de_dimension — son siempre dimensiones 'items',
    # nunca otra compuesta: sin anidar, sin ciclos posibles). A partir de
    # aquí un ResultadoDimension de una compuesta es indistinguible de uno
    # normal para interpretar_resultados_alumno y todo lo que viene después.
    dimensiones_compuestas = db.query(models.Dimension).filter(
        models.Dimension.prueba_id == prueba.id,
        models.Dimension.tipo_calculo == models.TIPO_CALCULO_SUMA_DE_DIMENSIONES,
    ).all()

    for compuesta in dimensiones_compuestas:
        filas_componentes = (
            db.query(models.DimensionComponente, models.Dimension.tipo_calculo)
            .join(models.Dimension, models.Dimension.id == models.DimensionComponente.dimension_componente_id)
            .filter(models.DimensionComponente.dimension_compuesta_id == compuesta.id)
            .all()
        )
        # 🔒 Defensa en profundidad: `reemplazar_componentes_de_dimension` ya
        # bloquea anidar una compuesta dentro de otra al ESCRIBIR (ver
        # _validar_componentes_dimension_compuesta), pero ese guard no vive
        # en este camino de CÁLCULO — una fila de DimensionComponente
        # insertada por fuera de ese endpoint (p. ej. un seed script por ORM
        # directo) apuntando a otra compuesta nunca aparece en
        # `puntajes_brutos` (solo lo llena la primera pasada, de dimensiones
        # 'items'), así que antes contribuía 0 de forma indistinguible de un
        # componente real sin respuestas, sin ningún aviso. Ahora se detecta
        # y se deja constancia en el log en vez de fallar en silencio.
        ids_componentes = []
        for fila_componente, tipo_calculo_componente in filas_componentes:
            if tipo_calculo_componente == models.TIPO_CALCULO_SUMA_DE_DIMENSIONES:
                logger.error(
                    "Dimensión compuesta %s tiene como componente a la dimensión %s, que es OTRA "
                    "dimensión compuesta (tipo_calculo='suma_de_dimensiones') -- anidamiento no "
                    "soportado. Se excluye del cálculo de la inscripción %s; corrige los "
                    "componentes de la dimensión %s.",
                    compuesta.id, fila_componente.dimension_componente_id, inscripcion_id, compuesta.id,
                )
                continue
            ids_componentes.append(fila_componente.dimension_componente_id)
        puntaje_total = sum(puntajes_brutos.get(cid, 0) for cid in ids_componentes)
        nuevo_resultado = models.ResultadoDimension(
            inscripcion_id=inscripcion_id,
            dimension_id=compuesta.id,
            puntaje_directo=round(float(puntaje_total), 4),
        )
        db.add(nuevo_resultado)
        nuevos_resultados.append(nuevo_resultado)

    # flush y NO commit: la finalización del examen abarca cálculo,
    # interpretación y cambio de estado en una sola transacción, de forma que
    # un fallo a mitad no deje al alumno marcado como "Finalizado" sin
    # resultados y sin poder reintentar (B-9).
    db.flush()
    return nuevos_resultados
    
# ==========================================
# CRUD PARA BAREMOS
# ==========================================
def _rangos_se_solapan(a_min, a_max, b_min, b_max) -> bool:
    """True si [a_min, a_max] y [b_min, b_max] comparten algún punto.

    None se trata como "sin tope" en ese extremo (-inf / +inf), igual que lo
    interpreta `interpretar_resultados_alumno` al filtrar por edad.
    """
    lo = max(a_min if a_min is not None else float("-inf"), b_min if b_min is not None else float("-inf"))
    hi = min(a_max if a_max is not None else float("inf"), b_max if b_max is not None else float("inf"))
    return lo <= hi


def _validar_solapamiento_baremo(db: Session, datos, dimension_id: int, excluir_id: int = None):
    """📏 Fase 4 — dos baremos que cubren el mismo tramo de puntaje para el
    mismo público (dimensión, sexo, edad, población) son ambiguos: la consulta
    de `interpretar_resultados_alumno` no garantiza cuál de los dos devuelve,
    así que dos psicólogos podrían ver interpretaciones distintas del mismo
    puntaje según qué fila llegó primero. Antes nada lo impedía. Un baremo
    "unisex" (sexo=None) o "todas las edades" (edad_min/max=None) se trata
    como aplicable a cualquier valor de esa variable al comparar.
    """
    # 🆕 Un tramo desactivado no ocupa su rango para efectos de solapamiento:
    # es justamente lo que permite crear un tramo de reemplazo que cubra el
    # mismo puntaje que uno ya retirado por tener diagnósticos asociados.
    candidatos = db.query(models.Baremo).filter(
        models.Baremo.dimension_id == dimension_id,
        models.Baremo.activo == True,  # noqa: E712
    )
    if excluir_id is not None:
        candidatos = candidatos.filter(models.Baremo.id != excluir_id)

    for existente in candidatos.all():
        # 🆕 Criterio normativo — a diferencia de sexo/poblacion (donde NULL
        # es "aplica a cualquiera"), aquí NULL significa "el único criterio"
        # y un valor concreto es un LENTE DISTINTO: solo chocan dos tramos
        # que comparten exactamente el mismo criterio_normativo (ambos NULL,
        # o ambos con la misma etiqueta). Dos tramos con criterios distintos
        # pueden solaparse en rango de puntaje sin ser un duplicado real —
        # ver Baremo.criterio_normativo.
        if getattr(existente, "criterio_normativo", None) != getattr(datos, "criterio_normativo", None):
            continue
        sexo_solapa = existente.sexo is None or datos.sexo is None or existente.sexo == datos.sexo
        if not sexo_solapa:
            continue
        poblacion_solapa = (
            existente.poblacion is None or datos.poblacion is None or existente.poblacion == datos.poblacion
        )
        if not poblacion_solapa:
            continue
        edad_solapa = _rangos_se_solapan(existente.edad_min, existente.edad_max, datos.edad_min, datos.edad_max)
        if not edad_solapa:
            continue
        puntaje_solapa = _rangos_se_solapan(
            existente.puntaje_minimo, existente.puntaje_maximo, datos.puntaje_minimo, datos.puntaje_maximo
        )
        if not puntaje_solapa:
            continue

        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                f"Este tramo se solapa con el baremo #{existente.id} "
                f"({existente.puntaje_minimo}–{existente.puntaje_maximo}, nivel '{existente.nivel}') "
                "para el mismo público. Ajusta los rangos o el filtro de sexo/edad/población "
                "para que no se crucen."
            ),
        )


def crear_baremo(db: Session, baremo: schemas.BaremoCreate):
    """Guarda una regla de interpretación clínica en la base de datos"""
    _validar_solapamiento_baremo(db, baremo, baremo.dimension_id)
    nuevo_baremo = models.Baremo(**baremo.model_dump())
    db.add(nuevo_baremo)
    db.commit()
    db.refresh(nuevo_baremo)
    return nuevo_baremo

# ==========================================
# MOTOR DE INTERPRETACIÓN DE BAREMOS
# ==========================================
def interpretar_resultados_alumno(db: Session, inscripcion_id: int):
    """
    Cruza los Puntajes Directos del alumno con la tabla de Baremos,
    aplicando filtros dinámicos por edad, sexo y nivel escolar si existen.

    Devuelve un diccionario con tres claves:
        diagnosticos             -> lista de DiagnosticoAlumno creados
        dimensiones_sin_baremo   -> nombres de dimensiones que no se pudieron
                                    interpretar por falta de baremo aplicable
        interpretacion_completa  -> True solo si TODAS las dimensiones con
                                    puntaje encontraron baremo

    La distinción importa: antes, una dimensión sin baremo simplemente no
    generaba diagnóstico, y aguas arriba la ausencia de diagnósticos se leía
    como "este alumno no tiene riesgo". Un error de configuración se convertía
    en un resultado psicométrico (B-6).
    """
    # 1. Buscamos los datos del alumno mediante su inscripción
    inscripcion = db.query(models.Inscripcion).filter(models.Inscripcion.id == inscripcion_id).first()

    # Extraemos las variables sociodemográficas (si el modelo Inscripcion o Alumno las tiene)
    # Usamos getattr de forma segura por si aún no están implementados esos campos en la tabla inscripciones
    edad_alumno = getattr(inscripcion, "edad", None)
    sexo_alumno = getattr(inscripcion, "sexo", None)
    poblacion_alumno = getattr(inscripcion, "poblacion", None) # Ej: "Secundaria"

    # 🆕 Criterio normativo activo de la prueba (ver Prueba.criterio_normativo_
    # activo) — None para toda prueba existente, que es exactamente lo único
    # que hay hoy (Baremo.criterio_normativo IS NULL), cero cambio de
    # comportamiento. Cuando una prueba tiene más de un criterio cargado,
    # este valor decide cuál se usa para interpretar ESTA inscripción.
    criterio_normativo_activo = None
    evaluacion = db.query(models.Evaluacion).filter(models.Evaluacion.id == inscripcion.evaluacion_id).first()
    if evaluacion:
        prueba = db.query(models.Prueba).filter(models.Prueba.id == evaluacion.prueba_id).first()
        criterio_normativo_activo = getattr(prueba, "criterio_normativo_activo", None) if prueba else None

    # 2. Traemos los Puntajes Directos que calculamos previamente para este alumno
    puntajes_directos = db.query(models.ResultadoDimension).filter(
        models.ResultadoDimension.inscripcion_id == inscripcion_id
    ).all()

    diagnosticos_finales = []
    dimensiones_sin_baremo = []

    if not puntajes_directos:
        # Sin ResultadoDimension no hay nada que interpretar — típicamente
        # porque finalizar_examen_seguro se llamó sin ninguna RespuestaAlumno
        # guardada. Antes esto caía en el vacío-es-verdadero de Python: un
        # bucle que nunca corre deja dimensiones_sin_baremo=[] y eso se leía
        # como "todas las dimensiones encontraron baremo", liberando un
        # resultado vacío al alumno como si fuera limpio. Es el mismo tipo de
        # hueco que un baremo faltante (B-6): se retiene, nunca se asume
        # ausencia de riesgo.
        logger.warning(
            "Inscripción %s: no se generó ningún ResultadoDimension (examen "
            "finalizado sin respuestas registradas). El informe queda "
            "retenido para revisión profesional.",
            inscripcion_id,
        )

    for pd in puntajes_directos:
        # 3. Construimos la consulta inteligente de Baremo
        query_baremo = db.query(models.Baremo).filter(
            models.Baremo.dimension_id == pd.dimension_id,
            models.Baremo.puntaje_minimo <= pd.puntaje_directo,
            models.Baremo.puntaje_maximo >= pd.puntaje_directo,
            # 🆕 Solo el tramo del criterio normativo ACTIVO de la prueba
            # (ver arriba) — en una prueba con un único criterio cargado
            # (todas hoy) esto es simplemente "IS NULL", idéntico al filtro
            # implícito de siempre.
            models.Baremo.criterio_normativo == criterio_normativo_activo,
            # 🆕 Un tramo desactivado (ver Baremo.activo) ya no debe aplicarse
            # a interpretaciones nuevas — es la contraparte de bloquear su
            # edición: la única forma de "retirar" un tramo en uso es
            # apagarlo, y eso debe dejar de tenerlo en cuenta a partir de
            # ahora. Si nadie configuró un reemplazo, el puntaje queda sin
            # baremo aplicable (dimensiones_sin_baremo), no "sin riesgo".
            models.Baremo.activo == True,  # noqa: E712
        )

        # Si el Baremo de la base de datos exige un género, filtramos. Si es Nulo (unisex), pasa.
        if sexo_alumno is not None:
            query_baremo = query_baremo.filter(
                (models.Baremo.sexo == None) | (models.Baremo.sexo == sexo_alumno)
            )

        # Si exige un rango de edad, filtramos.
        if edad_alumno is not None:
            query_baremo = query_baremo.filter(
                (models.Baremo.edad_min == None) | (models.Baremo.edad_min <= edad_alumno)
            ).filter(
                (models.Baremo.edad_max == None) | (models.Baremo.edad_max >= edad_alumno)
            )

        # Si exige una población o nivel escolar específico, filtramos.
        if poblacion_alumno is not None:
            query_baremo = query_baremo.filter(
                (models.Baremo.poblacion == None) | (models.Baremo.poblacion == poblacion_alumno)
            )

        # Ante empate, gana el baremo marcado como alerta: si una misma
        # puntuación encaja en dos tramos solapados y uno exige revisión
        # profesional, el alumno debe quedar protegido, no depender del orden
        # arbitrario que devuelva el motor.
        baremo_encontrado = query_baremo.order_by(
            models.Baremo.es_alerta.desc(), models.Baremo.id
        ).first()

        # 4. Si encontramos coincidencia, guardamos su diagnóstico
        if baremo_encontrado:
            es_alerta_baremo = bool(baremo_encontrado.es_alerta)
            nuevo_diagnostico = models.DiagnosticoAlumno(
                inscripcion_id=inscripcion_id,
                dimension_id=pd.dimension_id,
                nivel=baremo_encontrado.nivel,
                tipo_puntaje_usado=baremo_encontrado.tipo_puntaje,
                valor_convertido=baremo_encontrado.valor_convertido,
                interpretacion=baremo_encontrado.interpretacion,
                # 🆕 Informe profesional — snapshot del contenido configurado
                # para este tramo y de su rango de corte (ver "Resultados
                # cuantitativos"), mismo criterio de instantánea que nivel/
                # interpretacion/es_alerta: editar el baremo después nunca
                # reescribe un informe ya emitido.
                conclusion_sugerida=baremo_encontrado.conclusion_sugerida,
                recomendacion_sugerida=baremo_encontrado.recomendacion_sugerida,
                rango_minimo=baremo_encontrado.puntaje_minimo,
                rango_maximo=baremo_encontrado.puntaje_maximo,
                # 🆕 Puntero histórico — ver DiagnosticoAlumno.baremo_id.
                baremo_id=baremo_encontrado.id,
                # Instantánea del indicador explícito de alerta (B-5)
                es_alerta=es_alerta_baremo,
                # 🆕 Trazabilidad de origen — aplicar_reglas_a_diagnosticos()
                # puede añadir ",regla" después si una regla TAMBIÉN dispara
                # sobre esta misma dimensión; nunca al revés.
                alerta_por=("baremo" if es_alerta_baremo else None),
            )
            db.add(nuevo_diagnostico)
            diagnosticos_finales.append(nuevo_diagnostico)
        else:
            # 5. Sin baremo aplicable NO hay diagnóstico posible. Lo dejamos
            #    registrado por nombre para que el profesional sepa qué
            #    configurar; quien llama decide retener el informe (B-6).
            dimension = db.query(models.Dimension).filter(
                models.Dimension.id == pd.dimension_id
            ).first()
            nombre = dimension.nombre if dimension else f"dimension_{pd.dimension_id}"
            dimensiones_sin_baremo.append(nombre)
            logger.warning(
                "Inscripción %s: la dimensión '%s' obtuvo un puntaje directo de %s "
                "pero ningún baremo aplicable lo cubre. El informe queda retenido "
                "para revisión profesional.",
                inscripcion_id, nombre, pd.puntaje_directo,
            )

    # flush y NO commit: ver nota en calcular_resultados_dimensiones (B-9)
    db.flush()

    return {
        "diagnosticos": diagnosticos_finales,
        "dimensiones_sin_baremo": dimensiones_sin_baremo,
        "interpretacion_completa": bool(puntajes_directos) and not dimensiones_sin_baremo,
    }


# ==========================================
# INTERPRETACIÓN POR PERFIL (modo_interpretacion='perfil')
# ==========================================
def interpretar_perfil_alumno(db: Session, inscripcion_id: int) -> dict:
    """Interpretación por PERFIL — la segunda mitad del punto de extensión
    que `Prueba.modo_interpretacion` reservaba sin algoritmo. En vez de leer
    cada dimensión contra un rango absoluto de `Baremo` (lo que hace
    `interpretar_resultados_alumno`), compara las dimensiones de ESTA MISMA
    inscripción entre sí: a la de mayor `puntaje_directo` le asigna el
    puesto 1º, a la siguiente el 2º, y así sucesivamente — el caso real que
    lo necesita es IPIP-RIASEC (código Holland = las 3 dimensiones con
    mayor puntaje), pero esto es completamente genérico, no sabe que RIASEC
    existe.

    No consulta `Baremo` en absoluto (por diseño — ver el comentario en
    `Prueba.modo_interpretacion` en models.py): `nivel`/`interpretacion` se
    generan automáticamente a partir del puesto, nunca de un tramo
    configurado. `es_alerta` queda SIEMPRE False — un puesto relativo nunca
    implica riesgo por sí solo; si una prueba de perfil futura necesitara
    alertar sobre algo, es trabajo de una `ReglaInterpretacion` (que sigue
    evaluándose después, exactamente igual que para el modo 'rango'), no de
    este algoritmo. `baremo_id` queda NULL (no hubo ningún tramo que citar).

    Como no depende de que exista un tramo configurado, SIEMPRE interpreta
    completo: `dimensiones_sin_baremo` queda vacío y `interpretacion_completa`
    es True mientras haya al menos un ResultadoDimension — "sin baremo
    aplicable" es un concepto que no existe en este modo.

    Misma forma de retorno que `interpretar_resultados_alumno()` (dict con
    `diagnosticos`/`dimensiones_sin_baremo`/`interpretacion_completa`) para
    que `interpretar_segun_modo()` pueda despachar a cualquiera de las dos
    sin que el llamador necesite saber cuál corrió.
    """
    puntajes_directos = db.query(models.ResultadoDimension).filter(
        models.ResultadoDimension.inscripcion_id == inscripcion_id
    ).all()

    ids_dimension = [pd.dimension_id for pd in puntajes_directos]
    nombres_por_dimension = {}
    if ids_dimension:
        nombres_por_dimension = dict(
            db.query(models.Dimension.id, models.Dimension.nombre)
            .filter(models.Dimension.id.in_(ids_dimension))
            .all()
        )

    total = len(puntajes_directos)
    ordenados = sorted(puntajes_directos, key=lambda pd: pd.puntaje_directo, reverse=True)

    if not puntajes_directos:
        # Mismo caso que en interpretar_resultados_alumno: sin ningún
        # ResultadoDimension que rankear (examen finalizado sin respuestas),
        # no hay nada que interpretar. El docstring de esta función ya
        # prometía interpretacion_completa=False aquí ("mientras haya al
        # menos un ResultadoDimension"); el return de abajo lo cumple ahora.
        logger.warning(
            "Inscripción %s: no se generó ningún ResultadoDimension para "
            "interpretación por perfil (examen finalizado sin respuestas "
            "registradas). El informe queda retenido para revisión "
            "profesional.",
            inscripcion_id,
        )

    diagnosticos_finales = []
    for puesto, pd in enumerate(ordenados, start=1):
        nombre_dimension = nombres_por_dimension.get(pd.dimension_id, f"dimension_{pd.dimension_id}")
        nuevo_diagnostico = models.DiagnosticoAlumno(
            inscripcion_id=inscripcion_id,
            dimension_id=pd.dimension_id,
            nivel=f"{puesto}º de {total}",
            interpretacion=(
                f"Esta dimensión obtuvo el {puesto}º puntaje más alto de las {total} "
                f"dimensiones evaluadas en esta aplicación (puntaje directo: "
                f"{float(pd.puntaje_directo):g}). La interpretación de una prueba de "
                "perfil es comparativa: describe la posición relativa de esta dimensión "
                "frente a las demás evaluadas en la misma aplicación, no un nivel absoluto."
            ),
            es_alerta=False,
            alerta_por=None,
            baremo_id=None,
        )
        db.add(nuevo_diagnostico)
        diagnosticos_finales.append(nuevo_diagnostico)

    # flush y NO commit: mismo criterio que interpretar_resultados_alumno (B-9)
    db.flush()

    return {
        "diagnosticos": diagnosticos_finales,
        "dimensiones_sin_baremo": [],
        "interpretacion_completa": bool(puntajes_directos),
    }


def interpretar_segun_modo(db: Session, inscripcion_id: int, prueba_id: int) -> dict:
    """Único punto que decide si una inscripción se interpreta por 'rango'
    (interpretar_resultados_alumno, el único modo que existía antes) o por
    'perfil' (interpretar_perfil_alumno) — lee Prueba.modo_interpretacion
    una sola vez, para que finalizar_examen_seguro y reprocesar nunca
    repitan esta misma rama por separado ni puedan desincronizarse entre sí.
    Un prueba_id sin fila (no debería ocurrir) cae de vuelta a 'rango', el
    comportamiento de siempre."""
    modo = db.query(models.Prueba.modo_interpretacion).filter(models.Prueba.id == prueba_id).scalar()
    if modo == models.MODO_INTERPRETACION_PERFIL:
        return interpretar_perfil_alumno(db, inscripcion_id)
    return interpretar_resultados_alumno(db, inscripcion_id)


# ==========================================
# REGLAS DE INTERPRETACIÓN (genéricas, configurables por prueba)
# ==========================================
# Sin eval(): el operador de cada regla se resuelve contra este diccionario
# fijo de funciones, nunca contra una cadena ejecutada.
_OPERADORES_REGLA = {
    ">": _operator.gt, ">=": _operator.ge,
    "<": _operator.lt, "<=": _operator.le,
    "==": _operator.eq, "!=": _operator.ne,
}


def crear_regla_interpretacion(db: Session, prueba_id: int, datos: schemas.ReglaInterpretacionCreate):
    nueva = models.ReglaInterpretacion(prueba_id=prueba_id, **datos.model_dump())
    db.add(nueva)
    db.commit()
    db.refresh(nueva)
    return nueva


def listar_reglas_de_prueba(db: Session, prueba_id: int):
    return db.query(models.ReglaInterpretacion).filter(
        models.ReglaInterpretacion.prueba_id == prueba_id
    ).all()


def obtener_regla_interpretacion(db: Session, regla_id: int):
    return db.query(models.ReglaInterpretacion).filter(models.ReglaInterpretacion.id == regla_id).first()


def actualizar_regla_interpretacion(db: Session, regla: models.ReglaInterpretacion, datos: schemas.ReglaInterpretacionUpdate):
    for campo, valor in datos.model_dump(exclude_unset=True).items():
        setattr(regla, campo, valor)
    db.commit()
    db.refresh(regla)
    return regla


def evaluar_reglas_interpretacion(db: Session, inscripcion_id: int, prueba_id: int) -> dict:
    """Evalúa las ReglaInterpretacion ACTIVAS de `prueba_id` contra los
    resultados ya calculados de esta inscripción (RespuestaAlumno /
    ResultadoDimension). No toca DiagnosticoAlumno — solo dice qué disparó y
    qué no se pudo evaluar. Genérica: no conoce PHQ-9 ni ningún otro
    instrumento, solo lee `ambito`/`pregunta_id`/`dimension_id`/`operador`/
    `valor_comparacion` de la fila de regla.

    Devuelve:
      disparadas -> [{"dimension_id": int|None, "regla_id": int, "descripcion": str}]
                     dimension_id=None (ámbito 'puntaje_total') significa
                     "aplica a TODAS las dimensiones de esta prueba".
      con_error  -> [str] reglas que no se pudieron evaluar (mala config, o
                     referencian algo que esta inscripción no tiene). Un error
                     de regla NUNCA se trata como "no disparó" — quien llama
                     debe tratarlo igual que dimensiones_sin_baremo: retiene
                     el informe en vez de asumir ausencia de riesgo.
    """
    reglas = db.query(models.ReglaInterpretacion).filter(
        models.ReglaInterpretacion.prueba_id == prueba_id,
        models.ReglaInterpretacion.activa == True,  # noqa: E712
    ).all()

    disparadas = []
    con_error = []
    if not reglas:
        return {"disparadas": disparadas, "con_error": con_error}

    total_prueba = None  # se calcula perezosamente, como mucho una vez
    prueba_cache = None  # idem — solo se consulta si hace falta invertir un ítem

    for regla in reglas:
        try:
            if regla.ambito == models.AMBITO_ITEM:
                if not regla.pregunta_id:
                    raise ValueError("regla de ámbito 'item' sin pregunta_id")
                pregunta = db.query(models.Pregunta).filter(
                    models.Pregunta.id == regla.pregunta_id,
                    models.Pregunta.prueba_id == prueba_id,
                ).first()
                if not pregunta:
                    raise ValueError(f"la pregunta {regla.pregunta_id} no pertenece a esta prueba")
                respuesta = db.query(models.RespuestaAlumno).filter(
                    models.RespuestaAlumno.inscripcion_id == inscripcion_id,
                    models.RespuestaAlumno.pregunta_id == regla.pregunta_id,
                ).first()
                if not respuesta or respuesta.puntaje_obtenido is None:
                    raise ValueError(f"no hay respuesta registrada para la pregunta {regla.pregunta_id}")
                # 🆕 Auditoría de motor (sesión 2026-10-13): antes se comparaba
                # el valor CRUDO (puntaje_obtenido), sin aplicar la inversión
                # de Pregunta.es_inverso — una regla de seguridad sobre un
                # ítem invertido habría disparado en la dirección equivocada.
                # invertir_valor_item es la MISMA fórmula que usa el cálculo
                # de dimensión (core.algoritmo.calcular_likert_generico), así
                # que ambos caminos ahora interpretan el ítem en la misma
                # dirección clínica.
                if pregunta.es_inverso and prueba_cache is None:
                    prueba_cache = db.query(models.Prueba).filter(models.Prueba.id == prueba_id).first()
                escala_min_prueba, escala_max_prueba = resolver_escala(prueba_cache)
                valor_real = invertir_valor_item(
                    pregunta, respuesta.puntaje_obtenido, escala_min_prueba, escala_max_prueba,
                )
                dimension_objetivo = pregunta.dimension_id

            elif regla.ambito == models.AMBITO_DIMENSION:
                if not regla.dimension_id:
                    raise ValueError("regla de ámbito 'dimension' sin dimension_id")
                dimension = db.query(models.Dimension).filter(
                    models.Dimension.id == regla.dimension_id,
                    models.Dimension.prueba_id == prueba_id,
                ).first()
                if not dimension:
                    raise ValueError(f"la dimensión {regla.dimension_id} no pertenece a esta prueba")
                resultado = db.query(models.ResultadoDimension).filter(
                    models.ResultadoDimension.inscripcion_id == inscripcion_id,
                    models.ResultadoDimension.dimension_id == regla.dimension_id,
                ).first()
                if not resultado:
                    raise ValueError(f"no hay puntaje calculado para la dimensión {regla.dimension_id}")
                valor_real = resultado.puntaje_directo
                dimension_objetivo = regla.dimension_id

            elif regla.ambito == models.AMBITO_PUNTAJE_TOTAL:
                if total_prueba is None:
                    total_prueba = db.query(func.sum(models.ResultadoDimension.puntaje_directo)).join(
                        models.Dimension, models.ResultadoDimension.dimension_id == models.Dimension.id
                    ).filter(
                        models.ResultadoDimension.inscripcion_id == inscripcion_id,
                        models.Dimension.prueba_id == prueba_id,
                        # 🆕 Excluye dimensiones 'suma_de_dimensiones': su
                        # puntaje YA es la suma de otras dimensiones de esta
                        # misma prueba (ver Dimension.tipo_calculo). Sin este
                        # filtro, "puntaje_total" contaría cada componente
                        # dos veces (una directo, otra dentro de su
                        # compuesta) en cuanto exista una sola dimensión
                        # compuesta en la prueba.
                        models.Dimension.tipo_calculo != models.TIPO_CALCULO_SUMA_DE_DIMENSIONES,
                        # 🆕 Excluye también 'conteo_sintomas': su puntaje es
                        # una CANTIDAD de ítems sobre un umbral, no una
                        # puntuación bruta — sumarla junto con dimensiones
                        # 'items' de otra escala (ej. NICHQ Vanderbilt mezcla
                        # conteos 0-9/0-15 con una dimensión de funcionamiento
                        # 0-8) produciría un "total" sin ningún significado.
                        models.Dimension.tipo_calculo != models.TIPO_CALCULO_CONTEO_SINTOMAS,
                    ).scalar() or 0
                valor_real = total_prueba
                dimension_objetivo = None

            else:
                raise ValueError(f"ámbito de regla desconocido: {regla.ambito!r}")

            comparador = _OPERADORES_REGLA.get(regla.operador)
            if comparador is None:
                raise ValueError(f"operador de regla desconocido: {regla.operador!r}")

            if comparador(float(valor_real), float(regla.valor_comparacion)):
                disparadas.append({
                    "dimension_id": dimension_objetivo,
                    "regla_id": regla.id,
                    "descripcion": regla.descripcion or f"Regla #{regla.id} ({regla.ambito})",
                })
        except (ValueError, TypeError) as exc:
            logger.warning(
                "Regla de interpretación #%s (prueba %s) no se pudo evaluar para "
                "la inscripción %s: %s. El informe queda retenido para revisión "
                "profesional en vez de asumir que no hay riesgo.",
                regla.id, prueba_id, inscripcion_id, exc,
            )
            con_error.append(f"Regla #{regla.id}: {exc}")

    return {"disparadas": disparadas, "con_error": con_error}


def aplicar_reglas_a_diagnosticos(db: Session, inscripcion_id: int, diagnosticos: list, resultado_reglas: dict) -> list:
    """Sobrepone lo que disparó evaluar_reglas_interpretacion() sobre los
    DiagnosticoAlumno ya construidos por baremo (interpretar_resultados_alumno).

    La regla de alerta SIEMPRE prevalece: solo puede encender es_alerta si
    estaba apagado, nunca lo apaga si el baremo ya lo había encendido (ver
    "la regla de alerta debe prevalecer"). Si una regla dispara sobre una
    dimensión que no tenía diagnóstico (porque no había baremo aplicable),
    se crea un DiagnosticoAlumno nuevo para esa dimensión en vez de perder la
    alerta — así reutiliza exactamente el mismo mecanismo que ya consumen
    detectar_niveles_de_alerta() y obtener_casos_pendientes_de_seguimiento(),
    sin duplicar lógica.

    Devuelve la lista de diagnósticos (la original + los que se hayan creado).
    """
    disparadas = resultado_reglas["disparadas"]
    if not disparadas:
        return diagnosticos

    por_dimension = {}
    aplica_a_todas = []
    for d in disparadas:
        if d["dimension_id"] is None:
            aplica_a_todas.append(d)
        else:
            por_dimension.setdefault(d["dimension_id"], []).append(d)

    diagnosticos_por_dimension = {d.dimension_id: d for d in diagnosticos}

    def _marcar(diag, disparos):
        origenes = set(diag.alerta_por.split(",")) if diag.alerta_por else set()
        origenes.add("regla")
        diag.es_alerta = True
        diag.alerta_por = ",".join(sorted(origenes))
        if diag.regla_activada_id is None:
            diag.regla_activada_id = disparos[0]["regla_id"]
        db.add(diag)

    for diag in diagnosticos:
        disparos = list(por_dimension.get(diag.dimension_id, [])) + list(aplica_a_todas)
        if disparos:
            _marcar(diag, disparos)

    # Dimensiones que una regla alertó pero que no llegaron a tener
    # diagnóstico por baremo (típicamente: sin baremo aplicable todavía).
    for dimension_id, disparos in por_dimension.items():
        if dimension_id in diagnosticos_por_dimension:
            continue
        nuevo = models.DiagnosticoAlumno(
            inscripcion_id=inscripcion_id,
            dimension_id=dimension_id,
            nivel="Alerta por regla de interpretación",
            interpretacion=disparos[0]["descripcion"],
            es_alerta=True,
            alerta_por="regla",
            regla_activada_id=disparos[0]["regla_id"],
        )
        db.add(nuevo)
        diagnosticos.append(nuevo)
        diagnosticos_por_dimension[dimension_id] = nuevo

    db.flush()
    return diagnosticos


# ==========================================
# CRUD PARA ENLACES DE ACCESO Y CAMPAÑAS
# ==========================================
def validar_fecha_fin_futura(fecha_fin: Optional[datetime]) -> None:
    """🆕 Rechaza una fecha de cierre que ya pasó (sesión 2026-10-29).

    Antes se aceptaba: una campaña creada con `fecha_fin` en el pasado nacía
    ya vencida (campana_esta_expirada() la daba por cerrada desde el primer
    segundo, sin ningún aviso) — pasó de verdad con una campaña de prueba.
    Vive aquí y no en el field_validator de EvaluacionCreate/Update a
    propósito: esos schemas también se construyen en memoria (tests,
    normalización de zona horaria) con fechas pasadas legítimas; la regla es
    de negocio de "crear/editar vía API", no de forma del dato.
    `fecha_fin` ya llega normalizada a UTC naive (ver fecha_fin_a_utc_naive),
    así que se compara contra utcnow() igual que campana_esta_expirada()."""
    if fecha_fin is not None and fecha_fin <= datetime.utcnow():
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="La fecha de cierre debe ser posterior a la fecha y hora actual.",
        )


def crear_campana_evaluacion(db: Session, evaluacion: schemas.EvaluacionCreate, creador_id: int):
    """
    Crea la campaña/salón. Ya NO genera links por sí sola —
    los alumnos se matriculan después con matricular_roster_en_campana.

    `fecha_fin` es opcional: si el psicólogo la envía, queda registrada y a
    partir de ese momento los candados de expiración de las rutas del examen
    (cargar preguntas, guardar, finalizar, campos socio y matriculación)
    empiezan a rechazar el acceso. Si no la envía, la campaña no caduca.
    """
    validar_fecha_fin_futura(evaluacion.fecha_fin)
    nueva_evaluacion = models.Evaluacion(
        institucion_id=evaluacion.institucion_id,
        prueba_id=evaluacion.prueba_id,
        nombre_grupo=evaluacion.nombre_grupo,
        creador_id=creador_id,
        fecha_fin=evaluacion.fecha_fin,
    )
    db.add(nueva_evaluacion)
    db.commit()
    db.refresh(nueva_evaluacion)
    return nueva_evaluacion


# ==========================================
# CICLO DE VIDA DE CAMPAÑA: EDITAR / CERRAR / CANCELAR
# ==========================================
def contar_inscripciones_de_campana(db: Session, evaluacion_id: int) -> int:
    """Cuántas Inscripcion existen para esta campaña (roster o anónimas, en
    cualquier estado) — decide si `prueba_id` todavía es editable (ver
    actualizar_campana): un link mágico ya emitido asume las preguntas y
    baremos de la prueba original, cambiarla después rompería su cadena de
    autoridad (guardar_respuestas_examen la revalida en cada respuesta)."""
    return db.query(models.Inscripcion).filter(
        models.Inscripcion.evaluacion_id == evaluacion_id
    ).count()


def bloquear_evaluacion_para_editar(db: Session, evaluacion_id: int):
    """Trae la `Evaluacion` de esta campaña con FOR UPDATE, forzando estado fresco.

    🔒 Candado de concurrencia (hallazgo de auditoría, sesión 2026-10-16):
    sin esto, `actualizar_campana` (cambio de `prueba_id`) y
    `matricular_roster_en_campana` podían intercalarse — ambas leían
    `contar_inscripciones_de_campana()`/`evaluacion.prueba_id` sin bloqueo
    de fila, así que un cambio de prueba y una matriculación simultáneos
    podían pasar cada uno su propio chequeo antes de que el otro confirmara:
    el candado 409 que existe específicamente para impedir "cambiar la
    prueba de una campaña con inscripciones" quedaba desactivado por pura
    ventana de tiempo, y el alumno recién matriculado terminaba resolviendo
    la prueba NUEVA en cada endpoint del flujo de examen (que siempre lee
    `evaluacion.prueba_id` en vivo, nunca un snapshot tomado al matricular).
    Mismo mecanismo que `bloquear_inscripcion_para_finalizar`/
    `bloquear_inscripcion_para_reprocesar` (ver sus docstrings):
    `with_for_update()` serializa ambas rutas sobre la misma fila de
    `Evaluacion`, y `populate_existing()` fuerza el estado post-lock si esta
    `Evaluacion` ya estaba cargada en el identity map de la Session por una
    consulta previa sin bloqueo (p. ej. `_obtener_campana_de_mi_institucion`
    en el router, que valida tenencia/tenant antes de llegar aquí).
    """
    return db.query(models.Evaluacion).filter(
        models.Evaluacion.id == evaluacion_id
    ).with_for_update().populate_existing().first()


def actualizar_campana(db: Session, evaluacion: "models.Evaluacion", datos: schemas.EvaluacionUpdate):
    """PATCH parcial de una campaña.

    Solo mientras `estado == 'activa'`: una campaña cerrada o cancelada no se
    "reabre" implícitamente empujando `fecha_fin` hacia el futuro — hoy no
    existe una acción de reapertura, y esto es deliberado.
    """
    if evaluacion.estado != models.ESTADO_CAMPANA_ACTIVA:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Esta campaña está cerrada o cancelada; no admite ediciones.",
        )

    datos_enviados = datos.model_dump(exclude_unset=True)

    # Mismo candado que al crear: editar no puede dejar la campaña vencida
    # "por la puerta de atrás". `fecha_fin: null` (quitar la caducidad) sigue
    # permitido. Para terminarla a propósito está cerrar_campana().
    # Solo si la fecha CAMBIA de verdad: una campaña ya vencida por fecha
    # (estado "activa", fecha_fin pasada — ver lib/campanas.ts en el
    # frontend) tiene que poder renombrarse reenviando su misma fecha. El
    # margen de 60 s cubre el <input type="datetime-local">, que trunca a
    # minutos lo que el backend guardó con segundos.
    nueva_fecha_fin = datos_enviados.get("fecha_fin")
    if nueva_fecha_fin is not None and (
        evaluacion.fecha_fin is None or abs((nueva_fecha_fin - evaluacion.fecha_fin).total_seconds()) >= 60
    ):
        validar_fecha_fin_futura(nueva_fecha_fin)

    if "prueba_id" in datos_enviados and datos_enviados["prueba_id"] != evaluacion.prueba_id:
        # 🔒 Ver bloquear_evaluacion_para_editar: serializa contra una
        # matriculación de roster concurrente sobre esta misma campaña.
        evaluacion = bloquear_evaluacion_para_editar(db, evaluacion.id)
        if contar_inscripciones_de_campana(db, evaluacion.id) > 0:
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail="No se puede cambiar la prueba de una campaña que ya tiene inscripciones.",
            )

    for campo, valor in datos_enviados.items():
        setattr(evaluacion, campo, valor)

    db.commit()
    db.refresh(evaluacion)
    return evaluacion


def _cerrar_o_cancelar_campana(
    db: Session, evaluacion: "models.Evaluacion", usuario_id: int,
    motivo: Optional[str], nuevo_estado: str,
):
    """Función interna compartida por cerrar_campana/cancelar_campana: ambas
    mueven `estado`, dejan constancia de quién y cuándo, y fuerzan
    `fecha_fin` a "ahora" (si no estaba ya en el pasado) para reutilizar tal
    cual el candado de expiración que ya protege los 7 puntos del flujo de
    examen y matricular_roster_en_campana — cero endpoints nuevos que
    revisar por expiración."""
    # 🆕 Cierre de gap — naive a propósito: se compara contra
    # evaluacion.fecha_fin (DateTime sin zona horaria, ver
    # campana_esta_expirada() más arriba) en la misma línea de abajo.
    ahora = datetime.utcnow()
    evaluacion.estado = nuevo_estado
    evaluacion.fecha_cierre = ahora
    evaluacion.cerrada_por_usuario_id = usuario_id
    evaluacion.motivo_cierre = motivo
    if evaluacion.fecha_fin is None or evaluacion.fecha_fin > ahora:
        evaluacion.fecha_fin = ahora

    db.commit()
    db.refresh(evaluacion)
    return evaluacion


def cerrar_campana(db: Session, evaluacion: "models.Evaluacion", usuario_id: int, motivo: Optional[str] = None):
    """Cierre normal del ciclo de vida: deja de aceptar actividad nueva sin
    tocar ningún resultado ya emitido. Seguro en cualquier momento, tenga o
    no inscripciones o resultados finalizados."""
    if evaluacion.estado != models.ESTADO_CAMPANA_ACTIVA:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Esta campaña ya está cerrada o cancelada.",
        )
    return _cerrar_o_cancelar_campana(db, evaluacion, usuario_id, motivo, models.ESTADO_CAMPANA_CERRADA)


def cancelar_campana(db: Session, evaluacion: "models.Evaluacion", usuario_id: int, motivo: str):
    """Anula una campaña que nunca llegó a producir resultados reales.

    Bloqueada (409) si ya existe alguna Inscripcion Finalizada — cancelar no
    es una vía para ocultar o alterar un resultado ya emitido; para una
    campaña que ya tiene resultados reales, la acción correcta es cerrar, no
    cancelar.
    """
    if evaluacion.estado != models.ESTADO_CAMPANA_ACTIVA:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Esta campaña ya está cerrada o cancelada.",
        )

    tiene_finalizados = db.query(models.Inscripcion.id).filter(
        models.Inscripcion.evaluacion_id == evaluacion.id,
        models.Inscripcion.estado == "Finalizado",
    ).first() is not None
    if tiene_finalizados:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Esta campaña ya tiene evaluaciones finalizadas; no se puede cancelar, solo cerrar.",
        )
    return _cerrar_o_cancelar_campana(db, evaluacion, usuario_id, motivo, models.ESTADO_CAMPANA_CANCELADA)


# ==========================================
# 🆕 TEXTO VERSIONADO DEL ASENTIMIENTO INFORMADO
# ==========================================
TIPOS_TEXTO_CONSENTIMIENTO = ("apoderado", "asentimiento")


def _filtro_ambito_texto(institucion_id, tipo_texto: str, nivel_educativo_id):
    """Las 3 condiciones que identifican un ÁMBITO de versiones (ver
    AsentimientoInformadoTexto): institución o global, tipo, nivel o general."""
    T = models.AsentimientoInformadoTexto
    return (
        T.institucion_id.is_(None) if institucion_id is None else T.institucion_id == institucion_id,
        T.tipo_texto == tipo_texto,
        T.nivel_educativo_id.is_(None) if nivel_educativo_id is None else T.nivel_educativo_id == nivel_educativo_id,
    )


def resolver_texto_consentimiento_activo(
    db: Session, institucion_id, tipo_texto: str = "asentimiento", nivel_educativo_id=None,
):
    """🆕 (sesión 2026-10-30) El modelo ACTIVO que le corresponde a un
    alumno, en este orden: el de su nivel educativo -> el general de su
    colegio -> el global. Devuelve (texto, origen) con origen "nivel" |
    "colegio" | "global", o (None, None). Nunca lanza — mismo criterio que
    antes: un dato de configuración faltante no debe romper el flujo."""
    T = models.AsentimientoInformadoTexto
    candidatos = []
    if institucion_id is not None and nivel_educativo_id is not None:
        candidatos.append(("nivel", institucion_id, nivel_educativo_id))
    if institucion_id is not None:
        candidatos.append(("colegio", institucion_id, None))
    candidatos.append(("global", None, None))
    for origen, inst, nivel in candidatos:
        texto = db.query(T).filter(
            *_filtro_ambito_texto(inst, tipo_texto, nivel), T.activo == True,  # noqa: E712
        ).first()
        if texto:
            return texto, origen
    return None, None


def resolver_asentimiento_texto_activo(db: Session, institucion_id: int, nivel_educativo_id=None):
    """La versión activa del ASENTIMIENTO que le corresponde a esta
    institución (y, si se pasa, a este nivel) — la propia si tiene una
    activa, o la global si no. Nunca lanza: si ni siquiera existe una global
    (una base recién creada sin `seeds.seed_asentimiento_informado_default`
    todavía corrido), devuelve None y el router responde 404 — mismo
    criterio que el resto del flujo de examen ante un dato de configuración
    faltante. Envoltorio de resolver_texto_consentimiento_activo, conservado
    con su firma de siempre para los llamadores existentes."""
    texto, _ = resolver_texto_consentimiento_activo(db, institucion_id, "asentimiento", nivel_educativo_id)
    return texto


def nivel_educativo_de_seccion(db: Session, seccion_id):
    """Sección -> Grado -> nivel_educativo_id (o None si no hay sección)."""
    if seccion_id is None:
        return None
    from modulos.superadmin import models as sa_models
    fila = db.query(sa_models.Grado.nivel_educativo_id).join(
        sa_models.Seccion, sa_models.Seccion.grado_id == sa_models.Grado.id,
    ).filter(sa_models.Seccion.id == seccion_id).first()
    return fila[0] if fila else None


def nivel_educativo_actual_de_persona(db: Session, persona_id):
    """Nivel de la matrícula académica ACTIVA (fecha_fin IS NULL) de un
    alumno, o None si no está matriculado."""
    if persona_id is None:
        return None
    from modulos.superadmin import models as sa_models
    matricula = db.query(sa_models.MatriculaAcademica).filter(
        sa_models.MatriculaAcademica.persona_id == persona_id,
        sa_models.MatriculaAcademica.fecha_fin.is_(None),
    ).first()
    return nivel_educativo_de_seccion(db, matricula.seccion_id) if matricula else None


def nivel_educativo_de_inscripcion(db: Session, inscripcion):
    """El nivel con el que se evalúa una inscripción: el snapshot de su
    sección (`Inscripcion.seccion_id`, fijado al matricular) y, si no lo
    tiene, la matrícula activa actual del alumno."""
    nivel = nivel_educativo_de_seccion(db, getattr(inscripcion, "seccion_id", None))
    if nivel is None:
        nivel = nivel_educativo_actual_de_persona(db, getattr(inscripcion, "persona_id", None))
    return nivel


def ubicacion_academica_actual(db: Session, persona_id):
    """(nivel_educativo_id, "Secundaria · 3° · A") de la matrícula activa de
    un alumno, o (None, None). Para el PDF del modelo de consentimiento."""
    from modulos.superadmin import models as sa_models
    fila = db.query(
        sa_models.NivelEducativo.id, sa_models.NivelEducativo.nombre,
        sa_models.Grado.nombre, sa_models.Seccion.nombre,
    ).join(
        sa_models.Grado, sa_models.Grado.nivel_educativo_id == sa_models.NivelEducativo.id,
    ).join(
        sa_models.Seccion, sa_models.Seccion.grado_id == sa_models.Grado.id,
    ).join(
        sa_models.MatriculaAcademica, sa_models.MatriculaAcademica.seccion_id == sa_models.Seccion.id,
    ).filter(
        sa_models.MatriculaAcademica.persona_id == persona_id,
        sa_models.MatriculaAcademica.fecha_fin.is_(None),
    ).first()
    if not fila:
        return None, None
    nivel_id, nivel, grado, seccion = fila
    return nivel_id, " · ".join(x for x in (nivel, grado, seccion) if x)


def crear_version_asentimiento_texto(
    db: Session, institucion_id, contenido: str, usuario_id: int,
    tipo_texto: str = "asentimiento", nivel_educativo_id=None,
):
    """Crea una nueva versión (global si institucion_id es None, propia de
    una institución si no) y desactiva la anterior del MISMO ámbito en la
    misma transacción -- nunca se edita una fila ya existente (ver
    AsentimientoInformadoTexto, el snapshot de Inscripcion depende de que el
    contenido de una versión ya usada no cambie jamás). `version` es
    MAX(version existente de este ámbito) + 1, calculado en aplicación.

    🆕 El ámbito ahora también es (tipo_texto, nivel_educativo_id) — cada
    combinación numera sus versiones por separado. Los valores por defecto
    reproducen exactamente el comportamiento anterior.

    🔒 Dos POST simultáneos para el MISMO ámbito pueden calcular el mismo
    `version` (hallazgo real de tests/test_patron_concurrencia_unitario.py,
    no una precaución de más) -- el que pierde la carrera choca contra el
    índice único parcial (ver AsentimientoInformadoTexto.__table_args__) y
    commit_o_conflicto lo traduce a 409 en vez de dejar dos versiones
    activo=True a la vez. Quien pierde simplemente reintenta: el POST es
    una acción administrativa manual (editar el texto legal), no algo que
    de verdad necesite un lock de fila para esta frecuencia de uso."""
    filtro = _filtro_ambito_texto(institucion_id, tipo_texto, nivel_educativo_id)
    anteriores = db.query(models.AsentimientoInformadoTexto).filter(*filtro)
    ultima_version = db.query(func.max(models.AsentimientoInformadoTexto.version)).filter(*filtro).scalar()
    anteriores.filter(models.AsentimientoInformadoTexto.activo == True).update(  # noqa: E712
        {"activo": False}, synchronize_session=False,
    )
    nueva = models.AsentimientoInformadoTexto(
        institucion_id=institucion_id,
        tipo_texto=tipo_texto,
        nivel_educativo_id=nivel_educativo_id,
        version=(ultima_version or 0) + 1,
        contenido=contenido,
        activo=True,
        creado_por_usuario_id=usuario_id,
    )
    db.add(nueva)
    commit_o_conflicto(
        db, "Otra versión se creó al mismo tiempo para este mismo ámbito. Vuelve a intentarlo.",
    )
    db.refresh(nueva)
    return nueva


def listar_versiones_asentimiento_texto(db: Session, institucion_id, tipo_texto=None):
    """Historial completo (activas e inactivas) de una institución (o de la
    global) -- para que el equipo administrativo pueda ver qué decía cada
    versión y desde cuándo, no solo la activa. 🆕 Incluye todos los niveles
    y, salvo que se filtre por `tipo_texto`, ambos tipos — el cliente agrupa
    por ámbito."""
    T = models.AsentimientoInformadoTexto
    q = db.query(T).filter(T.institucion_id.is_(None) if institucion_id is None else T.institucion_id == institucion_id)
    if tipo_texto is not None:
        q = q.filter(T.tipo_texto == tipo_texto)
    return q.order_by(T.tipo_texto, T.nivel_educativo_id, T.version.desc()).all()


def obtener_intake_previo_por_token(db: Session, token: str):
    """Qué datos de filiación YA tiene esta inscripción antes de que el
    estudiante escriba nada — para que el formulario de "Antes de comenzar"
    del frontend no vuelva a pedir lo que el colegio ya cargó al matricular
    por padrón (ver matricular_roster_en_campana, que copia nombre/edad/
    sexo/población de la Persona real). `nombre_conocido=True` marca que
    `nombre_completo` viene del padrón y no debería ser editable — mismo
    criterio que `validar_y_registrar_alumno` ya usa para decidir qué
    rellenar y qué no (`if inscripcion.persona_id is None and not
    inscripcion.nombre_completo`). Para un link anónimo, todo viene en
    None/False y el formulario se ve exactamente igual que hoy.

    Devuelve None si el token no existe, para que el router decida el 404
    igual que en el resto del flujo de examen."""
    inscripcion = db.query(models.Inscripcion).filter(models.Inscripcion.token == token).first()
    if not inscripcion:
        return None
    return schemas.IntakeExamenResponse(
        nombre_completo=inscripcion.nombre_completo,
        nombre_conocido=inscripcion.persona_id is not None and bool(inscripcion.nombre_completo),
        edad=inscripcion.edad,
        sexo=inscripcion.sexo,
        poblacion=inscripcion.poblacion,
        informante_tipo=inscripcion.informante_tipo,
        ya_acepto_asentimiento=inscripcion.acepto_asentimiento,
    )


def validar_y_registrar_alumno(
    db: Session, token: str, datos: schemas.RegistroAlumnoDemograficos,
    ip_asentimiento: str = None, institucion_id: int = None,
):
    if not datos.acepta_terminos:
        return None

    inscripcion = db.query(models.Inscripcion).filter(
        models.Inscripcion.token == token,
        # 🔁 C-5 — también "En Progreso": el alumno que recarga la página o
        # pierde la conexión REANUDA en lugar de quedarse fuera. "Finalizado"
        # sigue excluido (el router lo rechaza antes con un 409).
        models.Inscripcion.estado.in_(["Pendiente", "En Progreso"]),
        # 🔒 C-8: un cupo eliminado o anonimizado ya no admite alumnos
        models.Inscripcion.is_deleted == False,      # noqa: E712
        models.Inscripcion.is_anonymized == False,   # noqa: E712
    ).first()

    if not inscripcion:
        return None

    # 🔒 Solo se completa lo que falte, nunca se pisa. Al REANUDAR esto importa
    # doblemente: edad, sexo y población deciden qué baremo se aplica, y no
    # deben poder cambiarse a mitad de examen reenviando el formulario.
    # Si el link viene ligado a una Persona del padrón, su nombre es el oficial.
    if inscripcion.persona_id is None and not inscripcion.nombre_completo:
        inscripcion.nombre_completo = datos.nombre_completo

    if inscripcion.edad is None:
        inscripcion.edad = datos.edad
    if inscripcion.sexo is None:
        inscripcion.sexo = datos.sexo
    if inscripcion.poblacion is None:
        inscripcion.poblacion = datos.poblacion

    # 🆕 Informante — el router ya exigió informante_nombre cuando
    # informante_tipo != 'estudiante' (ver validar_y_comenzar_evaluacion);
    # aquí solo se completa lo que falte, mismo criterio que el resto de
    # este bloque: reanudar un examen "En Progreso" nunca reescribe quién
    # ya había quedado registrado como informante.
    if inscripcion.informante_tipo != models.INFORMANTE_ESTUDIANTE:
        if not inscripcion.informante_nombre and datos.informante_nombre:
            inscripcion.informante_nombre = datos.informante_nombre.strip()
        if not inscripcion.informante_relacion and datos.informante_relacion:
            inscripcion.informante_relacion = datos.informante_relacion.strip()

    inscripcion.estado = "En Progreso"

    # ✅ FIX CRÍTICO #2: Marcamos el asentimiento como aceptado
    # para no bloquear las rutas de rendición de examen (cargar, progreso, finalizar).
    #
    # 🏛️ B-11 — Y dejamos constancia de CUÁNDO y DESDE DÓNDE lo aceptó. Sin
    # esto, `acepto_asentimiento = True` era una afirmación sin respaldo:
    # ante una fiscalización de la Ley 29733 no se podía probar el momento ni
    # el origen de la aceptación. Solo se escribe la primera vez, para no
    # pisar el evento original si el alumno vuelve a pasar por aquí.
    if not inscripcion.acepto_asentimiento:
        inscripcion.acepto_asentimiento = True
        inscripcion.fecha_asentimiento = datetime.now(timezone.utc)
        inscripcion.ip_asentimiento = ip_asentimiento
        # 🆕 Snapshot de QUÉ versión del texto vio y aceptó el alumno (ver
        # AsentimientoInformadoTexto) -- se resuelve una sola vez, en este
        # mismo momento, nunca se vuelve a tocar después aunque el texto
        # activo cambie. institucion_id=None (llamador no lo pasó, o el
        # cupo es anónimo sin Evaluacion resoluble) deja esto en NULL sin
        # romper el resto del flujo -- un faltante de configuración no debe
        # impedir que el alumno rinda su examen.
        if institucion_id is not None:
            texto_activo = resolver_asentimiento_texto_activo(
                db, institucion_id, nivel_educativo_de_inscripcion(db, inscripcion),
            )
            if texto_activo:
                inscripcion.asentimiento_texto_version_id = texto_activo.id

    db.commit()
    db.refresh(inscripcion)
    return inscripcion

# ==========================================
# RENDERIZADO DEL EXAMEN AL ALUMNO
# ==========================================
def obtener_preguntas_por_token(db: Session, token: str):
    """Busca a qué prueba pertenece este token y devuelve todas las preguntas con sus opciones"""
    # 1. Validamos que el link mágico exista
    inscripcion = db.query(models.Inscripcion).filter(models.Inscripcion.token == token).first()
    if not inscripcion:
        return None
    
    # 2. Buscamos a qué campaña/evaluación pertenece
    evaluacion = db.query(models.Evaluacion).filter(models.Evaluacion.id == inscripcion.evaluacion_id).first()
    
    # 3. Traemos todas las preguntas de ese test ordenadas
    preguntas_db = db.query(models.Pregunta).filter(
        models.Pregunta.prueba_id == evaluacion.prueba_id
    ).order_by(models.Pregunta.orden).all()
    
    # 4. Empaquetamos todo con sus opciones correspondientes
    resultado_final = []
    for p in preguntas_db:
        opciones_db = db.query(models.OpcionRespuesta).filter(
            models.OpcionRespuesta.pregunta_id == p.id
        ).all()
        
        resultado_final.append({
            "id": p.id,
            "texto_pregunta": p.texto_pregunta,
            "tipo_respuesta": p.tipo_respuesta,
            # 🆕 Solo la CANTIDAD de valores esperados — nunca la clave de
            # corrección (respuesta_correcta_numerica se queda en el servidor).
            "cantidad_valores_numericos": (
                len(p.respuesta_correcta_numerica) if p.respuesta_correcta_numerica else None
            ),
            "opciones": [{"id": opt.id, "texto_opcion": opt.texto_opcion} for opt in opciones_db]
        })
        
    return resultado_final

# ==========================================
# GUARDADO DE RESPUESTAS ANTI-TRAMPAS
# ==========================================
def obtener_pares_validos_de_prueba(db: Session, prueba_id: int):
    """Mapa {(pregunta_id, opcion_id): valor_puntuacion} de UNA prueba concreta.

    Es la lista blanca contra la que se valida todo lo que manda el alumno. Se
    construye recorriendo la relación real en base de datos
    (prueba -> pregunta -> opción), nunca a partir de los IDs del cliente.
    """
    filas = db.query(
        models.Pregunta.id,
        models.OpcionRespuesta.id,
        models.OpcionRespuesta.valor_puntuacion,
    ).join(
        models.OpcionRespuesta, models.OpcionRespuesta.pregunta_id == models.Pregunta.id
    ).filter(
        models.Pregunta.prueba_id == prueba_id
    ).all()

    return {(pregunta_id, opcion_id): valor for pregunta_id, opcion_id, valor in filas}


def obtener_preguntas_numericas_de_prueba(db: Session, prueba_id: int):
    """Mapa {pregunta_id: respuesta_correcta_numerica} de los ítems
    Pregunta.tipo_respuesta='numerico' de UNA prueba concreta — lista blanca
    equivalente a `obtener_pares_validos_de_prueba`, pero para ítems de
    respuesta construida (sin OpcionRespuesta). Se construye recorriendo la
    relación real en BD (prueba -> pregunta), nunca a partir de lo que mande
    el alumno — misma garantía de seguridad."""
    filas = db.query(
        models.Pregunta.id, models.Pregunta.respuesta_correcta_numerica,
    ).filter(
        models.Pregunta.prueba_id == prueba_id,
        models.Pregunta.tipo_respuesta == models.TIPO_RESPUESTA_NUMERICO,
    ).all()
    return {pregunta_id: respuesta_correcta for pregunta_id, respuesta_correcta in filas}


def guardar_respuestas_examen(db: Session, token: str, respuestas: List[schemas.RespuestaAlumnoSubmit]):
    """
    Recibe las opciones marcadas por el alumno, verifica que cada par
    (pregunta, opción) pertenezca realmente a SU prueba, y guarda el puntaje
    que dice la base de datos. Si el alumno cambió de opinión, sobrescribe.

    🆕 También soporta ítems de respuesta numérica construida
    (Pregunta.tipo_respuesta='numerico', ver BARSIT "Razonamiento numérico")
    — el alumno no elige una OpcionRespuesta, escribe una lista de números
    que se compara aquí mismo contra `Pregunta.respuesta_correcta_numerica`.
    El puntaje (1/0) se calcula SIEMPRE en el servidor, nunca se confía en
    nada que mande el cliente más allá de los números que escribió — misma
    garantía que ya tenía el camino de opción múltiple.

    🔒 B-7 — Antes solo se comprobaba que `opcion_id` existiera en la tabla.
    Como `/examen/{token}/guardar` es público, un alumno podía enviar la
    opción de otra pregunta —o de otra prueba, o de otro colegio— y quedarse
    con su `valor_puntuacion`. Además las respuestas ajenas contaminaban el
    cálculo, porque los puntajes se agrupan por la dimensión de la pregunta.

    Ahora la cadena de autoridad se recorre entera en BD:
        inscripción -> evaluación/campaña -> prueba -> pregunta -> opción
    (o, para un ítem numérico: ... -> prueba -> pregunta -> clave numérica)
    y cualquier par fuera de esa cadena invalida el lote COMPLETO.
    """
    inscripcion = db.query(models.Inscripcion).filter(
        models.Inscripcion.token == token
    ).first()
    if not inscripcion:
        raise HTTPException(status_code=404, detail="Link inválido.")

    if not respuestas:
        return {"mensaje": "Se guardaron 0 respuestas correctamente."}

    # 1. Resolvemos la prueba a la que este alumno tiene derecho realmente.
    evaluacion = db.query(models.Evaluacion).filter(
        models.Evaluacion.id == inscripcion.evaluacion_id
    ).first()
    if not evaluacion or not evaluacion.prueba_id:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Esta evaluación no tiene una prueba asociada. Avisa a tu colegio.",
        )

    pares_validos = obtener_pares_validos_de_prueba(db, evaluacion.prueba_id)
    preguntas_numericas = obtener_preguntas_numericas_de_prueba(db, evaluacion.prueba_id)

    # 2. VALIDACIÓN COMPLETA ANTES DE TOCAR NADA.
    #    Se comprueba todo el lote primero: así un par manipulado no llega a
    #    borrar las respuestas anteriores (antes el DELETE iba delante del
    #    bucle y una opción inválida provocaba pérdida silenciosa de datos).
    vistos = set()
    for r in respuestas:
        if r.pregunta_id in vistos:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=(
                    f"El envío trae dos respuestas para la pregunta {r.pregunta_id}. "
                    "Manda una sola respuesta por pregunta."
                ),
            )
        vistos.add(r.pregunta_id)

        if r.pregunta_id in preguntas_numericas:
            # 🆕 Ítem 'numerico': no hay opcion_id que validar contra
            # OpcionRespuesta — se valida que el envío tenga la FORMA
            # correcta (ninguna opcion_id, y tantos números como la clave
            # espera). El contenido exacto se compara recién al guardar.
            respuesta_correcta = preguntas_numericas[r.pregunta_id]
            if (
                r.opcion_id is not None
                or r.respuesta_numerica is None
                or len(r.respuesta_numerica) != len(respuesta_correcta)
            ):
                logger.warning(
                    "Respuesta numérica rechazada en la inscripción %s: la pregunta %s "
                    "esperaba %s valor(es) numérico(s) y no se recibió el formato esperado.",
                    inscripcion.id, r.pregunta_id, len(respuesta_correcta),
                )
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=(
                        "Una de las respuestas enviadas no corresponde a esta evaluación. "
                        "No se guardó ninguna: recarga la página e inténtalo de nuevo."
                    ),
                )
            continue

        if r.respuesta_numerica is not None or (r.pregunta_id, r.opcion_id) not in pares_validos:
            # Mensaje deliberadamente uniforme: no distingue "la pregunta no es
            # de tu prueba" de "la opción es de otra pregunta", para no ayudar
            # a mapear el catálogo ajeno a base de sondeos.
            logger.warning(
                "Respuesta rechazada en la inscripción %s: el par "
                "(pregunta=%s, opcion=%s) no pertenece a la prueba %s.",
                inscripcion.id, r.pregunta_id, r.opcion_id, evaluacion.prueba_id,
            )
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=(
                    "Una de las respuestas enviadas no corresponde a esta evaluación. "
                    "No se guardó ninguna: recarga la página e inténtalo de nuevo."
                ),
            )

    # 3. Todo el lote es válido: sobrescribimos de forma atómica.
    try:
        preguntas_ids = [r.pregunta_id for r in respuestas]
        db.query(models.RespuestaAlumno).filter(
            models.RespuestaAlumno.inscripcion_id == inscripcion.id,
            models.RespuestaAlumno.pregunta_id.in_(preguntas_ids),
        ).delete(synchronize_session=False)

        nuevas_respuestas = []
        for r in respuestas:
            if r.pregunta_id in preguntas_numericas:
                respuesta_correcta = preguntas_numericas[r.pregunta_id]
                nuevas_respuestas.append(models.RespuestaAlumno(
                    inscripcion_id=inscripcion.id,
                    pregunta_id=r.pregunta_id,
                    opcion_id=None,
                    respuesta_numerica=r.respuesta_numerica,
                    # El puntaje SIEMPRE se calcula aquí contra la clave guardada
                    # en Pregunta: el alumno nunca lo envía ni lo elige.
                    puntaje_obtenido=1 if r.respuesta_numerica == respuesta_correcta else 0,
                ))
            else:
                nuevas_respuestas.append(models.RespuestaAlumno(
                    inscripcion_id=inscripcion.id,
                    pregunta_id=r.pregunta_id,
                    opcion_id=r.opcion_id,
                    # El puntaje SIEMPRE se relee de la BD: el alumno nunca lo envía.
                    puntaje_obtenido=pares_validos[(r.pregunta_id, r.opcion_id)],
                ))
        db.add_all(nuevas_respuestas)
        db.commit()
    except Exception:
        db.rollback()
        logger.error(
            "Fallo guardando el lote de respuestas de la inscripción %s; "
            "se revirtió entero.", inscripcion.id, exc_info=True,
        )
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="No se pudieron guardar tus respuestas. Vuelve a intentarlo.",
        )

    return {"mensaje": f"Se guardaron {len(respuestas)} respuestas correctamente."}

def obtener_links_por_campana(db: Session, evaluacion_id: int):
    return db.query(models.Inscripcion).filter(
        models.Inscripcion.evaluacion_id == evaluacion_id,
        models.Inscripcion.is_deleted == False,
    ).all()

def cargar_test_masivo(db: Session, datos: schemas.CargaMasivaTest):
    """Inserta de golpe las dimensiones, preguntas y opciones de un test.

    Todo o nada: antes hacía un `commit()` por cada dimensión/pregunta/lote de
    opciones, así que una fila inválida a mitad de la carga (p. ej. una
    opción de una prueba con `dimension_id` que en realidad no le pertenece)
    dejaba el test permanentemente medio cargado en la BD, sin forma de
    reintentar limpio. Ahora usa `flush()` para obtener los IDs que necesita
    encadenar y un único `commit()` al final, con rollback completo ante
    cualquier error — el mismo patrón que `aplicar_importacion`/
    `guardar_respuestas_examen`.
    """
    try:
        # 1. Guardar primero las dimensiones y guardar sus IDs
        dimensiones_creadas = {}
        dimensiones_tipo_calculo = {}
        for dim in datos.dimensiones:
            db_dim = models.Dimension(
                prueba_id=datos.prueba_id, nombre=dim.nombre, descripcion=dim.descripcion,
                tipo_calculo=dim.tipo_calculo,
            )
            db.add(db_dim)
            db.flush()
            dimensiones_creadas[dim.nombre] = db_dim.id  # Guardamos el nombre y su nuevo ID
            dimensiones_tipo_calculo[dim.nombre] = dim.tipo_calculo

        # 2. Guardar las preguntas y sus opciones
        for preg in datos.preguntas:
            # Buscamos el ID de la dimensión que acabamos de crear
            dim_id = dimensiones_creadas.get(preg.nombre_dimension) if preg.nombre_dimension else None

            # 🆕 Una dimensión 'suma_de_dimensiones' recién creada en ESTE
            # mismo lote tampoco puede recibir preguntas propias — mismo
            # guard que crear_pregunta/actualizar_pregunta, ver
            # _validar_dimension_no_compuesta_para_pregunta.
            if preg.nombre_dimension and dimensiones_tipo_calculo.get(preg.nombre_dimension) == models.TIPO_CALCULO_SUMA_DE_DIMENSIONES:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=(
                        f"La dimensión '{preg.nombre_dimension}' es 'suma_de_dimensiones': "
                        "no puede tener preguntas propias."
                    ),
                )

            db_preg = models.Pregunta(
                prueba_id=datos.prueba_id,
                dimension_id=dim_id,
                texto_pregunta=preg.texto_pregunta,
                orden=preg.orden,
                es_inverso=preg.es_inverso,  # 🆕 Pasa el flag para recodificación
                incluir_en_informe=preg.incluir_en_informe,
                etiqueta_informe=preg.etiqueta_informe,
            )
            db.add(db_preg)
            db.flush()

            # 3. Guardar las alternativas ("Sí", "No", etc.) de esta pregunta
            for opc in preg.opciones:
                db_opc = models.OpcionRespuesta(
                    pregunta_id=db_preg.id,
                    texto_opcion=opc.texto_opcion,
                    valor_puntuacion=opc.valor_puntuacion
                )
                db.add(db_opc)
            db.flush()

        db.commit()
    except Exception:
        db.rollback()
        raise

    return {"mensaje": "¡Test cargado exitosamente en la base de datos!", "total_preguntas": len(datos.preguntas)}

# ==========================================
# ARMADO DE DATOS PARA EL INFORME PDF
# ==========================================
_CARGO_POR_ROL = {
    "psicologo": "Psicólogo(a)",
    "admin": "Administrador(a)",
    "superadmin": "Equipo AuroraSoft",
    "colegio": "Colegio",
}


def _resolver_contenido(override, configurado):
    """Las tres fuentes del informe (ver 'FUENTES DE CONTENIDO DEL INFORME'):
    el texto del PROFESIONAL (`override`) siempre prevalece sobre el
    CONFIGURADO por la prueba/baremo; si ninguno existe, no hay texto — nunca
    se inventa contenido clínico (ver 'NO INVENTAR CONTENIDO')."""
    if override:
        return {"texto": override, "fuente": "profesional"}
    if configurado:
        return {"texto": configurado, "fuente": "configurado"}
    return {"texto": None, "fuente": None}


def _calcular_codigo_perfil(dimensiones_datos: list, top_n: int = 3):
    """"Código de perfil" (p. ej. el código Holland de IPIP-RIASEC): la
    primera letra del nombre de las `top_n` dimensiones con mayor puntaje,
    en orden descendente. Genérico — no conoce RIASEC ni ningún instrumento
    por nombre; asume que el nombre de cada dimensión empieza con la letra
    que la identifica (cierto para R-I-A-S-E-C), lo cual es responsabilidad
    de cómo se nombran las dimensiones al cargar la prueba, no del motor.
    None si no hay dimensiones con puntaje."""
    con_puntaje = [d for d in dimensiones_datos if d.get("puntaje") is not None]
    if not con_puntaje:
        return None
    ordenadas = sorted(con_puntaje, key=lambda d: d["puntaje"], reverse=True)
    return "".join(d["nombre"][0].upper() for d in ordenadas[:top_n] if d["nombre"])


def _prueba_puede_generar_alerta(db: Session, prueba_id: int) -> bool:
    """True si ESTA prueba tiene, hoy, algún mecanismo que pueda marcar
    es_alerta=True — un tramo de Baremo activo con es_alerta=True, o una
    ReglaInterpretacion activa (que puede forzar alerta aunque el baremo no
    la pida). Se deriva de la configuración real en vez de guardarse como
    columna: así nunca queda desactualizado si alguien agrega/desactiva un
    tramo o una regla después. Usado solo para decidir el TÍTULO/redacción
    de la sección 5.3 del informe (ver core/pdf_generator.py) — nunca para
    decidir si retener un resultado, eso sigue siendo es_alerta/alerta_por."""
    tiene_baremo_alerta = db.query(models.Baremo.id).join(
        models.Dimension, models.Baremo.dimension_id == models.Dimension.id
    ).filter(
        models.Dimension.prueba_id == prueba_id,
        models.Baremo.es_alerta == True,  # noqa: E712
        models.Baremo.activo == True,  # noqa: E712
    ).first() is not None
    if tiene_baremo_alerta:
        return True
    return db.query(models.ReglaInterpretacion.id).filter(
        models.ReglaInterpretacion.prueba_id == prueba_id,
        models.ReglaInterpretacion.activa == True,  # noqa: E712
    ).first() is not None


_PATRON_DIMENSION_AREA = re.compile(r"^(Intereses|Aptitudes)\s+—\s+(.+?):\s*(.+)$")


def _sin_duplicados(lista: list) -> list:
    """Quita duplicados EXACTOS preservando el orden de aparición — usado al
    agregar conclusion_sugerida/recomendacion_sugerida de varios tramos que
    pueden compartir el mismo texto (ej. CHASIDE: un área "Alta" a la vez en
    Intereses y en Aptitudes no debe imprimir su párrafo de carreras
    sugeridas dos veces en la sección 6)."""
    vistos = set()
    resultado = []
    for item in lista:
        if item not in vistos:
            vistos.add(item)
            resultado.append(item)
    return resultado


def _agrupar_dimensiones_por_area(dimensiones_datos: list):
    """Agrupa dimensiones PAREADAS por área a partir de la convención de
    nombre "Intereses — X: Nombre"/"Aptitudes — X: Nombre" (ver
    seed_chaside.py) — genérico por convención de nombre, igual que
    _calcular_codigo_perfil asume que el nombre de cada dimensión empieza
    con su letra identificadora; no conoce CHASIDE ni ningún instrumento.
    Devuelve None si ninguna dimensión sigue la convención (nada que
    agrupar) — solo se llama cuando Prueba.agrupar_dimensiones_por_area ya
    confirmó que esta prueba la usa.

    Un área cuenta como "destacada" en Intereses/Aptitudes cuando su nivel
    (el de DiagnosticoAlumno, vía el Baremo del tramo alcanzado) es
    exactamente "Alto" — el nombre de tramo con el que seed_chaside.py carga
    el nivel superior de cada sub-escala. Las "Áreas Vocacionales
    Compatibles" son la UNIÓN de ambos conjuntos, no la intersección: un
    área que destaca en Intereses O en Aptitudes ya es una opción vocacional
    a considerar, no solo cuando coincide en ambas — ver docstring del
    módulo para el ejemplo que confirma esta regla contra un informe real."""
    areas: dict = {}
    for d in dimensiones_datos:
        m = _PATRON_DIMENSION_AREA.match(d.get("nombre") or "")
        if not m:
            continue
        tipo, codigo, nombre_area = m.group(1), m.group(2).strip(), m.group(3).strip()
        entrada = areas.setdefault(codigo, {"nombre_area": nombre_area, "interes": None, "aptitud": None})
        if tipo == "Intereses":
            entrada["interes"] = d
        else:
            entrada["aptitud"] = d

    if not areas:
        return None

    destacadas_interes = [
        a["nombre_area"] for a in areas.values() if a["interes"] and a["interes"].get("nivel") == "Alto"
    ]
    destacadas_aptitud = [
        a["nombre_area"] for a in areas.values() if a["aptitud"] and a["aptitud"].get("nivel") == "Alto"
    ]
    return {
        "areas": list(areas.values()),
        "destacadas_interes": destacadas_interes,
        "destacadas_aptitud": destacadas_aptitud,
        "compatibles": _sin_duplicados(destacadas_interes + destacadas_aptitud),
    }


def _formatear_lista_areas(nombres: list) -> str:
    """"A, B y C" — enumeración en español con "y" antes del último elemento,
    para el párrafo de síntesis del perfil vocacional."""
    if not nombres:
        return ""
    if len(nombres) == 1:
        return nombres[0]
    return f"{', '.join(nombres[:-1])} y {nombres[-1]}"


def _resumen_perfil_vocacional(compatibles: list):
    """Párrafo de síntesis (se antepone a las conclusiones, sección 6)
    generado a partir de las áreas vocacionales compatibles YA calculadas —
    nunca texto fijo cargado en un Baremo, así que no requiere contenido
    adicional por prueba. None si ninguna área llegó a "Alto" en ningún lado
    (perfil sin predominancia clara) — en ese caso la sección 6 cae al resto
    de la cadena normal (profesional > agregado por tramo > nada)."""
    if not compatibles:
        return None
    return (
        "El perfil vocacional evidenciado en esta aplicación se orienta principalmente hacia "
        f"las áreas de {_formatear_lista_areas(compatibles)}, de acuerdo con la integración de "
        "los intereses y las aptitudes registrados en el instrumento aplicado."
    )


def _resumen_codigo_perfil(dimensiones_ordenadas: list, codigo_perfil: str, top_n: int = 3):
    """Párrafo de síntesis (se antepone a las conclusiones, sección 6) para
    pruebas modo_interpretacion='perfil' (ej. IPIP-RIASEC) — generado a
    partir del código de perfil YA calculado (`_calcular_codigo_perfil`) y
    de las `top_n` dimensiones con mayor puntaje de ESTA aplicación, ya
    ordenadas por `dimensiones_ordenadas`. Genérico: no sabe que "SIA" es
    un código Holland, solo arma la frase con los nombres reales de las
    dimensiones top-n. None si no hay ninguna dimensión con puntaje."""
    top = [d["nombre"] for d in dimensiones_ordenadas[:top_n] if d.get("puntaje") is not None]
    if not top:
        return None
    resto = f", seguida de {_formatear_lista_areas(top[1:])}" if len(top) > 1 else ""
    return (
        f"El perfil vocacional obtenido corresponde al código {codigo_perfil} "
        f"({'–'.join(top)}). El área de mayor predominancia es {top[0]}{resto}."
    )


def _resolver_items_informativos(db: Session, prueba_id: int, inscripcion_id: int) -> list:
    """Ítems marcados `Pregunta.incluir_en_informe=True` (típicamente sin
    dimension_id, por lo que no puntúan — ver calcular_resultados_
    dimensiones/algoritmo.calcular_likert_generico, que ya los ignora): su
    respuesta se imprime literal en el informe bajo `etiqueta_informe`, ej.
    el ítem 10 del PHQ-9 ("qué tan difícil le han resultado estos
    problemas"). Genérica — no sabe que "Indicador de Disfuncionalidad" es
    un concepto de PHQ-9, solo lee la bandera y la etiqueta configuradas en
    la Pregunta; para cualquier otra prueba (ninguna la usa hoy) es [] y no
    cambia nada en el informe."""
    filas = db.query(models.Pregunta, models.OpcionRespuesta).outerjoin(
        models.RespuestaAlumno,
        (models.RespuestaAlumno.pregunta_id == models.Pregunta.id)
        & (models.RespuestaAlumno.inscripcion_id == inscripcion_id),
    ).outerjoin(
        models.OpcionRespuesta, models.OpcionRespuesta.id == models.RespuestaAlumno.opcion_id,
    ).filter(
        models.Pregunta.prueba_id == prueba_id,
        models.Pregunta.incluir_en_informe == True,  # noqa: E712
    ).order_by(models.Pregunta.orden).all()

    return [
        {
            "etiqueta": pregunta.etiqueta_informe or pregunta.texto_pregunta,
            "pregunta_texto": pregunta.texto_pregunta,
            "respuesta_texto": opcion.texto_opcion if opcion else None,
        }
        for pregunta, opcion in filas
    ]


def _resolver_items_monitoreados_riesgo(db: Session, prueba_id: int, inscripcion_id: int) -> list:
    """Para cada ReglaInterpretacion de ámbito 'item' (ítems "centinela" de
    seguridad, ej. el ítem 9 del PHQ-9 — ideación autolítica), reporta
    SIEMPRE la respuesta real del estudiante en la sección 5.3 del informe,
    haya disparado la regla o no — en vez de mencionar el ítem solo cuando
    sí hay alerta. Genérica — no sabe qué instrumento es, solo lee
    ambito/pregunta_id/operador/valor_comparacion de la regla, igual que
    evaluar_reglas_interpretacion; para cualquier prueba sin reglas de
    ámbito 'item' es [] y no cambia nada en el informe."""
    reglas = db.query(models.ReglaInterpretacion).filter(
        models.ReglaInterpretacion.prueba_id == prueba_id,
        models.ReglaInterpretacion.activa == True,  # noqa: E712
        models.ReglaInterpretacion.ambito == models.AMBITO_ITEM,
    ).all()
    if not reglas:
        return []

    prueba_cache = None  # se consulta como mucho una vez, solo si hace falta invertir algún ítem

    resultado = []
    for regla in reglas:
        pregunta = db.query(models.Pregunta).filter(models.Pregunta.id == regla.pregunta_id).first()
        if not pregunta:
            continue
        fila = db.query(models.RespuestaAlumno, models.OpcionRespuesta).join(
            models.OpcionRespuesta, models.OpcionRespuesta.id == models.RespuestaAlumno.opcion_id,
        ).filter(
            models.RespuestaAlumno.inscripcion_id == inscripcion_id,
            models.RespuestaAlumno.pregunta_id == regla.pregunta_id,
        ).first()
        if not fila:
            continue  # sin respuesta registrada: evaluar_reglas_interpretacion ya marca esto como error aparte
        respuesta, opcion = fila
        comparador = _OPERADORES_REGLA.get(regla.operador)
        # 🆕 Auditoría de motor (sesión 2026-10-13): misma corrección que
        # evaluar_reglas_interpretacion — "disparo" aquí debe coincidir
        # SIEMPRE con lo que esa función decide, o la sección 5.3 del PDF
        # podría mostrar "no disparó" para un ítem que el sistema sí trató
        # como alerta en otro lado (o viceversa).
        if pregunta.es_inverso and prueba_cache is None:
            prueba_cache = db.query(models.Prueba).filter(models.Prueba.id == prueba_id).first()
        escala_min_prueba, escala_max_prueba = resolver_escala(prueba_cache)
        valor_corregido = invertir_valor_item(
            pregunta, respuesta.puntaje_obtenido, escala_min_prueba, escala_max_prueba,
        )
        disparo = (
            bool(comparador(float(valor_corregido), float(regla.valor_comparacion)))
            if comparador and valor_corregido is not None else False
        )
        resultado.append({
            "pregunta_texto": pregunta.texto_pregunta,
            "respuesta_texto": opcion.texto_opcion,
            "disparo": disparo,
        })
    return resultado


def obtener_datos_para_pdf(db: Session, inscripcion_id: int):
    """
    Junta TODO lo que pdf_generator.generar_pdf() necesita para el informe
    psicométrico profesional completo: filiación, motivo, técnicas,
    observación conductual, resultados (cuantitativos y cualitativos) por
    dimensión, análisis de riesgo, conclusiones y recomendaciones.

    Distingue en todo momento tres fuentes de contenido (A automático,
    B configurado por la prueba/baremo, C profesional) — nunca se inventa
    texto clínico que falte en ninguna de las tres.
    """
    inscripcion = db.query(models.Inscripcion).filter(models.Inscripcion.id == inscripcion_id).first()
    if not inscripcion or inscripcion.is_deleted:
        return None

    evaluacion = db.query(models.Evaluacion).filter(models.Evaluacion.id == inscripcion.evaluacion_id).first()
    prueba = db.query(models.Prueba).filter(models.Prueba.id == evaluacion.prueba_id).first() if evaluacion else None

    resultados = db.query(models.ResultadoDimension, models.Dimension).join(
        models.Dimension, models.ResultadoDimension.dimension_id == models.Dimension.id
    ).filter(models.ResultadoDimension.inscripcion_id == inscripcion_id).all()

    diagnosticos = db.query(models.DiagnosticoAlumno, models.Dimension).join(
        models.Dimension, models.DiagnosticoAlumno.dimension_id == models.Dimension.id
    ).filter(models.DiagnosticoAlumno.inscripcion_id == inscripcion_id).all()

    diag_por_dimension_id = {diag.dimension_id: diag for diag, _ in diagnosticos}

    ids_regla_activada = {diag.regla_activada_id for diag, _ in diagnosticos if diag.regla_activada_id}
    descripciones_regla = {}
    if ids_regla_activada:
        filas_regla = db.query(models.ReglaInterpretacion.id, models.ReglaInterpretacion.descripcion).filter(
            models.ReglaInterpretacion.id.in_(ids_regla_activada)
        ).all()
        descripciones_regla = {rid: desc for rid, desc in filas_regla}

    # 5.1/5.2 — Resultados cuantitativos Y cualitativos, ya fusionados por
    # dimensión: cada fila trae TODO lo que esa dimensión necesita mostrar,
    # sin que el PDF tenga que volver a cruzar tablas.
    dimensiones_datos = []
    for resultado, dim in resultados:
        diag = diag_por_dimension_id.get(dim.id)
        dimensiones_datos.append({
            "nombre": dim.nombre,
            "puntaje": resultado.puntaje_directo,
            "nivel": diag.nivel if diag else None,
            "rango_minimo": diag.rango_minimo if diag else None,
            "rango_maximo": diag.rango_maximo if diag else None,
            "interpretacion": diag.interpretacion if diag else None,
            "conclusion_sugerida": diag.conclusion_sugerida if diag else None,
            "recomendacion_sugerida": diag.recomendacion_sugerida if diag else None,
            "es_alerta": bool(diag.es_alerta) if diag else False,
            "alerta_por": diag.alerta_por if diag else None,
            "motivo_regla": descripciones_regla.get(diag.regla_activada_id) if diag and diag.regla_activada_id else None,
            "sin_baremo": diag is None,
            # 🆕 Solo se leen más abajo cuando Prueba.modo_interpretacion==
            # 'perfil' (ver Dimension.contenido_perfil_predominante en
            # models.py) — NULL para cualquier prueba en modo 'rango'.
            "_contenido_perfil_predominante": dim.contenido_perfil_predominante,
            "_contenido_perfil_secundaria": dim.contenido_perfil_secundaria,
            "_conclusion_sugerida_perfil": dim.conclusion_sugerida_perfil,
        })

    es_alerta_general = any(d["es_alerta"] for d in dimensiones_datos)

    # 🆕 "Diseño A" para pruebas de perfil (ver Prueba.ordenar_dimensiones_
    # por_puntaje en models.py) — pura presentación: reordena lo que ya se
    # calculó, nunca cambia qué se calculó ni cómo se interpretó.
    codigo_perfil = None
    if prueba and prueba.ordenar_dimensiones_por_puntaje:
        dimensiones_datos = sorted(
            dimensiones_datos,
            key=lambda d: d["puntaje"] if d["puntaje"] is not None else float("-inf"),
            reverse=True,
        )
        codigo_perfil = _calcular_codigo_perfil(dimensiones_datos)

    # 🆕 Contenido de informe para modo_interpretacion='perfil' (ver
    # Dimension.contenido_perfil_predominante en models.py) — interpretar_
    # perfil_alumno nunca consulta Baremo, así que el texto automático que
    # ya escribió en DiagnosticoAlumno.interpretacion es solo una frase de
    # posición ("1º de 6"); si la prueba configuró contenido propio por
    # dimensión, ese contenido SOBRESCRIBE la frase automática aquí — nunca
    # al revés, y nunca si el campo está vacío (fallback a la frase
    # automática, igual que "sin_baremo" cae al mensaje genérico en modo
    # 'rango'). resumen_perfil_codigo se antepone a las conclusiones más
    # abajo, mismo mecanismo que CHASIDE usa con perfil_vocacional.
    resumen_perfil_codigo = None
    if prueba and prueba.modo_interpretacion == models.MODO_INTERPRETACION_PERFIL and codigo_perfil:
        top_n = 3
        for indice, d in enumerate(dimensiones_datos):
            es_predominante = indice < top_n
            predominante_texto = d.pop("_contenido_perfil_predominante")
            secundaria_texto = d.pop("_contenido_perfil_secundaria")
            conclusion_perfil = d.pop("_conclusion_sugerida_perfil")
            contenido_texto = predominante_texto if es_predominante else secundaria_texto
            if contenido_texto:
                d["interpretacion"] = contenido_texto
            if es_predominante and conclusion_perfil:
                d["conclusion_sugerida"] = conclusion_perfil
        resumen_perfil_codigo = _resumen_codigo_perfil(dimensiones_datos, codigo_perfil, top_n=top_n)
    else:
        for d in dimensiones_datos:
            d.pop("_contenido_perfil_predominante", None)
            d.pop("_contenido_perfil_secundaria", None)
            d.pop("_conclusion_sugerida_perfil", None)

    puede_alertar = _prueba_puede_generar_alerta(db, prueba.id) if prueba else True

    # 🆕 Agrupamiento por área (ver Prueba.agrupar_dimensiones_por_area en
    # models.py) — pura síntesis de lo ya calculado, para pruebas con
    # dimensiones pareadas Intereses/Aptitudes (CHASIDE). None para
    # cualquier otra prueba, cero cambio en su informe.
    perfil_vocacional = (
        _agrupar_dimensiones_por_area(dimensiones_datos)
        if prueba and prueba.agrupar_dimensiones_por_area else None
    )

    # 1. DATOS DE FILIACIÓN — todo del snapshot de Inscripcion, nunca de la
    # ficha viva de Persona (ver "Fuente de los datos").
    persona = (
        db.query(superadmin_models.Persona).filter(superadmin_models.Persona.id == inscripcion.persona_id).first()
        if inscripcion.persona_id else None
    )
    ocupacion_actual = "Estudiante" if persona and persona.tipo_perfil == "estudiante" else None

    evaluador = None
    if inscripcion.revisado_por_usuario_id:
        evaluador = db.query(auth_models.Usuario.nombre_completo).filter(
            auth_models.Usuario.id == inscripcion.revisado_por_usuario_id
        ).scalar()
    elif evaluacion and evaluacion.creador_id:
        evaluador = db.query(auth_models.Usuario.nombre_completo).filter(
            auth_models.Usuario.id == evaluacion.creador_id
        ).scalar()

    filiacion = {
        "apellidos_nombres": inscripcion.nombre_completo or None,
        "edad": inscripcion.edad,
        "sexo": inscripcion.sexo,
        "poblacion": inscripcion.poblacion,
        # "Grado de instrucción" en nuestro contexto escolar = la sección
        # histórica de la aplicación (nunca la actual — ver CONTEXTO HISTÓRICO).
        "grado_instruccion": (
            f"{inscripcion.poblacion} - {inscripcion.grado_seccion_historico}"
            if inscripcion.poblacion and inscripcion.grado_seccion_historico
            else (inscripcion.grado_seccion_historico or inscripcion.poblacion or None)
        ),
        "ocupacion_actual": ocupacion_actual,
        # "Vive con" no tiene ninguna fuente de datos en el sistema hoy (no
        # se recoge en ningún formulario) — se deja señalado explícitamente
        # en vez de omitirse en silencio, ver "NO INVENTAR CONTENIDO".
        "vive_con": None,
        "fecha_evaluacion": inscripcion.fecha_resolucion,
        "evaluador": evaluador,
        # 🆕 Informante (ver models.INFORMANTES_VALIDOS) — None para el
        # autorreporte de siempre (informante_tipo='estudiante'), así el PDF
        # no imprime nada nuevo para ninguna prueba existente. Cuando SÍ es
        # padre/madre o docente, el informe debe decir quién respondió en
        # realidad — nunca puede leerse como si lo hubiera contestado el
        # propio estudiante.
        "informante_tipo": (
            inscripcion.informante_tipo if inscripcion.informante_tipo != models.INFORMANTE_ESTUDIANTE else None
        ),
        "informante_nombre": inscripcion.informante_nombre,
        "informante_relacion": inscripcion.informante_relacion,
    }

    # 2. MOTIVO DE CONSULTA — profesional > configurado por la prueba > nada.
    motivo_consulta = _resolver_contenido(
        inscripcion.motivo_consulta, prueba.motivo_consulta_default if prueba else None,
    )

    # 3. TÉCNICAS E INSTRUMENTO APLICADO — el instrumento SIEMPRE viene de
    # Prueba (nunca se reemplaza); las técnicas adicionales son profesional
    # (esta aplicación puntual) > configurado por la prueba > nada, mismo
    # orden de prioridad que motivo_consulta/observacion_conductual.
    tecnicas_texto = inscripcion.tecnicas_override or (prueba.tecnicas_adicionales if prueba else None)
    tecnicas = []
    if tecnicas_texto:
        tecnicas.extend(t.strip() for t in tecnicas_texto.split(",") if t.strip())
    tecnicas.append(prueba.nombre_publico if prueba else "Instrumento psicométrico")

    # 4. OBSERVACIÓN CONDUCTUAL — profesional > configurado por la prueba > nada
    # (🆕 ya no es "solo profesional" — ver Prueba.observacion_conductual_default).
    observacion_conductual = _resolver_contenido(
        inscripcion.observacion_conductual, prueba.observacion_conductual_default if prueba else None,
    )

    # 6/7. CONCLUSIONES / RECOMENDACIONES — profesional > agregado de lo
    # configurado por tramo en cada dimensión alcanzada > nada.
    # _sin_duplicados: un área "Alta" a la vez en Intereses y en Aptitudes
    # (ej. CHASIDE) puede cargar el MISMO texto de carreras sugeridas en los
    # dos tramos — se imprime una sola vez, no dos (ver su propio docstring).
    conclusiones_configuradas = _sin_duplicados(
        [d["conclusion_sugerida"] for d in dimensiones_datos if d["conclusion_sugerida"]]
    )
    recomendaciones_configuradas = _sin_duplicados(
        [d["recomendacion_sugerida"] for d in dimensiones_datos if d["recomendacion_sugerida"]]
    )

    # 🆕 Párrafo de síntesis del perfil vocacional (ver perfil_vocacional
    # arriba) — antepuesto a las conclusiones configuradas por tramo, nunca
    # las reemplaza. None para cualquier prueba sin agrupar_dimensiones_por_
    # area, cero cambio.
    if perfil_vocacional:
        resumen_perfil = _resumen_perfil_vocacional(perfil_vocacional["compatibles"])
        if resumen_perfil:
            conclusiones_configuradas = [resumen_perfil] + conclusiones_configuradas

    # 🆕 Mismo mecanismo que perfil_vocacional arriba, para modo_
    # interpretacion='perfil' (ver resumen_perfil_codigo calculado junto a
    # codigo_perfil). None para cualquier prueba en modo 'rango', cero
    # cambio.
    if resumen_perfil_codigo:
        conclusiones_configuradas = [resumen_perfil_codigo] + conclusiones_configuradas

    conclusiones = _resolver_contenido(
        inscripcion.conclusion_override,
        "\n\n".join(conclusiones_configuradas) if conclusiones_configuradas else None,
    )
    # 🆕 recomendacion_generica_default (ver Prueba en models.py) es el
    # último fallback, debajo del agregado por tramo — toda prueba que ya
    # usa recomendacion_sugerida por tramo (PHQ-9/GAD-7) nunca llega a leer
    # este campo porque esa lista nunca queda vacía para ellas; CHASIDE no
    # carga recomendacion_sugerida en ningún tramo (su recomendación no
    # depende del área destacada) y usa solo este campo.
    recomendaciones = _resolver_contenido(
        inscripcion.recomendaciones_override,
        (
            "\n\n".join(recomendaciones_configuradas) if recomendaciones_configuradas
            else (prueba.recomendacion_generica_default if prueba else None)
        ),
    )

    # 🆕 firma_cargo_override (ver models.py) SIEMPRE prevalece sobre la
    # etiqueta genérica por rol — mismo criterio "profesional > configurado"
    # que conclusion_override/recomendaciones_override, aplicado a la firma.
    profesional_cargo = inscripcion.firma_cargo_override
    if not profesional_cargo and inscripcion.revisado_por_usuario_id:
        rol_revisor = db.query(auth_models.Usuario.rol).filter(
            auth_models.Usuario.id == inscripcion.revisado_por_usuario_id
        ).scalar()
        profesional_cargo = _CARGO_POR_ROL.get(rol_revisor, rol_revisor)

    # 🆕 firma_nombre_override (auditoría de UX, sesión 2026-10-13): mismo
    # criterio "profesional > automático" — por defecto el nombre sigue
    # siendo el derivado de revisado_por_usuario_id/creador (ver `evaluador`
    # arriba), pero un profesional puede corregirlo para esta aplicación
    # puntual (ej. un supervisor firma en lugar de quien hizo la revisión
    # técnica). firma_dni_override no tiene ningún valor automático posible
    # (el DNI del profesional no vive en ningún otro dato del sistema) — si
    # está vacío, el PDF imprime el placeholder en blanco de siempre.
    profesional_responsable = inscripcion.firma_nombre_override or (
        evaluador if inscripcion.revisado_por_usuario_id else None
    )
    profesional_dni = inscripcion.firma_dni_override

    return {
        "inscripcion_id": inscripcion.id,
        "persona_id": inscripcion.persona_id,
        "prueba_id": prueba.id if prueba else None,
        "prueba_nombre_tecnico": prueba.nombre_tecnico if prueba else None,
        "nombre_prueba": prueba.nombre_publico if prueba else "Evaluación Psicométrica",
        "institucion_id": evaluacion.institucion_id if evaluacion else None,
        "evaluacion_id": evaluacion.id if evaluacion else None,
        "nombre_grupo": evaluacion.nombre_grupo if evaluacion else None,
        "fecha_aplicacion": inscripcion.fecha_resolucion,

        "filiacion": filiacion,
        "motivo_consulta": motivo_consulta,
        "tecnicas": tecnicas,
        "observacion_conductual": observacion_conductual,
        "dimensiones": dimensiones_datos,
        "items_informativos": _resolver_items_informativos(db, prueba.id, inscripcion_id) if prueba else [],

        "es_alerta": es_alerta_general,
        "puede_alertar": puede_alertar,
        "codigo_perfil": codigo_perfil,
        "perfil_vocacional": perfil_vocacional,
        "items_monitoreados_riesgo": _resolver_items_monitoreados_riesgo(db, prueba.id, inscripcion_id) if prueba else [],
        "estado_interpretacion": inscripcion.estado_interpretacion,
        "dimensiones_sin_baremo": (
            [d.strip() for d in inscripcion.dimensiones_sin_baremo.split(",")]
            if inscripcion.dimensiones_sin_baremo else []
        ),
        "reglas_con_error": (
            [r.strip() for r in inscripcion.reglas_con_error.split(",")]
            if inscripcion.reglas_con_error else []
        ),

        "conclusiones": conclusiones,
        "recomendaciones": recomendaciones,

        "estado_revision": {
            "resultado_visible_alumno": inscripcion.resultado_visible_alumno,
            "informe_liberado": inscripcion.informe_liberado,
            "profesional_responsable": profesional_responsable,
            "profesional_cargo": profesional_cargo,
            "profesional_dni": profesional_dni,
            "fecha_revision_informe": inscripcion.fecha_revision_informe,
            "interpretacion_clinica": inscripcion.interpretacion_clinica,
        },
    }

# ==========================================
# INFORMES — Informe/InformeVersion + STORAGE
# ==========================================
def obtener_informe_de_inscripcion(db: Session, inscripcion_id: int):
    return db.query(models.Informe).filter(models.Informe.inscripcion_id == inscripcion_id).first()


def verificar_puede_ver_informe(usuario_actual, institucion_id: int, informe_liberado: bool) -> None:
    """🔒 Remediación de permisos, paso 3 — candado único para el acceso
    STAFF a un informe puntual (PDF de descarga, metadata).

    - superadmin: siempre.
    - admin/psicologo: solo de su propia institución, como siempre.
    - colegio: de su propia institución, Y SOLO si `informe_liberado` es
      True — la MISMA condición que ya protege al propio alumno en
      GET /examen/{token}/informe-pdf (nunca menos estricta que esa). Antes,
      `colegio` descargaba/leía cualquier informe de su institución sin
      importar si seguía retenido para revisión — ver CLAUDE.md, "Permission
      surface: what the matrix test can't see".
    """
    if usuario_actual.rol == "superadmin":
        return
    if institucion_id != usuario_actual.institucion_id:
        raise HTTPException(status_code=403, detail="No puedes ver informes de otro colegio.")
    if usuario_actual.rol == "colegio" and not informe_liberado:
        raise HTTPException(
            status_code=403,
            detail="Este informe todavía está en revisión; no está disponible para el colegio.",
        )


def _bloquear_o_crear_informe_para_version(
    db: Session, inscripcion_id: int, institucion_id: int, prueba_id: int, persona_id,
) -> "models.Informe":
    """Get-or-create CON LOCK del `Informe` de esta inscripción, para que
    `generar_y_almacenar_informe` pueda leer `version_actual` fresco y sin
    carrera — mismo mecanismo exacto que
    `modulos.documentos.crud.bloquear_correlativo_para_emision` usa para
    `CorrelativoDocumento` (ver ese docstring).

    Sin esto, dos llamadas casi simultáneas a `generar_y_almacenar_informe`
    para la misma inscripción (doble clic en "Reemitir informe", o una
    reemisión manual coincidiendo con la tarea automática que ya dispara
    `finalizar_examen_seguro`, o con `reprocesar`) podían leer el MISMO
    `version_actual`, generar y subir a storage dos PDFs con el mismo
    `numero_version`, y solo chocar al `commit()` final contra
    `uq_edu_informe_versiones_numero` — con un `IntegrityError` crudo que
    ningún call site capturaba (ver generar_y_almacenar_informe).

    `with_for_update()` basta para el caso "el Informe ya existe" (se
    bloquea la fila hasta el commit/rollback de quien llamó). Para el caso
    "todavía no existe" (la primera versión de un informe nuevo), dos
    transacciones concurrentes pueden ver ambas "no existe" e intentar
    crearlo a la vez — se usa el mismo truco de SAVEPOINT
    (`db.begin_nested()`) que `bloquear_correlativo_para_emision`: la que
    pierde la carrera contra `uq_edu_informes_inscripcion` se recupera sin
    tumbar la transacción completa y vuelve a consultar YA con lock.
    """
    informe = db.query(models.Informe).filter(
        models.Informe.inscripcion_id == inscripcion_id
    ).with_for_update().populate_existing().first()
    if informe:
        return informe

    nested = db.begin_nested()
    try:
        informe = models.Informe(
            inscripcion_id=inscripcion_id,
            institucion_id=institucion_id,
            prueba_id=prueba_id,
            persona_id=persona_id,
            version_actual=0,  # generar_y_almacenar_informe lo sube a 1 en la misma transacción
        )
        db.add(informe)
        nested.commit()  # flush hacia la transacción exterior, NO un commit real — asigna informe.id
        return informe
    except IntegrityError:
        nested.rollback()
        return db.query(models.Informe).filter(
            models.Informe.inscripcion_id == inscripcion_id
        ).with_for_update().populate_existing().first()


def _iso(valor):
    return valor.isoformat() if isinstance(valor, datetime) else valor


def _snapshot_json_seguro(datos: dict) -> dict:
    """Copia TODO lo que arma obtener_datos_para_pdf() a un dict apto para la
    columna JSON (edu_informe_versiones.resultados_snapshot) — el contenido
    completo del informe (filiación, motivo, técnicas, observación,
    resultados, análisis de riesgo, conclusiones, recomendaciones) queda
    congelado aquí exactamente como se usó para construir ESA versión del
    PDF, sin importar qué cambie después en Prueba/Baremo/Inscripcion (ver
    "SNAPSHOT DE LA VERSIÓN").

    El tipo JSON de SQLAlchemy serializa con json.dumps, que no sabe
    convertir datetime: aquí se pasan a isoformat() explícitamente para no
    reventar el commit de la versión con un TypeError.
    """
    filiacion = dict(datos["filiacion"])
    filiacion["fecha_evaluacion"] = _iso(filiacion.get("fecha_evaluacion"))

    estado_revision = dict(datos["estado_revision"])
    estado_revision["fecha_revision_informe"] = _iso(estado_revision.get("fecha_revision_informe"))

    return {
        "inscripcion_id": datos["inscripcion_id"],
        "persona_id": datos["persona_id"],
        "prueba_id": datos["prueba_id"],
        "prueba_nombre_tecnico": datos["prueba_nombre_tecnico"],
        "nombre_prueba": datos["nombre_prueba"],
        "evaluacion_id": datos["evaluacion_id"],
        "nombre_grupo": datos["nombre_grupo"],
        "fecha_aplicacion": _iso(datos["fecha_aplicacion"]),
        "filiacion": filiacion,
        "motivo_consulta": datos["motivo_consulta"],
        "tecnicas": datos["tecnicas"],
        "observacion_conductual": datos["observacion_conductual"],
        "dimensiones": datos["dimensiones"],
        "items_informativos": datos.get("items_informativos", []),
        "es_alerta": datos["es_alerta"],
        "estado_interpretacion": datos["estado_interpretacion"],
        "dimensiones_sin_baremo": datos["dimensiones_sin_baremo"],
        "reglas_con_error": datos["reglas_con_error"],
        "items_monitoreados_riesgo": datos.get("items_monitoreados_riesgo", []),
        "perfil_vocacional": datos.get("perfil_vocacional"),
        "conclusiones": datos["conclusiones"],
        "recomendaciones": datos["recomendaciones"],
        "estado_revision": estado_revision,
    }


def obtener_anexos_evidencia_informe(db: Session, informe_id: int) -> list[dict]:
    """Evidencia del expediente clínico (EntradaExpediente tipo='evidencia')
    marcada para incluirse en el PDF de ESTE informe puntual —
    informe_id=informe_id e incluir_en_pdf_informe=True. Se consulta en CADA
    generación de versión (automática/reemisión/reproceso) y nunca se
    congela en el snapshot: una versión ya emitida es inmutable como
    siempre, pero si se agrega o se desmarca evidencia después, la PRÓXIMA
    versión (reemitir/reprocesar) sí refleja el estado actual.

    Devuelve el formato exacto que pdf_generator.generar_pdf() espera en
    contenido["anexos"]: tipo ('imagen'/'pdf', por content_type),
    titulo/descripcion/fecha y los bytes ya leídos de storage. Una
    evidencia cuyo archivo ya no está en storage se omite (con un warning
    en logs) en vez de reventar la generación de todo el informe por un
    solo anexo faltante."""
    entradas = db.query(superadmin_models.EntradaExpediente).filter(
        superadmin_models.EntradaExpediente.tipo_entrada == superadmin_crud.TIPO_ENTRADA_EVIDENCIA,
        superadmin_models.EntradaExpediente.informe_id == informe_id,
        superadmin_models.EntradaExpediente.incluir_en_pdf_informe.is_(True),
    ).order_by(superadmin_models.EntradaExpediente.fecha_entrada.asc()).all()
    if not entradas:
        return []

    from core.storage import obtener_storage_service
    storage = obtener_storage_service()
    anexos = []
    for entrada in entradas:
        try:
            contenido_bytes = storage.leer(entrada.storage_key)
        except FileNotFoundError:
            logger.warning(f"Evidencia {entrada.id} marcada para el PDF del informe {informe_id} pero sin archivo en storage; se omite.")
            continue
        tipo = "pdf" if (entrada.content_type or "").lower() == "application/pdf" else "imagen"
        anexos.append({
            "tipo": tipo,
            "titulo": entrada.titulo,
            "descripcion": entrada.descripcion,
            "fecha": entrada.fecha_entrada,
            "bytes": contenido_bytes,
        })
    return anexos


def construir_contenido_pdf(
    datos: dict, institucion=None, numero_version: int | None = None, motivo_version: str | None = None,
    anexos: list[dict] | None = None,
) -> dict:
    """Arma el dict `contenido` que espera pdf_generator.generar_pdf() a
    partir de lo que devuelve obtener_datos_para_pdf(). Centralizado para que
    el PDF ya almacenado y el generado al vuelo (fallback de inscripciones
    previas a esta fase) muestren exactamente el mismo contenido."""
    return {
        "institucion": {
            "nombre": institucion.nombre if institucion else None,
            "logo_url": institucion.logo_url if institucion else None,
        },
        "nombre_prueba": datos["nombre_prueba"],
        "version": {"numero": numero_version, "motivo": motivo_version},
        "filiacion": datos["filiacion"],
        "motivo_consulta": datos["motivo_consulta"],
        "tecnicas": datos["tecnicas"],
        "observacion_conductual": datos["observacion_conductual"],
        "dimensiones": datos["dimensiones"],
        "items_informativos": datos.get("items_informativos", []),
        "codigo_perfil": datos.get("codigo_perfil"),
        "perfil_vocacional": datos.get("perfil_vocacional"),
        "analisis_riesgo": {
            "es_alerta": datos["es_alerta"],
            "puede_alertar": datos.get("puede_alertar", True),
            "estado_interpretacion": datos["estado_interpretacion"],
            "dimensiones_sin_baremo": datos["dimensiones_sin_baremo"],
            "reglas_con_error": datos["reglas_con_error"],
            "items_monitoreados": datos.get("items_monitoreados_riesgo", []),
        },
        "conclusiones": datos["conclusiones"],
        "recomendaciones": datos["recomendaciones"],
        "firma": {
            "nombre": datos["estado_revision"]["profesional_responsable"],
            "cargo": datos["estado_revision"]["profesional_cargo"],
            "dni": datos["estado_revision"].get("profesional_dni"),
        },
        "interpretacion_clinica": datos["estado_revision"]["interpretacion_clinica"],
        "nombre_grupo": datos["nombre_grupo"],
        "fecha_aplicacion": datos["fecha_aplicacion"],
        "anexos": anexos or [],
    }


def generar_y_almacenar_informe(
    db: Session, inscripcion_id: int, usuario_id: int | None = None, motivo: str = "Emisión inicial",
    origen: str = "automatica", retracta_version_previa_liberada: bool = False,
) -> "models.InformeVersion":
    """Genera el PDF desde el snapshot de la aplicación, lo guarda en storage y registra su metadata.

    Se invoca automáticamente (motivo="Emisión inicial", usuario_id=None,
    origen="automatica") justo después de que `finalizar_examen_seguro`
    confirme su transacción — se genera y almacena SIEMPRE, tenga o no alerta
    el resultado, porque la autorización de descarga la deciden
    `informe_liberado`/`resultado_visible_alumno` en el momento de servirlo,
    nunca la mera existencia del archivo. También se invoca manualmente desde
    el endpoint de reemisión (origen="reemision": solo vuelve a renderizar lo
    que YA está en DiagnosticoAlumno) y desde el de reproceso
    (origen="reproceso": la llamada llega DESPUÉS de recalcular
    DiagnosticoAlumno contra baremos/reglas corregidos — ver
    reprocesar_informe_de_inscripcion). Nunca reescribe una versión
    existente: siempre inserta una InformeVersion nueva
    (numero_version = version_actual + 1) y solo entonces mueve el puntero —
    ver "VERSIONADO DEL INFORME".

    `retracta_version_previa_liberada` lo pasa SOLO reprocesar_informe_de_
    inscripcion, y SOLO cuando la inscripción estaba `informe_liberado=True`
    Y `resultado_visible_alumno=True` ANTES de recalcular, pero el resultado
    corregido vuelve a exigir revisión — es decir, cuando esta misma versión
    nueva reemplaza a una que el alumno pudo haber llegado a descargar. Ver
    docs/gaps-pendientes.md.
    """
    from core.pdf_generator import generar_pdf
    from core.storage import construir_storage_key, obtener_storage_service

    datos = obtener_datos_para_pdf(db=db, inscripcion_id=inscripcion_id)
    if not datos:
        raise ValueError(f"No hay datos para generar el informe de la inscripción {inscripcion_id}")

    institucion = db.query(inst_models.Institucion).filter(
        inst_models.Institucion.id == datos["institucion_id"]
    ).first()

    # El número de versión se resuelve ANTES de generar el PDF (no después)
    # para poder imprimirlo dentro del propio documento — "Informe v1"/"v2"
    # debe poder leerse en el PDF mismo, no solo inferirse de la metadata.
    #
    # `_bloquear_o_crear_informe_para_version` toma el lock de fila aquí
    # mismo y lo mantiene hasta el commit/rollback de abajo — todo lo que
    # sigue (generar el PDF, subirlo a storage, insertar la InformeVersion)
    # corre bajo ese lock, mismo criterio que `generar_version_documento`
    # ya usa para el correlativo+documento del módulo `documentos`.
    informe = _bloquear_o_crear_informe_para_version(
        db, inscripcion_id=inscripcion_id, institucion_id=datos["institucion_id"],
        prueba_id=datos["prueba_id"], persona_id=datos["persona_id"],
    )
    numero_version = informe.version_actual + 1

    try:
        anexos = obtener_anexos_evidencia_informe(db, informe.id)
        contenido_pdf = construir_contenido_pdf(
            datos, institucion=institucion, numero_version=numero_version, motivo_version=motivo, anexos=anexos,
        )
        pdf_buffer = generar_pdf(contenido_pdf)
        contenido_bytes = pdf_buffer.getvalue()
        checksum = hashlib.sha256(contenido_bytes).hexdigest()

        storage_key = construir_storage_key(datos["institucion_id"], informe.id, numero_version)

        obtener_storage_service().guardar(storage_key, contenido_bytes, "application/pdf")

        nombre_base = (datos["filiacion"].get("apellidos_nombres") or "evaluado").replace(" ", "_")
        version = models.InformeVersion(
            informe_id=informe.id,
            numero_version=numero_version,
            resultados_snapshot=_snapshot_json_seguro(datos),
            storage_key=storage_key,
            nombre_archivo=f"informe_{nombre_base}_v{numero_version}.pdf",
            mime_type="application/pdf",
            tamano_bytes=len(contenido_bytes),
            checksum_sha256=checksum,
            generado_por_usuario_id=usuario_id,
            motivo_version=motivo,
            origen=origen,
            retracta_version_previa_liberada=retracta_version_previa_liberada,
        )
        db.add(version)
        informe.version_actual = numero_version
        informe.actualizado_en = datetime.now(timezone.utc)
        db.commit()
    except Exception:
        # Libera el lock de fila limpiamente y deja `informe.version_actual`
        # sin la mutación a medias — importa especialmente para
        # `reprocesar_informe_de_inscripcion`, que reutiliza esta MISMA
        # sesión para seguir trabajando (reemitir_documentos_vinculados_a_
        # inscripcion) justo después de esta llamada.
        db.rollback()
        raise
    db.refresh(version)
    return version


@celery_app.task(bind=True, max_retries=3, default_retry_delay=60)
def tarea_generar_informe_en_segundo_plano(self, inscripcion_id: int):
    """Tarea Celery real (Fase 2, ver core/cola_tareas.py): abre su PROPIA
    sesión de DB.

    Un worker de Celery es un proceso distinto del que sirvió la petición
    HTTP que finalizó el examen, así que reutilizar la sesión de esa
    petición aquí no es una opción — mismo motivo por el que
    enviar_alerta_riesgo tampoco recibe `db`.

    🆕 PLAN_REMODELACION_ARQUITECTURA.md, fase 3 (bitácora) — la llamada
    real sigue siendo exactamente la misma; `ejecutar_tarea_registrada`
    solo envuelve con un RegistroTareaAsincrona antes/después, para que un
    operador pueda ver qué corrió sin grepear logs.

    🆕 Fase 2 (Celery) — a diferencia de la bitácora-solamente de arriba,
    el `except` de aquí ahora SÍ cambia: antes tragaba la excepción entera
    (no había nada del otro lado de BackgroundTasks que la recogiera); con
    un worker Celery real, `self.retry(exc=exc)` reintenta automáticamente
    (hasta `max_retries`) en vez de dejar la tarea fallida en silencio. El
    rollback/log de antes se mantiene igual, solo se agrega el retry.
    """
    from core.database import SessionLocal
    from core.tareas_async import ejecutar_tarea_registrada
    db = SessionLocal()
    try:
        ejecutar_tarea_registrada(
            db, tipo_tarea="generar_informe", referencia_id=inscripcion_id,
            fn=lambda: generar_y_almacenar_informe(
                db, inscripcion_id=inscripcion_id, usuario_id=None, motivo="Emisión inicial",
            ),
            job_id=self.request.id,
        )
    except Exception as exc:
        db.rollback()
        logger.error(
            "Fallo generando/almacenando el informe de la inscripción %s en segundo plano "
            "(intento %s/%s).",
            inscripcion_id, self.request.retries + 1, self.max_retries + 1, exc_info=True,
        )
        raise self.retry(exc=exc)
    finally:
        db.close()


def obtener_version_actual_de_informe(db: Session, informe: "models.Informe"):
    return db.query(models.InformeVersion).filter(
        models.InformeVersion.informe_id == informe.id,
        models.InformeVersion.numero_version == informe.version_actual,
    ).first()


def listar_versiones_de_informe(db: Session, informe_id: int):
    return db.query(models.InformeVersion).filter(
        models.InformeVersion.informe_id == informe_id
    ).order_by(models.InformeVersion.numero_version.asc()).all()


def leer_pdf_de_version(version: "models.InformeVersion") -> bytes:
    from core.storage import obtener_storage_service
    return obtener_storage_service().leer(version.storage_key)


def listar_historial_informes_de_persona(db: Session, persona_id: int, institucion_id: int):
    """Historial de informes de un estudiante, SIEMPRE acotado a una institución.

    Mismo candado que el resto de la ficha del alumno: un informe de un
    colegio no debe poder consultarse a través del persona_id desde otro
    colegio, aunque compartan el mismo `Persona` (no debería, pero
    institucion_id se re-verifica igual, defensa en profundidad).
    """
    filas = db.query(models.Informe, models.Inscripcion, models.Evaluacion, models.Prueba).join(
        models.Inscripcion, models.Informe.inscripcion_id == models.Inscripcion.id
    ).join(
        models.Evaluacion, models.Inscripcion.evaluacion_id == models.Evaluacion.id
    ).join(
        models.Prueba, models.Informe.prueba_id == models.Prueba.id
    ).filter(
        models.Informe.persona_id == persona_id,
        models.Informe.institucion_id == institucion_id,
    ).order_by(models.Inscripcion.fecha_resolucion.desc()).all()

    # 🆕 Precarga en bloque de qué inscripciones tienen al menos un
    # DiagnosticoAlumno con es_alerta=True — antes era una query POR FILA
    # del historial (hallazgo de auditoría, sesión 2026-10-16; bajo impacto
    # real porque está acotado al historial de UN alumno, pero igual de
    # fácil de evitar con una sola query IN en vez de N).
    inscripcion_ids = [inscripcion.id for _, inscripcion, _, _ in filas]
    inscripciones_con_alerta = {
        fila[0] for fila in db.query(models.DiagnosticoAlumno.inscripcion_id).filter(
            models.DiagnosticoAlumno.inscripcion_id.in_(inscripcion_ids),
            models.DiagnosticoAlumno.es_alerta == True,
        ).distinct().all()
    } if inscripcion_ids else set()

    resultado = []
    for informe, inscripcion, evaluacion, prueba in filas:
        resultado.append({
            "informe_id": informe.id,
            "inscripcion_id": inscripcion.id,
            "prueba_nombre_publico": prueba.nombre_publico,
            "nombre_grupo": evaluacion.nombre_grupo,
            "fecha_aplicacion": inscripcion.fecha_resolucion,
            "version_actual": informe.version_actual,
            "es_alerta": inscripcion.id in inscripciones_con_alerta,
            "estado_interpretacion": inscripcion.estado_interpretacion,
            "informe_liberado": inscripcion.informe_liberado,
        })
    return resultado


# ==========================================
# DASHBOARD AGREGADO POR SALÓN (CAMPAÑA)
# ==========================================
def _armar_alumnos_y_distribucion(db: Session, inscripciones: list, incluir_seccion: bool = False) -> dict:
    """Dado un conjunto de `Inscripcion` (de UNA campaña, o de varias — ver
    `obtener_dashboard_academico` más abajo), arma `alumnos`/`distribucion_
    por_dimension`/los tres totales. Extraído de `obtener_dashboard_campana`
    para que ambos dashboards compartan exactamente la misma agregación en
    vez de mantener dos copias que puedan divergir.

    `incluir_seccion=True` agrega el nombre de la sección a cada fila de
    alumno — solo tiene sentido cuando las inscripciones pueden venir de
    secciones distintas (el dashboard académico, ámbito "grado"); el
    dashboard de campaña nunca lo necesita, todos sus alumnos ya comparten
    la campaña (y normalmente la sección) por construcción.
    """
    alumnos = []
    distribucion: dict = {}

    # 🐢 Rendimiento (Fase 4) — una sola consulta para todas las inscripciones
    # en vez de una por alumno (ver comentario original en el dashboard de
    # campaña, motivo idéntico aquí: con un grado de 150 alumnos serían 150
    # round-trips si se hiciera dentro del bucle).
    inscripcion_ids = [insc.id for insc in inscripciones]
    filas_por_inscripcion: dict = {}
    if inscripcion_ids:
        todas_las_filas = db.query(models.ResultadoDimension, models.Dimension, models.DiagnosticoAlumno).join(
            models.Dimension, models.ResultadoDimension.dimension_id == models.Dimension.id
        ).outerjoin(
            models.DiagnosticoAlumno,
            and_(
                models.DiagnosticoAlumno.inscripcion_id == models.ResultadoDimension.inscripcion_id,
                models.DiagnosticoAlumno.dimension_id == models.ResultadoDimension.dimension_id,
            )
        ).filter(models.ResultadoDimension.inscripcion_id.in_(inscripcion_ids)).all()
        for resultado, dim, diag in todas_las_filas:
            filas_por_inscripcion.setdefault(resultado.inscripcion_id, []).append((resultado, dim, diag))

    nombres_seccion = {}
    if incluir_seccion:
        seccion_ids = {insc.seccion_id for insc in inscripciones if insc.seccion_id}
        if seccion_ids:
            nombres_seccion = {
                s.id: s.nombre for s in db.query(superadmin_models.Seccion).filter(
                    superadmin_models.Seccion.id.in_(seccion_ids)
                ).all()
            }

    for insc in inscripciones:
        lista_resultados = []
        for resultado, dim, diag in filas_por_inscripcion.get(insc.id, []):
            nivel = diag.nivel if diag else None
            lista_resultados.append({
                "dimension": dim.nombre,
                "puntaje_directo": resultado.puntaje_directo,
                "nivel": nivel,
                # 🆕 Ya estaba en memoria (mismo `diag` del join de arriba) —
                # ver nota en schemas.ResultadoAlumnoDashboard.
                "minimo": diag.rango_minimo if diag else None,
                "maximo": diag.rango_maximo if diag else None,
                "es_alerta": diag.es_alerta if diag else None,
            })
            if nivel:
                distribucion.setdefault(dim.nombre, {})
                distribucion[dim.nombre][nivel] = distribucion[dim.nombre].get(nivel, 0) + 1

        fila = {
            "inscripcion_id": insc.id,
            "persona_id": insc.persona_id,   # usado por obtener_dashboard_proceso para no re-consultar
            "nombre_completo": insc.nombre_completo,
            "estado": insc.estado,
            "fecha_resolucion": insc.fecha_resolucion,
            "resultados": lista_resultados,
        }
        if incluir_seccion:
            fila["seccion"] = (
                insc.grado_seccion_historico or nombres_seccion.get(insc.seccion_id) or None
            )
        alumnos.append(fila)

    total_inscritos = len(inscripciones)
    total_finalizados = sum(1 for i in inscripciones if i.estado == "Finalizado")

    return {
        "alumnos": alumnos,
        "distribucion_por_dimension": distribucion,
        "total_inscritos": total_inscritos,
        "total_finalizados": total_finalizados,
        "total_pendientes": total_inscritos - total_finalizados,
    }


def _calcular_estadisticas_dimension(resultados_dimension: list, modo_interpretacion: str) -> dict:
    """Estadística descriptiva de UNA dimensión a partir de los resultados ya
    resueltos de todos los alumnos del grupo (una entrada por alumno que SÍ
    tiene `ResultadoDimension` en esta dimensión — mismo shape que
    `ResultadoAlumnoDashboard`: dimension/puntaje_directo/nivel/minimo/
    maximo/es_alerta). n/media/mediana/mínimo/máximo se calculan siempre,
    sobre `puntaje_directo` tal cual (una dimensión compuesta o de conteo de
    síntomas sigue siendo un número como cualquier otro). La distribución
    por nivel y el conteo en alerta solo tienen sentido en modo_interpretacion
    'rango' — en 'perfil' (interpretar_perfil_alumno, ver docs/motor-scoring.md)
    nunca hay niveles categóricos ni alertas, así que quedan vacíos/None."""
    puntajes = [r["puntaje_directo"] for r in resultados_dimension]
    n = len(puntajes)
    ordenados = sorted(puntajes)
    if n == 0:
        media = mediana = minimo = maximo = None
    else:
        media = sum(puntajes) / n
        mitad = n // 2
        mediana = ordenados[mitad] if n % 2 else (ordenados[mitad - 1] + ordenados[mitad]) / 2
        minimo, maximo = ordenados[0], ordenados[-1]
    # El máximo del tramo de Baremo más alto alcanzado entre los alumnos del
    # grupo — ninguno tiene por qué haber caído en el mismo tramo, así que
    # se toma el mayor `maximo` visto, no el primero que aparezca.
    maximos_vistos = [r["maximo"] for r in resultados_dimension if r.get("maximo") is not None]
    maximo_escala = max(maximos_vistos) if maximos_vistos else None

    salida = {
        "n": n, "media": media, "mediana": mediana, "minimo": minimo, "maximo": maximo,
        "maximo_escala": maximo_escala, "distribucion": [], "en_alerta": None, "dominante_cantidad": None,
    }
    if modo_interpretacion == models.MODO_INTERPRETACION_PERFIL:
        return salida

    conteos: dict = {}
    for r in resultados_dimension:
        nivel = r.get("nivel")
        if nivel is None:
            continue
        actual = conteos.setdefault(nivel, {"cantidad": 0, "es_alerta": False})
        actual["cantidad"] += 1
        if r.get("es_alerta"):
            actual["es_alerta"] = True
    total_con_nivel = sum(v["cantidad"] for v in conteos.values())
    salida["distribucion"] = sorted(
        [
            {
                "nivel": nivel, "cantidad": v["cantidad"],
                "pct": (v["cantidad"] / total_con_nivel * 100) if total_con_nivel else 0.0,
                "es_alerta": v["es_alerta"],
            }
            for nivel, v in conteos.items()
        ],
        key=lambda f: -f["cantidad"],
    )
    salida["en_alerta"] = sum(1 for r in resultados_dimension if r.get("es_alerta"))
    return salida


def calcular_resumen_estadistico_dashboard(db: Session, prueba, alumnos: list) -> dict:
    """Resumen estadístico por dimensión para un dashboard (de campaña o
    académico) — ver ResumenEstadisticoPrueba en aurora_ui y la hoja
    "Resumen estadístico" de `generar_excel_dashboard` más abajo, ambos
    consumidores de esta misma estructura. Señales ESTRUCTURALES de la
    prueba (nunca por nombre/clave — mismo criterio que todo el motor de
    scoring, ver docs/motor-scoring.md): una prueba puede tener de 1 a 14
    dimensiones, interpretar por rango absoluto o por perfil comparativo
    (`modo_interpretacion`), ser clínica o no, y agrupar sus dimensiones por
    área (CHASIDE) o no.

    🔒 Se calcula SIEMPRE a partir de `alumnos` tal cual los devuelve
    `_armar_alumnos_y_distribucion` — ANTES de que el router aplique
    `redactar_resultados_dashboard_campana` para 'colegio'. Esa redacción
    solo vacía `alumnos[].resultados` (el diagnóstico puntual de cada
    alumno), nunca toca ningún agregado de nivel de grupo — mismo criterio
    exacto que ya protege `distribucion_por_dimension` ("ver participación/
    avance está bien, ver el diagnóstico puntual no"), así que este resumen
    sobrevive intacto para 'colegio' con el mismo criterio, sin necesitar su
    propia redacción.

    `prueba` puede ser None (una Evaluacion cuya Prueba ya no existe) — en
    ese caso se devuelven valores neutros en vez de reventar."""
    if not prueba:
        return {
            "modo_interpretacion": models.MODO_INTERPRETACION_RANGO,
            "agrupar_dimensiones_por_area": False,
            "tiene_alertas_posibles": False,
            "dimensiones": [],
            "areas_compatibles": [],
        }

    dimensiones_db = db.query(models.Dimension).filter(
        models.Dimension.prueba_id == prueba.id
    ).order_by(models.Dimension.id).all()  # Dimension no tiene columna `orden` propia — su id ascendente ya refleja el orden de carga de cada seed.

    resultados_por_dimension: dict = {}
    for alumno in alumnos:
        for r in alumno.get("resultados", []):
            resultados_por_dimension.setdefault(r["dimension"], []).append(r)

    dimensiones_stats = []
    for dim in dimensiones_db:
        stats = _calcular_estadisticas_dimension(resultados_por_dimension.get(dim.nombre, []), prueba.modo_interpretacion)
        dimensiones_stats.append({"nombre": dim.nombre, "tipo_calculo": dim.tipo_calculo, **stats})

    if prueba.modo_interpretacion == models.MODO_INTERPRETACION_PERFIL:
        # Frecuencia de "dimensión dominante" (la de mayor puntaje de CADA
        # alumno) — la lectura grupal de un ranking individual, ver
        # interpretar_perfil_alumno/docs/motor-scoring.md.
        conteo_dominante: dict = {}
        for alumno in alumnos:
            resultados = alumno.get("resultados", [])
            if not resultados:
                continue
            top = max(resultados, key=lambda r: r["puntaje_directo"])
            conteo_dominante[top["dimension"]] = conteo_dominante.get(top["dimension"], 0) + 1
        for d in dimensiones_stats:
            d["dominante_cantidad"] = conteo_dominante.get(d["nombre"], 0)

    areas_compatibles = []
    if prueba.agrupar_dimensiones_por_area:
        # Reusa la MISMA convención/función de agrupamiento que ya usa el
        # informe PDF (_agrupar_dimensiones_por_area) aplicada alumno por
        # alumno, tabulada a nivel de grupo — nunca una segunda
        # implementación de la convención de nombre "Intereses — X"/
        # "Aptitudes — X". Esa función espera {"nombre", "nivel"}, distinto
        # del {"dimension", "nivel"} que trae `resultados` aquí.
        frecuencia_area: dict = {}
        for alumno in alumnos:
            datos_adaptados = [{"nombre": r["dimension"], "nivel": r.get("nivel")} for r in alumno.get("resultados", [])]
            agrupado = _agrupar_dimensiones_por_area(datos_adaptados)
            if not agrupado:
                continue
            for area in agrupado["compatibles"]:
                frecuencia_area[area] = frecuencia_area.get(area, 0) + 1
        areas_compatibles = sorted(
            [{"area": area, "cantidad": cantidad} for area, cantidad in frecuencia_area.items()],
            key=lambda f: -f["cantidad"],
        )

    return {
        "modo_interpretacion": prueba.modo_interpretacion,
        "agrupar_dimensiones_por_area": prueba.agrupar_dimensiones_por_area,
        "tiene_alertas_posibles": _prueba_puede_generar_alerta(db, prueba.id),
        "dimensiones": dimensiones_stats,
        "areas_compatibles": areas_compatibles,
    }


def obtener_dashboard_campana(db: Session, evaluacion_id: int):
    """
    Arma el resumen completo de una campaña: cuántos alumnos han
    terminado, sus resultados por dimensión, y la distribución de
    niveles para graficar.
    """
    evaluacion = db.query(models.Evaluacion).filter(models.Evaluacion.id == evaluacion_id).first()
    if not evaluacion:
        return None

    prueba = db.query(models.Prueba).filter(models.Prueba.id == evaluacion.prueba_id).first()
    inscripciones = db.query(models.Inscripcion).filter(
    models.Inscripcion.evaluacion_id == evaluacion_id,
    models.Inscripcion.is_deleted == False,
    ).all()

    agregado = _armar_alumnos_y_distribucion(db, inscripciones)

    return {
        "evaluacion_id": evaluacion.id,
        "nombre_grupo": evaluacion.nombre_grupo,
        "nombre_prueba": prueba.nombre_publico if prueba else "N/A",
        # 🆕 Título corto para la UI — ver Prueba.nombre_corto en models.py.
        "nombre_corto_prueba": (prueba.nombre_corto or prueba.nombre_tecnico) if prueba else None,
        **agregado,
        # 🆕 Resumen estadístico por dimensión (ver ResumenEstadisticoPrueba
        # en aurora_ui y la hoja "Resumen estadístico" de generar_excel_
        # dashboard) — calculado ANTES de cualquier redacción para 'colegio',
        # ver docstring de calcular_resumen_estadistico_dashboard.
        "resumen_estadistico": calcular_resumen_estadistico_dashboard(db, prueba, agregado["alumnos"]),
        # 🆕 Para "Editar campaña" en el dashboard — ver DashboardCampanaResponse.
        "estado": evaluacion.estado,
        "fecha_fin": evaluacion.fecha_fin,
        "prueba_id": evaluacion.prueba_id,
        "institucion_id": evaluacion.institucion_id,  # se usa solo para el candado, no viaja en la respuesta
    }


def obtener_dashboard_academico(
    db: Session, prueba_id: int, seccion_id: int | None = None, grado_id: int | None = None,
    institucion_id: int | None = None,
):
    """Mismo tipo de resumen que `obtener_dashboard_campana`, pero por SALÓN,
    GRADO o INSTITUCIÓN COMPLETA en vez de por campaña puntual — cruza TODAS
    las campañas en que se aplicó `prueba_id`, no una sola. Responde "¿cómo
    está 3° B en PHQ-9?" (o "¿cómo está TODO el colegio en PHQ-9?") en vez
    de "¿cómo salió esta campaña de marzo?".

    Usa el snapshot histórico `Inscripcion.seccion_id` (la sección vigente
    al momento de ESA aplicación puntual — ver `matricular_roster_en_
    campana`), nunca la matrícula académica actual: mismo criterio que el
    resto del sistema favorece snapshot sobre estado vivo (ver "CONTEXTO
    HISTÓRICO" en docs/estudiante-nucleo-y-privacidad.md). Un alumno que
    cambió de sección entre dos aplicaciones de la misma prueba aparece bajo
    la sección que tenía en CADA aplicación, no bajo la actual — igual que
    ya pasa con el dashboard de campaña para `grado_seccion_historico`.

    🆕 Ámbito "institución" (sesión 2026-10-25, a pedido del usuario — ver
    docs/gaps-pendientes.md): antes esta función exigía exactamente uno de
    `seccion_id`/`grado_id`, así que ver una prueba a escala de todo el
    colegio (aplicada en varios grados a la vez) significaba revisar cada
    salón/grado por separado. Reusa la MISMA agregación (`_armar_alumnos_y_
    distribucion`/`calcular_resumen_estadistico_dashboard`) que ya corre
    para "grado" — cero lógica nueva de estadística, solo un filtro más
    amplio (`Evaluacion.institucion_id` en vez de `seccion_id`/`grado_id`).
    `incluir_seccion=True` reusa `grado_seccion_historico` tal cual (ver
    `_armar_alumnos_y_distribucion`), que ya combina grado+sección en un
    solo string — necesario acá porque, a diferencia del ámbito "grado"
    (una sola sección letra por fila alcanza), varios grados a la vez sí
    necesitan distinguir "3° A" de "4° A".

    Como mucho uno de `seccion_id`/`grado_id`/`institucion_id` debe venir
    con valor — lo valida el router, no esta función. Devuelve None si la
    prueba o el salón/grado/institución no existen.
    """
    prueba = db.query(models.Prueba).filter(models.Prueba.id == prueba_id).first()
    if not prueba:
        return None

    query = db.query(models.Inscripcion).join(
        models.Evaluacion, models.Inscripcion.evaluacion_id == models.Evaluacion.id
    ).filter(
        models.Evaluacion.prueba_id == prueba_id,
        models.Inscripcion.is_deleted == False,  # noqa: E712
    )

    if seccion_id is not None:
        seccion = db.query(superadmin_models.Seccion).filter(superadmin_models.Seccion.id == seccion_id).first()
        if not seccion:
            return None
        grado = db.query(superadmin_models.Grado).filter(superadmin_models.Grado.id == seccion.grado_id).first()
        resuelto_institucion_id = seccion.institucion_id
        nombre_ambito = f'{grado.nombre + " " if grado else ""}"{seccion.nombre}"'.strip()
        tipo_ambito = "seccion"
        id_ambito = seccion_id
        query = query.filter(models.Inscripcion.seccion_id == seccion_id)
        incluir_seccion = False  # un solo salón: la columna sería redundante en cada fila
    elif grado_id is not None:
        grado = db.query(superadmin_models.Grado).filter(superadmin_models.Grado.id == grado_id).first()
        if not grado:
            return None
        resuelto_institucion_id = grado.institucion_id
        nombre_ambito = grado.nombre
        tipo_ambito = "grado"
        id_ambito = grado_id
        secciones_del_grado = db.query(superadmin_models.Seccion.id).filter(
            superadmin_models.Seccion.grado_id == grado_id
        )
        query = query.filter(models.Inscripcion.seccion_id.in_(secciones_del_grado))
        incluir_seccion = True  # varias secciones combinadas: hace falta distinguirlas
    elif institucion_id is not None:
        institucion = db.query(inst_models.Institucion).filter(inst_models.Institucion.id == institucion_id).first()
        if not institucion:
            return None
        resuelto_institucion_id = institucion_id
        nombre_ambito = institucion.nombre
        tipo_ambito = "institucion"
        id_ambito = institucion_id
        query = query.filter(models.Evaluacion.institucion_id == institucion_id)
        incluir_seccion = True  # varios grados/secciones combinados: hace falta distinguirlos
    else:
        return None

    inscripciones = query.order_by(models.Inscripcion.fecha_resolucion.desc()).all()
    agregado = _armar_alumnos_y_distribucion(db, inscripciones, incluir_seccion=incluir_seccion)

    return {
        "prueba_id": prueba.id,
        "nombre_prueba": prueba.nombre_publico,
        # 🆕 Título corto para la UI — ver Prueba.nombre_corto en models.py.
        "nombre_corto_prueba": prueba.nombre_corto or prueba.nombre_tecnico,
        "ambito": {
            "tipo": tipo_ambito,
            "id": id_ambito,
            "nombre": nombre_ambito,
        },
        **agregado,
        # 🆕 Ver calcular_resumen_estadistico_dashboard / obtener_dashboard_
        # campana arriba — mismo criterio, calculado antes de la redacción.
        "resumen_estadistico": calcular_resumen_estadistico_dashboard(db, prueba, agregado["alumnos"]),
        "institucion_id": resuelto_institucion_id,  # se usa solo para el candado, no viaja en la respuesta
    }


def obtener_tendencia_seccion(db: Session, prueba_id: int, seccion_id: int):
    """Evolución de UN salón en el tiempo para `prueba_id` — a diferencia de
    `obtener_dashboard_academico` (que funde TODAS las campañas en un solo
    agregado), esto parte las mismas inscripciones por `evaluacion_id` y
    devuelve un punto por campaña, ordenado cronológicamente — responde
    "¿este salón mejoró o empeoró entre marzo y agosto?", algo que el
    agregado único no puede mostrar.

    Reusa `_armar_alumnos_y_distribucion`/`calcular_resumen_estadistico_
    dashboard` tal cual, una vez por campaña — ninguna lógica de agregación
    nueva, solo la misma función de siempre llamada varias veces en vez de
    sobre la lista entera. `fecha_creacion` de la Evaluacion es el ancla
    cronológica (siempre existe, a diferencia de `fecha_resolucion` por
    alumno); dentro de cada campaña los alumnos siguen viniendo ordenados
    como ya lo hace `obtener_dashboard_campana`.

    Deliberadamente solo `seccion_id` (nunca `grado_id`): un grado agrupa
    secciones que pudieron rendir la misma prueba en fechas distintas por
    campañas distintas, así que una sola línea de tiempo por grado mezclaría
    cohortes distintos — queda fuera de alcance, ver docs/frontend.md."""
    prueba = db.query(models.Prueba).filter(models.Prueba.id == prueba_id).first()
    if not prueba:
        return None
    seccion = db.query(superadmin_models.Seccion).filter(superadmin_models.Seccion.id == seccion_id).first()
    if not seccion:
        return None
    grado = db.query(superadmin_models.Grado).filter(superadmin_models.Grado.id == seccion.grado_id).first()
    nombre_ambito = f'{grado.nombre + " " if grado else ""}"{seccion.nombre}"'.strip()

    evaluaciones = db.query(models.Evaluacion).filter(
        models.Evaluacion.prueba_id == prueba_id,
        models.Evaluacion.id.in_(
            db.query(models.Inscripcion.evaluacion_id).filter(
                models.Inscripcion.seccion_id == seccion_id,
                models.Inscripcion.is_deleted == False,  # noqa: E712
            )
        ),
    ).order_by(models.Evaluacion.fecha_creacion).all()

    puntos = []
    for evaluacion in evaluaciones:
        inscripciones = db.query(models.Inscripcion).filter(
            models.Inscripcion.evaluacion_id == evaluacion.id,
            models.Inscripcion.seccion_id == seccion_id,
            models.Inscripcion.is_deleted == False,  # noqa: E712
        ).all()
        agregado = _armar_alumnos_y_distribucion(db, inscripciones)
        if agregado["total_finalizados"] == 0:
            continue  # campaña sin ningún resultado real para este salón -- nada que graficar
        puntos.append({
            "evaluacion_id": evaluacion.id,
            "nombre_grupo": evaluacion.nombre_grupo,
            "fecha": evaluacion.fecha_creacion,
            "total_finalizados": agregado["total_finalizados"],
            "resumen_estadistico": calcular_resumen_estadistico_dashboard(db, prueba, agregado["alumnos"]),
        })

    return {
        "prueba_id": prueba.id,
        "nombre_prueba": prueba.nombre_publico,
        "nombre_corto_prueba": prueba.nombre_corto or prueba.nombre_tecnico,
        "ambito": {"tipo": "seccion", "id": seccion_id, "nombre": nombre_ambito},
        "puntos": puntos,
        "institucion_id": seccion.institucion_id,  # se usa solo para el candado, no viaja en la respuesta
    }


_ALIAS_SEXO_A_ENUM = {
    "m": schemas.SexoEnum.MASCULINO.value,
    "h": schemas.SexoEnum.MASCULINO.value,
    "masculino": schemas.SexoEnum.MASCULINO.value,
    "hombre": schemas.SexoEnum.MASCULINO.value,
    "f": schemas.SexoEnum.FEMENINO.value,
    "mujer": schemas.SexoEnum.FEMENINO.value,
    "femenino": schemas.SexoEnum.FEMENINO.value,
}


def _normalizar_sexo_para_baremo(genero: Optional[str]) -> Optional[str]:
    """`Persona.genero` es texto libre — un colegio puede haberlo cargado como
    "M", "Hombre" o "masculino" — pero `Baremo.sexo` y
    `RegistroAlumnoDemograficos.sexo` (la otra vía de captura, para un examen
    anónimo) usan siempre los literales exactos de `schemas.SexoEnum`
    ("Masculino"/"Femenino"), y `interpretar_resultados_alumno` los compara
    por igualdad exacta de texto. Sin esta normalización, un alumno
    matriculado por padrón cuyo `genero` no coincide carácter por carácter
    nunca encuentra un tramo de baremo específico por sexo — aunque exista
    uno aplicable — y queda forzado en silencio a "sin baremo aplicable".
    Un valor no reconocido se deja tal cual, igual que antes: esta función
    solo amplía qué se reconoce, nunca descarta información."""
    if not genero:
        return genero
    return _ALIAS_SEXO_A_ENUM.get(genero.strip().lower(), genero)


def matricular_roster_en_campana(db: Session, evaluacion_id: int, persona_ids: List[int]):
    """
    Genera un link mágico NOMBRADO (ligado a una Persona ya registrada)
    por cada estudiante del padrón que se pase en persona_ids.
    
    Aplica filtros de consentimiento y buena práctica clínica, devolviendo
    tanto las inscripciones generadas como la lista de rechazados con su motivo.
    """
    # 1. Obtenemos la evaluación para conocer la prueba y la institución.
    # 🔒 Con FOR UPDATE — ver bloquear_evaluacion_para_editar: serializa
    # contra un cambio concurrente de `prueba_id` vía actualizar_campana,
    # para que ningún alumno quede matriculado bajo una prueba que el PATCH
    # de edición cambió a mitad de esta misma operación.
    evaluacion = bloquear_evaluacion_para_editar(db, evaluacion_id)
    if not evaluacion:
        return {"generados": [], "rechazados": [], "error": f"Evaluación ID {evaluacion_id} no encontrada."}

    # 🆕 Informante — se lee UNA sola vez (mismo criterio de evitar N+1 que
    # exigir_asentimiento_estudiante más abajo) y se copia como snapshot a
    # cada Inscripcion nueva: es un atributo de qué VERSIÓN de la prueba es
    # esta campaña, fijo desde que la campaña existe, nunca vuelto a leer de
    # Prueba después (igual que grado_seccion_historico).
    prueba_de_campana = db.query(models.Prueba).filter(models.Prueba.id == evaluacion.prueba_id).first()
    informante_tipo_campana = prueba_de_campana.informante_tipo if prueba_de_campana else models.INFORMANTE_ESTUDIANTE

    # 🔒 CANDADO DE EXPIRACIÓN POR FECHA LÍMITE (Paso 2.2)
    if campana_esta_expirada(evaluacion):
        return {
            "generados": [], 
            "rechazados": [], 
            "error": "No se pueden matricular alumnos: la campaña ya ha alcanzado su fecha de cierre."
        }

    # 🆕 Política institucional de consentimientos — se carga UNA sola vez
    # antes del bucle (nunca dentro, para no convertir esto en un N+1 por
    # cada alumno del padrón). Con el flag apagado (default de toda
    # institución existente) el chequeo de abajo nunca se activa.
    institucion = db.query(inst_models.Institucion).filter(
        inst_models.Institucion.id == evaluacion.institucion_id
    ).first()
    exigir_asentimiento_estudiante = bool(institucion and institucion.exigir_asentimiento_estudiante)

    inscripciones_creadas = []
    rechazados = []

    for pid in persona_ids:
        persona = db.query(superadmin_models.Persona).filter(superadmin_models.Persona.id == pid).first()
        if not persona:
            rechazados.append({"persona_id": pid, "motivo": "Persona no encontrada en la base de datos."})
            continue

        #   🔒 CANDADO MULTI-TENANT — EL FIX CRÍTICO
        # Nunca confíes en que el persona_id que llega por body pertenece a
        # la institución dueña de la campaña. Se valida explícitamente contra
        # el registro real en BD, no contra nada que venga del frontend.
        if persona.institucion_id != evaluacion.institucion_id:
            rechazados.append({
                "persona_id": pid,
                "motivo": "Esta persona no pertenece a la institución de esta campaña."
            })
            continue

        # Validar consentimiento — 🔒 B-8: acotado a la institución de la campaña,
        # para que un consentimiento de otro colegio no habilite esta matrícula.
        #
        # 🆕 Consentimientos del estudiante/apoderado — se pide la FILA, no
        # solo el booleano: se necesita su id para dejar constancia en la
        # inscripción de EXACTAMENTE qué consentimiento la amparó (punto 7-9
        # de esa fase). El consentimiento de APODERADO bloquea SIEMPRE; el
        # del propio estudiante se registra y se guarda si existe, y ADEMÁS
        # bloquea cuando `Institucion.exigir_asentimiento_estudiante` está
        # activo (política institucional configurable — ver el chequeo justo
        # debajo, antes esto era un "punto 6" sin implementar).
        consentimiento_apoderado = superadmin_crud.obtener_consentimiento_vigente(
            db, persona.id, evaluacion.institucion_id,
            tipo_consentimiento=superadmin_models.TIPO_CONSENTIMIENTO_APODERADO,
        )
        if not consentimiento_apoderado:
            rechazados.append({
                "persona_id": pid,
                "nombre": f"{persona.nombres} {persona.apellidos}",
                "motivo": "Sin consentimiento vigente."
            })
            continue
        consentimiento_asentimiento = superadmin_crud.obtener_consentimiento_vigente(
            db, persona.id, evaluacion.institucion_id,
            tipo_consentimiento=superadmin_models.TIPO_CONSENTIMIENTO_ASENTIMIENTO,
        )
        if exigir_asentimiento_estudiante and not consentimiento_asentimiento:
            rechazados.append({
                "persona_id": pid,
                "nombre": f"{persona.nombres} {persona.apellidos}",
                "motivo": "Sin asentimiento del propio estudiante vigente.",
            })
            continue

        # 🔒 C-15: un alumno no puede tener dos cupos en la misma campaña.
        # La restricción UNIQUE de la BD es la garantía real; esto la
        # convierte en un rechazo explicado en vez de un IntegrityError.
        ya_matriculado = db.query(models.Inscripcion).filter(
            models.Inscripcion.evaluacion_id == evaluacion_id,
            models.Inscripcion.persona_id == persona.id,
        ).first()
        if ya_matriculado:
            rechazados.append({
                "persona_id": pid,
                "nombre": f"{persona.nombres} {persona.apellidos}",
                "motivo": "Este alumno ya está matriculado en esta campaña.",
            })
            continue

        # 🆕 Candado de buena práctica clínica
        puede, motivo = puede_aplicar_prueba(db, persona.id, evaluacion.prueba_id, evaluacion.institucion_id)
        if not puede:
            rechazados.append({
                "persona_id": pid, 
                "nombre": f"{persona.nombres} {persona.apellidos}",
                "motivo": motivo
            })
            continue

        # Calculamos la edad a partir de la fecha de nacimiento, si la tiene
        edad_calculada = None
        if persona.fecha_nacimiento:
            hoy = date.today()
            nacimiento = persona.fecha_nacimiento.date() if hasattr(persona.fecha_nacimiento, "date") else persona.fecha_nacimiento
            edad_calculada = hoy.year - nacimiento.year - ((hoy.month, hoy.day) < (nacimiento.month, nacimiento.day))

        # 🆕 Núcleo del estudiante — Punto 8: un alumno ya matriculado por
        # padrón no debería tener que re-ingresar sus datos al abrir el
        # examen. `nombre_completo`/`edad`/`sexo` ya venían de la Persona;
        # ahora también se deriva `poblacion` de su nivel educativo VIGENTE
        # (no de lo que el alumno tipeé) cuando el nombre del nivel coincide
        # con una de las poblaciones que entiende el motor de baremos
        # (Primaria/Secundaria/Universitario). Si no coincide o el alumno no
        # tiene matrícula académica activa, queda en None — el formulario del
        # examen sigue pidiéndola, como hasta ahora, sin romper nada.
        #
        # También se copia la sección vigente como snapshot INMUTABLE: si el
        # alumno cambia de sección después, esta inscripción ya aplicada no
        # se entera (ver comentario en Inscripcion.seccion_id).
        poblacion_derivada = None
        seccion_id_actual = None
        grado_seccion_snapshot = None
        matricula = superadmin_crud.obtener_matricula_activa(db=db, persona_id=persona.id)
        if matricula:
            info = superadmin_crud._describir_matricula(db=db, matricula=matricula)
            seccion_id_actual = info["seccion_id"]
            if info["grado"] or info["seccion"]:
                grado_seccion_snapshot = f'{info["grado"] or ""} {info["seccion"] or ""}'.strip() or None
            nivel_nombre = (info["nivel_educativo"] or "").strip().lower()
            for valor in schemas.PoblacionEnum:
                if valor.value.lower() == nivel_nombre:
                    poblacion_derivada = valor.value
                    break

        nueva_inscripcion = models.Inscripcion(
            evaluacion_id=evaluacion_id,
            persona_id=persona.id,
            token=str(uuid_lib.uuid4()),
            nombre_completo=f"{persona.nombres} {persona.apellidos}",
            edad=edad_calculada,
            sexo=_normalizar_sexo_para_baremo(persona.genero),
            poblacion=poblacion_derivada,
            seccion_id=seccion_id_actual,
            grado_seccion_historico=grado_seccion_snapshot,
            informante_tipo=informante_tipo_campana,
            # 🆕 Consentimientos del estudiante/apoderado — snapshot
            # INMUTABLE de qué consentimiento amparó esta aplicación
            # concreta. Una revocación posterior no reescribe estas
            # columnas: la aplicación sigue señalando al mismo registro
            # histórico que estaba vigente cuando se creó (puntos 7-9).
            consentimiento_apoderado_id=consentimiento_apoderado.id,
            consentimiento_asentimiento_id=(
                consentimiento_asentimiento.id if consentimiento_asentimiento else None
            ),
            estado="Pendiente",
        )
        db.add(nueva_inscripcion)
        inscripciones_creadas.append(nueva_inscripcion)

    # 🔒 Auditoría de concurrencia (PLAN_REMODELACION_ARQUITECTURA.md, 1.3) —
    # el precheck `ya_matriculado` de arriba solo cierra la ventana
    # SECUENCIAL. Dos llamadas casi simultáneas a matricular-roster con
    # persona_ids superpuestos (doble clic, o un reintento del frontend)
    # pueden pasar ambas ese precheck antes de que cualquiera haga commit, y
    # solo el índice único `uq_edu_inscripciones_evaluacion_persona` detiene
    # a la segunda — como IntegrityError crudo (500) sin esto. Mismo patrón
    # que `registrar_consentimiento`.
    commit_o_conflicto(
        db,
        "Una de estas personas ya fue matriculada en esta campaña por otra solicitud "
        "simultánea. Vuelve a intentar la matriculación para los alumnos restantes.",
    )
    for i in inscripciones_creadas:
        db.refresh(i)

    return {
        "generados": inscripciones_creadas, 
        "rechazados": rechazados
    }

def verificar_propiedad_de_prueba(db: Session, prueba_id: int, usuario_actual):
    """
    None = pública (solo superadmin la toca).
    Con institucion_id = privada de un colegio (solo ese colegio la toca).
    """
    prueba = db.query(models.Prueba).filter(models.Prueba.id == prueba_id).first()
    if not prueba:
        raise HTTPException(status_code=404, detail="Prueba no encontrada.")
    if usuario_actual.rol == "superadmin":
        return prueba
    if prueba.institucion_id is None or prueba.institucion_id != usuario_actual.institucion_id:
        raise HTTPException(status_code=403, detail="No puedes modificar una prueba que no es tuya.")
    return prueba

def generar_excel_dashboard(dashboard: dict) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resultados"
    ws.append(["Alumno", "Estado", "Dimensión", "Puntaje Directo", "Nivel"])

    for alumno in dashboard["alumnos"]:
        if not alumno["resultados"]:
            ws.append([alumno["nombre_completo"], alumno["estado"], "-", "-", "-"])
        for r in alumno["resultados"]:
            ws.append([alumno["nombre_completo"], alumno["estado"], r["dimension"], r["puntaje_directo"], r["nivel"]])

    # 🆕 Hoja "Resumen estadístico" — mismo cálculo que ResumenEstadisticoPrueba
    # en aurora_ui (calcular_resumen_estadistico_dashboard, una sola vez,
    # compartido por el JSON del dashboard y este Excel). Layout de columnas
    # FIJO para cualquier prueba (igual que la hoja "Resultados" de arriba,
    # ya uniforme para cualquier instrumento) — una celda queda vacía cuando
    # esa columna no aplica al modo_interpretacion de esta prueba (ej. "Nivel
    # mayoritario" en modo 'perfil', "Alumnos... dominante" en modo 'rango').
    resumen = dashboard.get("resumen_estadistico") or {}
    ws2 = wb.create_sheet("Resumen estadístico")
    ws2.append([
        "Dimensión", "Tipo", "N", "Media", "Mediana", "Mínimo", "Máximo", "Máximo de la escala",
        "Nivel mayoritario", "% mayoría", "En alerta", "% alerta", "Alumnos con esta como dominante",
    ])
    for d in resumen.get("dimensiones", []):
        distribucion = d.get("distribucion") or []
        mayoria = distribucion[0] if distribucion else None
        pct_alerta = (
            round(d["en_alerta"] / d["n"] * 100, 1)
            if d.get("en_alerta") is not None and d.get("n")
            else None
        )
        ws2.append([
            d["nombre"], d["tipo_calculo"], d["n"],
            round(d["media"], 2) if d.get("media") is not None else None,
            round(d["mediana"], 2) if d.get("mediana") is not None else None,
            d.get("minimo"), d.get("maximo"), d.get("maximo_escala"),
            mayoria["nivel"] if mayoria else None,
            round(mayoria["pct"], 1) if mayoria else None,
            d.get("en_alerta"), pct_alerta,
            d.get("dominante_cantidad"),
        ])

    # 🆕 Hoja "Distribución por nivel" — el desglose COMPLETO (no solo la
    # mayoría de la hoja anterior), una fila por (dimensión, nivel). No se
    # crea en modo_interpretacion='perfil' (nunca hay niveles categóricos,
    # ver docs/motor-scoring.md) ni para una prueba sin ningún baremo
    # aplicado todavía — mismo criterio "no mostrar una hoja vacía".
    filas_distribucion = [
        (d["nombre"], niv)
        for d in resumen.get("dimensiones", [])
        for niv in (d.get("distribucion") or [])
    ]
    if filas_distribucion:
        ws3 = wb.create_sheet("Distribución por nivel")
        ws3.append(["Dimensión", "Nivel", "Cantidad", "%", "Es alerta"])
        for nombre_dim, niv in filas_distribucion:
            ws3.append([nombre_dim, niv["nivel"], niv["cantidad"], round(niv["pct"], 1), "Sí" if niv["es_alerta"] else "No"])

    # 🆕 Hoja "Áreas vocacionales" — solo cuando la prueba agrupa por área
    # (CHASIDE, ver Prueba.agrupar_dimensiones_por_area) y el grupo realmente
    # tiene alguna área destacada en al menos un alumno.
    areas = resumen.get("areas_compatibles") or []
    if areas:
        total_alumnos = dashboard.get("total_finalizados") or len(dashboard["alumnos"])
        ws4 = wb.create_sheet("Áreas vocacionales")
        ws4.append(["Área", "Alumnos compatibles", "% del total"])
        for a in areas:
            pct = round(a["cantidad"] / total_alumnos * 100, 1) if total_alumnos else None
            ws4.append([a["area"], a["cantidad"], pct])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

def obtener_dashboard_proceso(db: Session, institucion_id: int):
    """
    Consolida las evaluaciones de una institución o campaña global, 
    agrupando por estudiante (persona_id o inscripcion_id).
    """
    evaluaciones = db.query(models.Evaluacion).filter(models.Evaluacion.institucion_id == institucion_id).all()
    if not evaluaciones:
        return None

    consolidado = {}
    for evaluacion in evaluaciones:
        dash = obtener_dashboard_campana(db, evaluacion.id)
        if not dash or "alumnos" not in dash:
            continue

        for alumno in dash["alumnos"]:
            # 🐢 Rendimiento (Fase 4) — antes se volvía a consultar la Inscripción
            # aquí solo para leer persona_id, aunque obtener_dashboard_campana ya
            # la había traído: con N campañas x M alumnos eran N×M consultas
            # adicionales. persona_id ahora viaja en el dict que ya devolvió.
            clave = alumno.get("persona_id") or alumno["inscripcion_id"]

            if clave not in consolidado:
                consolidado[clave] = {
                    "nombre_completo": alumno["nombre_completo"], 
                    "pruebas": []
                }

            consolidado[clave]["pruebas"].append({
                "prueba": dash.get("nombre_prueba", "Evaluación"),
                "estado": alumno["estado"],
                "resultados": alumno.get("resultados", []),
            })

    return {
        "institucion_id": institucion_id,
        "alumnos": list(consolidado.values()),
    }


def redactar_resultados_dashboard_campana(dashboard: dict) -> dict:
    """🔒 Remediación de permisos, paso 5 — copia el dashboard de
    `obtener_dashboard_campana` vaciando `resultados` (puntaje/nivel real)
    de cada alumno, para `colegio`. Deja intactos nombre/estado/fecha (quién
    rindió, quién falta) y los conteos agregados (`total_*`,
    `distribucion_por_dimension`) — mismo criterio que `ficha-completa`: ver
    participación/avance está bien, ver el diagnóstico puntual no."""
    return {
        **dashboard,
        "alumnos": [{**alumno, "resultados": []} for alumno in dashboard["alumnos"]],
    }


def redactar_resultados_dashboard_institucional(dashboard: dict) -> dict:
    """Misma redacción que `redactar_resultados_dashboard_campana`, pero para
    la forma consolidada por institución de `obtener_dashboard_proceso`
    (`alumnos[].pruebas[].resultados`, no `alumnos[].resultados`)."""
    return {
        **dashboard,
        "alumnos": [
            {**alumno, "pruebas": [{**p, "resultados": []} for p in alumno["pruebas"]]}
            for alumno in dashboard["alumnos"]
        ],
    }


def obtener_resumen_inicio(db: Session, institucion_id: int) -> dict:
    """🆕 Resumen operativo del mes para `/inicio` (sesión 2026-10-25) — ver
    docstring del router. `inicio_mes` es naive-UTC, mismo criterio que
    `campana_esta_expirada`/`Inscripcion.fecha_resolucion` (ambos sin zona
    horaria), para no mezclar aware/naive en la comparación."""
    ahora = datetime.utcnow()
    inicio_mes = ahora.replace(day=1, hour=0, minute=0, second=0, microsecond=0)

    total_alumnos = db.query(superadmin_models.Persona).filter(
        superadmin_models.Persona.institucion_id == institucion_id,
        superadmin_models.Persona.tipo_perfil == "estudiante",
        superadmin_models.Persona.activo == True,  # noqa: E712
    ).count()

    pruebas_finalizadas_mes = db.query(models.Inscripcion).join(
        models.Evaluacion, models.Inscripcion.evaluacion_id == models.Evaluacion.id
    ).filter(
        models.Evaluacion.institucion_id == institucion_id,
        models.Inscripcion.estado == "Finalizado",
        models.Inscripcion.is_deleted == False,  # noqa: E712
        models.Inscripcion.fecha_resolucion >= inicio_mes,
    ).count()

    # DISTINCT persona_id — un alumno con dos sesiones el mismo mes cuenta
    # una sola vez ("cuántos alumnos" recibieron seguimiento, no cuántas
    # sesiones hubo).
    estudiantes_con_seguimiento_mes = db.query(
        func.count(func.distinct(models.SesionSeguimiento.persona_id))
    ).filter(
        models.SesionSeguimiento.institucion_id == institucion_id,
        models.SesionSeguimiento.fecha >= inicio_mes,
    ).scalar() or 0

    campanas_activas = db.query(models.Evaluacion).filter(
        models.Evaluacion.institucion_id == institucion_id,
        models.Evaluacion.estado == models.ESTADO_CAMPANA_ACTIVA,
    ).count()

    return {
        "total_alumnos": total_alumnos,
        "pruebas_finalizadas_mes": pruebas_finalizadas_mes,
        "estudiantes_con_seguimiento_mes": estudiantes_con_seguimiento_mes,
        "campanas_activas": campanas_activas,
        "mes": inicio_mes.strftime("%Y-%m"),
    }


def _meses_hacia_atras(cantidad: int):
    """Primer día (naive-UTC) de los últimos `cantidad` meses, del más
    antiguo al actual inclusive."""
    hoy = datetime.utcnow()
    anio, mes = hoy.year, hoy.month
    salida = []
    for _ in range(cantidad):
        salida.append(datetime(anio, mes, 1))
        mes -= 1
        if mes == 0:
            mes, anio = 12, anio - 1
    return list(reversed(salida))


def obtener_series_inicio(db: Session, institucion_id: int, usuario_id: int, meses: int = 6) -> dict:
    """🆕 Series mensuales + avance de campañas para los gráficos de `/inicio`
    (sesión 2026-10-30). Una consulta agrupada por mes por serie (nunca una
    por mes), todo acotado a la institución:

    - evaluaciones: inscripciones finalizadas por mes de fecha_resolucion, y
      cuántas de ellas tienen al menos un diagnóstico con es_alerta (el
      indicador autoritativo — nunca el texto del nivel).
    - sesiones: sesiones de seguimiento registradas por mes.
    - citas: mismo alcance que la agenda del usuario (sus citas individuales
      + todas las actividades grupales del colegio; las individuales de otro
      psicólogo son privadas y no se cuentan), sin las canceladas.
    - campanas: hasta 6 campañas activas con inscritos/finalizados, las que
      cierran antes primero.
    """
    from modulos.superadmin import models as sa_models

    inicios = _meses_hacia_atras(meses)
    desde = inicios[0]
    clave = lambda d: d.strftime("%Y-%m")  # noqa: E731
    mes_de = lambda col: func.date_trunc("month", col)  # noqa: E731

    finalizadas = dict(
        (clave(m), n) for m, n in db.query(
            mes_de(models.Inscripcion.fecha_resolucion), func.count(models.Inscripcion.id),
        ).join(
            models.Evaluacion, models.Inscripcion.evaluacion_id == models.Evaluacion.id,
        ).filter(
            models.Evaluacion.institucion_id == institucion_id,
            models.Inscripcion.estado == "Finalizado",
            models.Inscripcion.is_deleted == False,  # noqa: E712
            models.Inscripcion.fecha_resolucion >= desde,
        ).group_by(mes_de(models.Inscripcion.fecha_resolucion)).all()
    )

    hay_alerta = db.query(models.DiagnosticoAlumno.id).filter(
        models.DiagnosticoAlumno.inscripcion_id == models.Inscripcion.id,
        models.DiagnosticoAlumno.es_alerta == True,  # noqa: E712
    ).exists()
    con_alerta = dict(
        (clave(m), n) for m, n in db.query(
            mes_de(models.Inscripcion.fecha_resolucion), func.count(models.Inscripcion.id),
        ).join(
            models.Evaluacion, models.Inscripcion.evaluacion_id == models.Evaluacion.id,
        ).filter(
            models.Evaluacion.institucion_id == institucion_id,
            models.Inscripcion.estado == "Finalizado",
            models.Inscripcion.is_deleted == False,  # noqa: E712
            models.Inscripcion.fecha_resolucion >= desde,
            hay_alerta,
        ).group_by(mes_de(models.Inscripcion.fecha_resolucion)).all()
    )

    sesiones = dict(
        (clave(m), n) for m, n in db.query(
            mes_de(models.SesionSeguimiento.fecha), func.count(models.SesionSeguimiento.id),
        ).filter(
            models.SesionSeguimiento.institucion_id == institucion_id,
            models.SesionSeguimiento.fecha >= desde,
        ).group_by(mes_de(models.SesionSeguimiento.fecha)).all()
    )

    Cita = sa_models.CitaProgramada
    citas = dict(
        (clave(m), n) for m, n in db.query(
            mes_de(Cita.fecha_hora), func.count(Cita.id),
        ).filter(
            Cita.institucion_id == institucion_id,
            or_(Cita.psicologo_id == usuario_id, Cita.persona_id.is_(None)),
            Cita.estado != "Cancelada",
            Cita.fecha_hora >= desde,
        ).group_by(mes_de(Cita.fecha_hora)).all()
    )

    serie = []
    for inicio in inicios:
        k = clave(inicio)
        total = finalizadas.get(k, 0)
        alerta = con_alerta.get(k, 0)
        serie.append({
            "mes": k,
            "evaluaciones": total,
            "con_alerta": alerta,
            "sin_alerta": total - alerta,
            "sesiones_seguimiento": sesiones.get(k, 0),
            "citas": citas.get(k, 0),
        })

    activas = db.query(models.Evaluacion).filter(
        models.Evaluacion.institucion_id == institucion_id,
        models.Evaluacion.estado == models.ESTADO_CAMPANA_ACTIVA,
    ).order_by(models.Evaluacion.fecha_fin.asc().nullslast(), models.Evaluacion.id.desc()).limit(6).all()
    conteos = {}
    if activas:
        for evaluacion_id, estado, n in db.query(
            models.Inscripcion.evaluacion_id, models.Inscripcion.estado, func.count(models.Inscripcion.id),
        ).filter(
            models.Inscripcion.evaluacion_id.in_([e.id for e in activas]),
            models.Inscripcion.is_deleted == False,  # noqa: E712
        ).group_by(models.Inscripcion.evaluacion_id, models.Inscripcion.estado).all():
            c = conteos.setdefault(evaluacion_id, {"inscritos": 0, "finalizados": 0})
            c["inscritos"] += n
            if estado == "Finalizado":
                c["finalizados"] += n
    pruebas = {
        p.id: (p.nombre_corto or p.nombre_tecnico)
        for p in db.query(models.Prueba).filter(models.Prueba.id.in_({e.prueba_id for e in activas})).all()
    } if activas else {}
    campanas = [{
        "id": e.id,
        "nombre_grupo": e.nombre_grupo,
        "prueba": pruebas.get(e.prueba_id),
        "fecha_fin": e.fecha_fin.isoformat() if e.fecha_fin else None,
        "inscritos": conteos.get(e.id, {}).get("inscritos", 0),
        "finalizados": conteos.get(e.id, {}).get("finalizados", 0),
    } for e in activas]

    return {"meses": serie, "campanas": campanas}


def obtener_casos_pendientes_de_seguimiento(db: Session, institucion_id: int):
    # 🆕 Rediseño de /casos-riesgo (sesión 2026-10-30): la pantalla leía
    # `prueba`/`nivel`/`fecha`, claves que este listado nunca devolvió — se
    # veía el nombre del alumno con todo lo demás vacío. Se suman (solo
    # ADITIVO, ningún consumidor existente cambia) prueba/dimensión/campaña
    # e ids, resueltos en el mismo SELECT — Evaluacion ya estaba en el join,
    # Prueba/Dimension son 1:1 por fila, sin N+1 nuevo.
    alertas = db.query(
        models.DiagnosticoAlumno, models.Inscripcion, models.Evaluacion, models.Prueba, models.Dimension
    ).join(
        models.Inscripcion,
        models.DiagnosticoAlumno.inscripcion_id == models.Inscripcion.id
    ).join(
        models.Evaluacion,
        models.Inscripcion.evaluacion_id == models.Evaluacion.id
    ).outerjoin(
        models.Prueba,
        models.Evaluacion.prueba_id == models.Prueba.id
    ).outerjoin(
        models.Dimension,
        models.DiagnosticoAlumno.dimension_id == models.Dimension.id
    ).filter(
    models.Evaluacion.institucion_id == institucion_id,
    models.Inscripcion.persona_id.isnot(None),
    models.Inscripcion.is_deleted == False,   # 🔒 FIX
    ).all()

    # 🚨 Fase 4 — HALLAZGO CRÍTICO: este listado seguía filtrando por el texto
    # legacy de NIVELES_DE_ALERTA, el mismo criterio cerrado en toda la Fase 1
    # (C-13/B-6) a favor del indicador explícito `es_alerta`. Un baremo
    # marcado es_alerta=True con un nivel NO incluido en ese set de palabras
    # — exactamente el caso de prueba "Sin indicadores aparentes" de la Fase 1
    # — sí retenía el informe y sí disparaba el correo, pero jamás aparecía
    # aquí: el caso quedaba huérfano en el worklist de seguimiento del
    # psicólogo, invisible para quien tiene que darle continuidad clínica.
    pendientes = []
    for diag, insc, evaluacion, prueba, dimension in alertas:
        if not diag.es_alerta:
            continue
        tiene_seguimiento = db.query(models.SesionSeguimiento).filter(
            models.SesionSeguimiento.persona_id == insc.persona_id,
            models.SesionSeguimiento.fecha >= insc.fecha_resolucion,
        ).first()
        if not tiene_seguimiento:
            pendientes.append({
                "persona_id": insc.persona_id,
                "nombre_completo": insc.nombre_completo,
                "nivel_alerta": diag.nivel,
                "fecha_evaluacion": insc.fecha_resolucion,
                "inscripcion_id": insc.id,
                "evaluacion_id": evaluacion.id,
                "campana": evaluacion.nombre_grupo,
                "prueba": (prueba.nombre_corto or prueba.nombre_tecnico) if prueba else None,
                "dimension": dimension.nombre if dimension else None,
            })
    return pendientes

def puede_aplicar_prueba(db: Session, persona_id: int, prueba_id: int, institucion_id: int):
    """
    Dos candados de buena práctica clínica, no de rendimiento del servidor:
    1) Límite mensual de pruebas totales para ese alumno (fatiga de evaluación).
    2) Tiempo mínimo antes de reaplicar la MISMA prueba (efecto de práctica).
    Devuelve (True, "") si puede aplicarse, o (False, "motivo") si no.
    """
    institucion = db.query(inst_models.Institucion).filter(inst_models.Institucion.id == institucion_id).first()
    limite_mensual = institucion.limite_pruebas_mensual_por_alumno if institucion else 5

    # 🆕 Cierre de gap — naive a propósito: se usa como valor de un filtro
    # SQLAlchemy contra Inscripcion.fecha_resolucion, que es DateTime sin
    # zona horaria.
    hace_30_dias = datetime.utcnow() - timedelta(days=30)
    aplicadas_ultimo_mes = db.query(models.Inscripcion).join(
        models.Evaluacion, models.Inscripcion.evaluacion_id == models.Evaluacion.id
    ).filter(
        models.Inscripcion.persona_id == persona_id,
        models.Inscripcion.fecha_resolucion >= hace_30_dias,
        # 🔒 C-8: una evaluación dada de baja no debe seguir consumiendo el
        # cupo mensual del alumno ni bloquear una reaplicación legítima.
        models.Inscripcion.is_deleted == False,  # noqa: E712
    ).count()

    if aplicadas_ultimo_mes >= limite_mensual:
        return False, f"Este alumno ya alcanzó el límite de {limite_mensual} pruebas en los últimos 30 días."

    prueba = db.query(models.Prueba).filter(models.Prueba.id == prueba_id).first()
    dias_minimos = prueba.dias_minimos_entre_aplicaciones or 60
    # 🆕 Cierre de gap — mismo motivo que hace_30_dias más arriba: filtro
    # SQLAlchemy contra una columna DateTime sin zona horaria.
    limite_fecha = datetime.utcnow() - timedelta(days=dias_minimos)

    ya_la_rindio_reciente = db.query(models.Inscripcion).join(
        models.Evaluacion, models.Inscripcion.evaluacion_id == models.Evaluacion.id
    ).filter(
        models.Inscripcion.persona_id == persona_id,
        models.Evaluacion.prueba_id == prueba_id,
        models.Inscripcion.fecha_resolucion >= limite_fecha,
    ).first()

    if ya_la_rindio_reciente:
        return False, f"Esta prueba ya se le aplicó a este alumno hace menos de {dias_minimos} días."

    return True, ""

def soft_delete_alumno(
    db: Session, inscripcion_id: int, institucion_id: int, es_superadmin: bool = False,
    ejecutado_por=None, ip_solicitante: str = None, referencia_solicitud: str = None,
) -> bool:
    """
    🔒 Realiza el borrado lógico (Soft Delete) de la inscripción de un alumno.
    🏛️ B-12 — Deja constancia en la bitácora ARCO de quién y desde dónde.
    """
    query = db.query(models.Inscripcion).join(
        models.Evaluacion, models.Inscripcion.evaluacion_id == models.Evaluacion.id
    ).filter(models.Inscripcion.id == inscripcion_id)

    if not es_superadmin:
        query = query.filter(models.Evaluacion.institucion_id == institucion_id)

    inscripcion = query.first()
    if not inscripcion or inscripcion.is_deleted:
        return False

    evaluacion = db.query(models.Evaluacion).filter(
        models.Evaluacion.id == inscripcion.evaluacion_id
    ).first()

    if ejecutado_por is not None:
        superadmin_crud.registrar_auditoria_arco(
            db, accion="soft_delete_inscripcion", ejecutado_por=ejecutado_por,
            institucion_id=evaluacion.institucion_id if evaluacion else None,
            persona_id=inscripcion.persona_id, inscripcion_id=inscripcion.id,
            ip_solicitante=ip_solicitante,
            detalle=(
                f"Borrado lógico de la inscripción de '{inscripcion.nombre_completo}'"
                + (f". Solicitud: {referencia_solicitud}" if referencia_solicitud else "")
            ),
        )

    inscripcion.is_deleted = True
    inscripcion.deleted_at = datetime.now(timezone.utc)
    db.commit()
    return True


def anonimizar_datos_alumno_arco(
    db: Session, inscripcion_id: int, institucion_id: int, es_superadmin: bool = False,
    ejecutado_por=None, ip_solicitante: str = None, referencia_solicitud: str = None,
) -> bool:
    """
    🏛️ LEY 29733 (Derecho ARCO): Anonimiza irreversiblemente la PII del alumno,
    preservando el registro psicométrico anonimizado.

    Sigue sin ser un borrado físico: la fila y sus respuestas permanecen para
    conservar la validez psicométrica del grupo; solo se vacía la PII.

    🏛️ B-12 — Se anota en la bitácora ARCO antes de perder los datos.
    """
    query = db.query(models.Inscripcion).join(
        models.Evaluacion, models.Inscripcion.evaluacion_id == models.Evaluacion.id
    ).filter(models.Inscripcion.id == inscripcion_id)

    if not es_superadmin:
        query = query.filter(models.Evaluacion.institucion_id == institucion_id)

    inscripcion = query.first()
    if not inscripcion:
        return False

    evaluacion = db.query(models.Evaluacion).filter(
        models.Evaluacion.id == inscripcion.evaluacion_id
    ).first()

    if ejecutado_por is not None:
        superadmin_crud.registrar_auditoria_arco(
            db, accion="anonimizacion_inscripcion", ejecutado_por=ejecutado_por,
            institucion_id=evaluacion.institucion_id if evaluacion else None,
            persona_id=inscripcion.persona_id, inscripcion_id=inscripcion.id,
            ip_solicitante=ip_solicitante,
            detalle=(
                f"Anonimización irreversible de la inscripción de "
                f"'{inscripcion.nombre_completo}'"
                + (f". Solicitud: {referencia_solicitud}" if referencia_solicitud else "")
            ),
        )

    hash_arco = hashlib.sha256(f"ARCO_{inscripcion.id}_{datetime.now(timezone.utc)}".encode()).hexdigest()[:10]

    inscripcion.nombre_completo = f"ANONIMO_ARCO_{hash_arco}"
    inscripcion.persona_id = None          # 🔒 FIX: rompe el enlace hacia el registro maestro identificable
    inscripcion.is_anonymized = True
    inscripcion.anonymized_at = datetime.now(timezone.utc)
    inscripcion.is_deleted = True
    inscripcion.deleted_at = datetime.now(timezone.utc)

    db.commit()
    return True

def obtener_historial_de_persona(db: Session, persona_id: int):
    """
    Retorna todo el historial de seguimiento de sesiones de una persona.
    """
    return db.query(models.SesionSeguimiento).filter(
        models.SesionSeguimiento.persona_id == persona_id
    ).order_by(models.SesionSeguimiento.fecha.desc()).all()

def crear_sesion_seguimiento(
    db: Session, 
    datos: schemas.SesionSeguimientoCreate, 
    psicologo_id: int, 
    institucion_id: int
):
    nueva = models.SesionSeguimiento(
        persona_id=datos.persona_id,
        psicologo_id=psicologo_id,
        institucion_id=institucion_id,
        observaciones=datos.observaciones,
        accion_tomada=datos.accion_tomada,
    )
    db.add(nueva)
    db.commit()
    db.refresh(nueva)
    return nueva

def actualizar_sesion_seguimiento(db: Session, sesion_id: int, datos: schemas.SesionSeguimientoUpdate):
    sesion = db.query(models.SesionSeguimiento).filter(
        models.SesionSeguimiento.id == sesion_id
    ).first()
    if not sesion:
        return None
    for campo, valor in datos.model_dump(exclude_unset=True).items():
        setattr(sesion, campo, valor)
    db.commit()
    db.refresh(sesion)
    return sesion


def eliminar_sesion_seguimiento(db: Session, sesion_id: int) -> bool:
    sesion = db.query(models.SesionSeguimiento).filter(
        models.SesionSeguimiento.id == sesion_id
    ).first()
    if not sesion:
        return False
    db.delete(sesion)
    db.commit()
    return True

def actualizar_prueba(db: Session, prueba_id: int, datos: schemas.PruebaUpdate):
    prueba = db.query(models.Prueba).filter(models.Prueba.id == prueba_id).first()
    if not prueba:
        return None
    for campo, valor in datos.model_dump(exclude_unset=True).items():
        setattr(prueba, campo, valor)
    db.commit()
    db.refresh(prueba)
    return prueba

def obtener_baremos_por_dimension(db: Session, dimension_id: int):
    return db.query(models.Baremo).filter(models.Baremo.dimension_id == dimension_id).all()

def obtener_campos_socio_por_token(db: Session, token: str):
    inscripcion = db.query(models.Inscripcion).filter(models.Inscripcion.token == token).first()
    if not inscripcion:
        return None
    
    evaluacion = db.query(models.Evaluacion).filter(models.Evaluacion.id == inscripcion.evaluacion_id).first()
    
    return db.query(models.CampoSociodemografico).filter(
        models.CampoSociodemografico.cliente_id == evaluacion.institucion_id
    ).all()


def guardar_respuestas_socio(db: Session, token: str, respuestas: List[schemas.RespuestaSociodemograficaSubmit]):
    inscripcion = db.query(models.Inscripcion).filter(models.Inscripcion.token == token).first()
    if not inscripcion:
        return None
        
    campos_ids = [r.campo_id for r in respuestas]
    
    # Eliminamos respuestas anteriores para evitar duplicados si el usuario reenvía el formulario
    db.query(models.RespuestaSociodemografica).filter(
        models.RespuestaSociodemografica.inscripcion_id == inscripcion.id,
        models.RespuestaSociodemografica.campo_id.in_(campos_ids),
    ).delete(synchronize_session=False)
    
    # Creamos los nuevos registros
    nuevas = [
        models.RespuestaSociodemografica(inscripcion_id=inscripcion.id, campo_id=r.campo_id, valor=r.valor)
        for r in respuestas
    ]
    
    db.add_all(nuevas)
    db.commit()
    
    return {"mensaje": f"Se guardaron {len(nuevas)} respuestas sociodemográficas."}

def listar_campanas_de_institucion(db: Session, institucion_id: int, skip: int = 0, limit: int = 50):
    """🆕 Paginación real (docs/gaps-pendientes.md) — antes devolvía TODA
    campaña creada, sin tope, para siempre. `total` se cuenta antes de
    aplicar `offset/limit`."""
    query = db.query(models.Evaluacion).filter(
        models.Evaluacion.institucion_id == institucion_id
    )
    total = query.count()
    items = query.order_by(models.Evaluacion.fecha_creacion.desc()).offset(skip).limit(limit).all()
    return total, items

def obtener_prueba_completa(db: Session, prueba_id: int):
    prueba = db.query(models.Prueba).filter(models.Prueba.id == prueba_id).first()
    if not prueba:
        return None

    dimensiones = db.query(models.Dimension).filter(models.Dimension.prueba_id == prueba_id).all()
    preguntas = db.query(models.Pregunta).filter(
        models.Pregunta.prueba_id == prueba_id
    ).order_by(models.Pregunta.orden).all()
    baremos = db.query(models.Baremo).join(
        models.Dimension, models.Baremo.dimension_id == models.Dimension.id
    ).filter(models.Dimension.prueba_id == prueba_id).all()

    #   🆕 Traemos TODAS las opciones de esta prueba en una sola query
    pregunta_ids = [p.id for p in preguntas]
    opciones = db.query(models.OpcionRespuesta).filter(
        models.OpcionRespuesta.pregunta_id.in_(pregunta_ids)
    ).all() if pregunta_ids else []

    opciones_por_pregunta = {}
    for o in opciones:
        opciones_por_pregunta.setdefault(o.pregunta_id, []).append(o)

    #   🆕 Anidamos las opciones dentro de cada pregunta
    preguntas_por_dimension = {}
    for p in preguntas:
        p_dict = {
            "id": p.id,
            "prueba_id": p.prueba_id,
            "texto_pregunta": p.texto_pregunta,
            "orden": p.orden,
            "dimension_id": p.dimension_id,
            "es_inverso": p.es_inverso,
            "opciones": opciones_por_pregunta.get(p.id, []),
        }
        preguntas_por_dimension.setdefault(p.dimension_id, []).append(p_dict)

    dimensiones_con_preguntas = []
    for d in dimensiones:
        dimensiones_con_preguntas.append({
            "id": d.id,
            "prueba_id": d.prueba_id,
            "nombre": d.nombre,
            "descripcion": d.descripcion,
            "preguntas": preguntas_por_dimension.get(d.id, []),
        })

    reglas = db.query(models.ReglaInterpretacion).filter(
        models.ReglaInterpretacion.prueba_id == prueba_id
    ).all()

    return {**prueba.__dict__, "dimensiones": dimensiones_con_preguntas, "baremos": baremos, "reglas": reglas}

def actualizar_pregunta(db: Session, pregunta_id: int, datos: schemas.PreguntaUpdate):
    pregunta = db.query(models.Pregunta).filter(models.Pregunta.id == pregunta_id).first()
    if not pregunta:
        return None
    cambios = datos.model_dump(exclude_unset=True)
    if "dimension_id" in cambios:
        _validar_dimension_no_compuesta_para_pregunta(db, cambios["dimension_id"])
    for campo, valor in cambios.items():
        setattr(pregunta, campo, valor)
    db.commit()
    db.refresh(pregunta)
    return pregunta

def eliminar_pregunta(db: Session, pregunta_id: int) -> bool:
    pregunta = db.query(models.Pregunta).filter(models.Pregunta.id == pregunta_id).first()
    if not pregunta:
        return False
    db.delete(pregunta)
    db.commit()
    return True

def actualizar_dimension(db: Session, dimension_id: int, datos: schemas.DimensionUpdate):
    dimension = db.query(models.Dimension).filter(models.Dimension.id == dimension_id).first()
    if not dimension:
        return None
    cambios = datos.model_dump(exclude_unset=True)

    # 🆕 tipo_calculo decide CÓMO se llega al puntaje (ver Dimension.tipo_
    # calculo) — cambiarlo sobre una dimensión que ya tiene resultados
    # calculados reinterpretaría retroactivamente datos ya producidos con el
    # criterio anterior, mismo espíritu que bloquear rango/nivel/es_alerta en
    # un Baremo ya usado. conteo_umbral_item recibe el mismo candado: cambiar
    # el umbral de "qué cuenta como síntoma" es tan retroactivo como cambiar
    # el tipo_calculo mismo.
    cambia_tipo = "tipo_calculo" in cambios and cambios["tipo_calculo"] != dimension.tipo_calculo
    cambia_umbral = (
        "conteo_umbral_item" in cambios and cambios["conteo_umbral_item"] != dimension.conteo_umbral_item
    )
    if cambia_tipo or cambia_umbral:
        tiene_resultados = db.query(models.ResultadoDimension.id).filter(
            models.ResultadoDimension.dimension_id == dimension_id
        ).first() is not None
        if tiene_resultados:
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail="Esta dimensión ya tiene resultados calculados: no se puede cambiar tipo_calculo/conteo_umbral_item.",
            )
        if cambia_tipo and cambios["tipo_calculo"] == models.TIPO_CALCULO_SUMA_DE_DIMENSIONES:
            tiene_preguntas = db.query(models.Pregunta.id).filter(
                models.Pregunta.dimension_id == dimension_id
            ).first() is not None
            if tiene_preguntas:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Esta dimensión tiene preguntas propias: no puede convertirse en 'suma_de_dimensiones'.",
                )

    # 🆕 El estado FINAL (ya aplicados los cambios de este PATCH) nunca puede
    # dejar una dimensión 'conteo_sintomas' sin umbral — mismo requisito que
    # el CHECK de BD (ck_dimensiones_conteo_umbral_requerido), pero con un 400
    # legible en vez de un IntegrityError crudo al hacer commit.
    tipo_resultante = cambios.get("tipo_calculo", dimension.tipo_calculo)
    umbral_resultante = cambios.get("conteo_umbral_item", dimension.conteo_umbral_item)
    if tipo_resultante == models.TIPO_CALCULO_CONTEO_SINTOMAS and umbral_resultante is None:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Una dimensión 'conteo_sintomas' requiere conteo_umbral_item.",
        )

    for campo, valor in cambios.items():
        setattr(dimension, campo, valor)
    db.commit()
    db.refresh(dimension)
    return dimension

def eliminar_dimension(db: Session, dimension_id: int) -> bool:
    dimension = db.query(models.Dimension).filter(models.Dimension.id == dimension_id).first()
    if not dimension:
        return False

    #   🔒 Candado de integridad referencial — antes de tocar la BD
    tiene_preguntas = db.query(models.Pregunta).filter(
        models.Pregunta.dimension_id == dimension_id
    ).first()
    if tiene_preguntas:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="No se puede eliminar la dimensión: tiene preguntas asociadas. Elimina o reasigna las preguntas primero."
        )

    tiene_baremos = db.query(models.Baremo).filter(
        models.Baremo.dimension_id == dimension_id
    ).first()
    if tiene_baremos:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="No se puede eliminar la dimensión: tiene baremos de interpretación asociados."
        )

    tiene_resultados = db.query(models.ResultadoDimension).filter(
        models.ResultadoDimension.dimension_id == dimension_id
    ).first()
    if tiene_resultados:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="No se puede eliminar la dimensión: ya existen resultados de alumnos calculados sobre ella. Desactiva la prueba en su lugar."
        )

    # 🆕 Si esta dimensión alimenta a una dimensión compuesta (ver
    # DimensionComponente), borrarla desarmaría en silencio esa composición
    # ya configurada — a diferencia de ser ELLA la compuesta (cuyas propias
    # filas de composición sí tienen ondelete="CASCADE", porque no tienen
    # sentido sin su dueña).
    es_componente_de_otra = db.query(models.DimensionComponente).filter(
        models.DimensionComponente.dimension_componente_id == dimension_id
    ).first()
    if es_componente_de_otra:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                "No se puede eliminar la dimensión: es componente de una dimensión "
                "'suma_de_dimensiones'. Quítala primero de esa composición."
            ),
        )

    db.delete(dimension)
    db.commit()
    return True


# ==========================================
# COMPOSICIÓN DE DIMENSIONES COMPUESTAS (Dimension.tipo_calculo == 'suma_de_dimensiones')
# ==========================================
def _validar_componentes_dimension_compuesta(db: Session, dimension_compuesta: models.Dimension, ids_componentes: list):
    """Nunca confiar en los ids que llegaron del cliente: cada componente se
    revalida contra la BD (existe, pertenece a la MISMA prueba, y no es a su
    vez una dimensión compuesta — sin anidar, ver DimensionComponente)."""
    if dimension_compuesta.tipo_calculo != models.TIPO_CALCULO_SUMA_DE_DIMENSIONES:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Solo una dimensión 'suma_de_dimensiones' puede tener componentes.",
        )
    if dimension_compuesta.id in ids_componentes:
        raise HTTPException(status_code=400, detail="Una dimensión no puede ser su propio componente.")
    if len(set(ids_componentes)) != len(ids_componentes):
        raise HTTPException(status_code=400, detail="La lista de componentes tiene ids duplicados.")
    if not ids_componentes:
        return

    encontrados = {
        d.id: d for d in db.query(models.Dimension).filter(models.Dimension.id.in_(ids_componentes)).all()
    }
    tipos_calculo_vistos = set()
    for cid in ids_componentes:
        dim = encontrados.get(cid)
        if dim is None:
            raise HTTPException(status_code=404, detail=f"La dimensión componente #{cid} no existe.")
        if dim.prueba_id != dimension_compuesta.prueba_id:
            raise HTTPException(
                status_code=400,
                detail=f"La dimensión componente #{cid} no pertenece a la misma prueba que la dimensión compuesta.",
            )
        if dim.tipo_calculo == models.TIPO_CALCULO_SUMA_DE_DIMENSIONES:
            raise HTTPException(
                status_code=400,
                detail=f"La dimensión componente #{cid} es a su vez 'suma_de_dimensiones': no se permite anidar.",
            )
        tipos_calculo_vistos.add(dim.tipo_calculo)
    # 🆕 Auditoría de contenido (sesión 2026-10-12): mezclar un componente
    # 'items' (suma de puntajes) con uno 'conteo_sintomas' (cantidad de
    # ítems que alcanzan un umbral) produce un "total" sin unidad común —
    # exactamente la misma razón por la que evaluar_reglas_interpretacion
    # (ámbito 'puntaje_total') ya excluye ambos tipos de su propio SUM. Ese
    # candado existía ahí pero no aquí, donde de verdad se construye la
    # composición — cerrado ahora, en el mismo lugar donde se previene el
    # anidamiento.
    if len(tipos_calculo_vistos) > 1:
        raise HTTPException(
            status_code=400,
            detail=(
                "No se puede mezclar componentes 'items' (suma de puntajes) con 'conteo_sintomas' "
                "(cantidad de ítems que alcanzan un umbral) dentro de la misma dimensión compuesta: "
                "producirían un total sin unidad común. Todos los componentes deben compartir el "
                f"mismo tipo_calculo (se recibió: {sorted(tipos_calculo_vistos)})."
            ),
        )


def listar_componentes_de_dimension(db: Session, dimension_compuesta_id: int):
    return db.query(models.DimensionComponente).filter(
        models.DimensionComponente.dimension_compuesta_id == dimension_compuesta_id
    ).order_by(models.DimensionComponente.orden).all()


def reemplazar_componentes_de_dimension(db: Session, dimension_compuesta_id: int, componentes: list):
    """Reemplazo COMPLETO (no merge) — mismo criterio que
    DocumentoInstitucionalSeccionContenido: la lista enviada ES la nueva
    composición, no un parche sobre la anterior."""
    dimension = db.query(models.Dimension).filter(models.Dimension.id == dimension_compuesta_id).first()
    if not dimension:
        raise HTTPException(status_code=404, detail="Dimensión no encontrada.")

    ids_componentes = [c.dimension_componente_id for c in componentes]
    _validar_componentes_dimension_compuesta(db, dimension, ids_componentes)

    db.query(models.DimensionComponente).filter(
        models.DimensionComponente.dimension_compuesta_id == dimension_compuesta_id
    ).delete(synchronize_session=False)
    for c in componentes:
        db.add(models.DimensionComponente(
            dimension_compuesta_id=dimension_compuesta_id,
            dimension_componente_id=c.dimension_componente_id,
            orden=c.orden,
        ))
    # 🔒 Auditoría de concurrencia — este reemplazo completo (delete+insert)
    # no toma ningún lock de fila: dos PUT .../componentes casi simultáneos
    # para la MISMA dimensión compuesta podían chocar contra
    # `uq_edu_dim_componentes_par` con un IntegrityError crudo (500) en vez
    # de un 409 explicado.
    commit_o_conflicto(
        db,
        "La composición de esta dimensión se modificó al mismo tiempo desde otra "
        "solicitud. Vuelve a intentarlo.",
    )
    return listar_componentes_de_dimension(db, dimension_compuesta_id)


def actualizar_opcion(db: Session, opcion_id: int, datos: schemas.OpcionRespuestaUpdate):
    opcion = db.query(models.OpcionRespuesta).filter(models.OpcionRespuesta.id == opcion_id).first()
    if not opcion:
        return None
    for campo, valor in datos.model_dump(exclude_unset=True).items():
        setattr(opcion, campo, valor)
    db.commit()
    db.refresh(opcion)
    return opcion

def eliminar_opcion(db: Session, opcion_id: int) -> bool:
    opcion = db.query(models.OpcionRespuesta).filter(models.OpcionRespuesta.id == opcion_id).first()
    if not opcion:
        return False

    #   🔒 Candado de integridad referencial (centralizado aquí, ya no solo en el router)
    tiene_respuestas = db.query(models.RespuestaAlumno).filter(
        models.RespuestaAlumno.opcion_id == opcion_id
    ).first()
    if tiene_respuestas:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="No se puede eliminar la opción: ha sido seleccionada por alumnos en sus respuestas."
        )

    db.delete(opcion)
    db.commit()
    return True

# 🆕 Auditoría de arquitectura genérica, punto 1 — estos campos definen QUÉ
# significa el tramo y A QUIÉN aplica (su rango de puntaje/edad, su nivel, si
# alerta). Editarlos sobre un baremo que YA interpretó a algún alumno real
# rompería la garantía de que DiagnosticoAlumno.baremo_id sigue apuntando a
# "lo que este tramo decía cuando se usó" — el contenido snapshotted en el
# diagnóstico seguiría siendo correcto, pero el puntero dejaría de servir
# para auditar. `sexo`/`poblacion` y el contenido informativo (interpretacion/
# conclusion_sugerida/recomendacion_sugerida/tipo_puntaje/valor_convertido)
# quedan fuera a propósito: no son el criterio de clasificación en sí.
CAMPOS_BAREMO_BLOQUEADOS_SI_EN_USO = {
    "puntaje_minimo", "puntaje_maximo", "edad_min", "edad_max", "nivel", "es_alerta",
}


def _baremo_tiene_diagnosticos(db: Session, baremo_id: int) -> bool:
    return db.query(models.DiagnosticoAlumno.id).filter(
        models.DiagnosticoAlumno.baremo_id == baremo_id
    ).first() is not None


def actualizar_baremo(db: Session, baremo_id: int, datos: schemas.BaremoUpdate):
    baremo = db.query(models.Baremo).filter(models.Baremo.id == baremo_id).first()
    if not baremo:
        return None

    cambios = datos.model_dump(exclude_unset=True)

    campos_bloqueados_pedidos = CAMPOS_BAREMO_BLOQUEADOS_SI_EN_USO & cambios.keys()
    if campos_bloqueados_pedidos and _baremo_tiene_diagnosticos(db, baremo_id):
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                f"Este baremo ya interpretó al menos un resultado real: no se puede editar "
                f"{', '.join(sorted(campos_bloqueados_pedidos))}. Desactívalo (\"activo\": false) "
                "y crea un tramo nuevo si necesitas cambiar el criterio."
            ),
        )

    # 📏 Fase 4 — BaremoUpdate es un PATCH parcial: hay que validar el
    # resultado FINAL (lo que ya tenía el baremo + el parche), no solo los
    # campos que llegaron en este request.
    class _Fusion:
        pass
    fusion = _Fusion()
    for campo in ("edad_min", "edad_max", "sexo", "poblacion", "puntaje_minimo", "puntaje_maximo", "criterio_normativo"):
        setattr(fusion, campo, cambios.get(campo, getattr(baremo, campo)))

    if fusion.puntaje_minimo > fusion.puntaje_maximo:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="puntaje_minimo no puede ser mayor que puntaje_maximo.",
        )
    if fusion.edad_min is not None and fusion.edad_max is not None and fusion.edad_min > fusion.edad_max:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="edad_min no puede ser mayor que edad_max.",
        )
    _validar_solapamiento_baremo(db, fusion, baremo.dimension_id, excluir_id=baremo.id)

    for campo, valor in cambios.items():
        setattr(baremo, campo, valor)
    db.commit()
    db.refresh(baremo)
    return baremo

def eliminar_baremo(db: Session, baremo_id: int) -> Optional[str]:
    """Borra un baremo SIN uso; desactiva (en vez de borrar) uno que ya
    interpretó algún resultado real — mismo espíritu que el resto del
    sistema (soft delete de Inscripcion/Persona): un tramo que ya produjo un
    DiagnosticoAlumno no puede desaparecer físicamente sin dejar su
    `baremo_id` en NULL (ver ondelete="SET NULL"), lo que rompería la
    trazabilidad que el punto 1 de esta fase existe para dar.

    Devuelve "eliminado" | "desactivado" | None (no existía)."""
    baremo = db.query(models.Baremo).filter(models.Baremo.id == baremo_id).first()
    if not baremo:
        return None
    if _baremo_tiene_diagnosticos(db, baremo_id):
        baremo.activo = False
        db.commit()
        return "desactivado"
    db.delete(baremo)
    db.commit()
    return "eliminado"