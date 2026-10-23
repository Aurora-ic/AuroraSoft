"""Verificación end-to-end — Remediación de permisos: 'colegio' fuera de lo clínico.

Crece junto con el plan de remediación (ver CLAUDE.md, "Permission surface:
what the matrix test can't see"): cada paso agrega su propia sección aquí en
vez de un archivo nuevo por paso.

Paso 2 — cierre mecánico de rol (`requerir_staff_clinico`): confirma contra
un servidor real que un usuario `colegio` autenticado recibe 403 en los ~24
endpoints cerrados, que `admin`/`psicologo` NO se ven afectados (deben poder
llegar más allá del chequeo de rol — pueden seguir dando 404 por un id
inexistente, eso es esperado y no es lo que esta prueba mide), y que el
catálogo de lectura (pruebas/baremos/campañas) sigue abierto a `colegio`.

Paso 3 — informes (`crud.verificar_puede_ver_informe`): monta un caso REAL
completo (campaña PHQ-9 -> matrícula -> examen -> finalizar con puntaje alto,
que deja `informe_liberado=False` al finalizar) para probar la transición de
verdad: `colegio` recibe 403 mientras está retenido y ve la fila desaparecer
del historial, `admin`/`psicologo` lo siguen viendo sin restricción, y tras
liberarlo con el endpoint real (`PATCH liberar-informe`) `colegio` pasa a ver
los tres endpoints (PDF, metadata, historial) sin ningún cambio de código —
la misma condición, evaluada de nuevo.

    uvicorn main:app --port 8010         (SMTP deshabilitado)
    python tests/verificar_permisos_colegio_e2e.py
    python tests/verificar_permisos_colegio_e2e.py --limpiar
"""

import argparse
import io
import json
import os
import sys
import time
from datetime import datetime

import requests
from dotenv import load_dotenv
from openpyxl import load_workbook
from sqlalchemy import create_engine, text

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
load_dotenv()

PREFIJO = "ZZCOLE"
API = "http://127.0.0.1:8010/api"
MOTOR = None
_resultados = []


def comprobar(nombre, condicion, detalle=""):
    _resultados.append((nombre, bool(condicion), detalle))
    print(f"  [{'PASA' if condicion else 'FALLA'}] {nombre}" + (f" — {detalle}" if detalle else ""))


def exigir(r, esperado, contexto):
    if r.status_code != esperado:
        print(f"\n!! {contexto}: esperaba {esperado}, recibió {r.status_code}\n   {r.text[:500]}")
        sys.exit(1)
    return r.json() if r.text else None


def pedir(metodo, url, contexto="", **kw):
    for _ in range(8):
        r = metodo(url, timeout=30, **kw)
        if r.status_code != 429 or "/auth/" not in url:
            return r
        print(f"    (429 en {contexto}; esperando 13s)")
        time.sleep(13)
    return r


def login(email, clave):
    return pedir(requests.post, f"{API}/auth/login", "login", data={"username": email, "password": clave})


def sql_uno(consulta, **p):
    with MOTOR.connect() as c:
        return c.execute(text(consulta), p).fetchone()


def esperar_informe(headers, inscripcion_id, timeout=20):
    """El informe se genera vía BackgroundTasks después de que /finalizar ya
    respondió — sin este poll, un chequeo inmediato de historial/metadata
    puede ver 'no existe' para CUALQUIER rol (falso positivo de 'colegio
    filtrado' cuando en realidad nadie tenía nada que ver todavía). Mismo
    patrón que `verificar_informe_pdf_e2e.py::esperar_informe`."""
    limite = time.time() + timeout
    while time.time() < limite:
        r = requests.get(f"{API}/educativo/inscripciones/{inscripcion_id}/informe", headers=headers, timeout=30)
        if r.status_code == 200 and r.json().get("existe"):
            return r.json()
        time.sleep(0.5)
    return requests.get(f"{API}/educativo/inscripciones/{inscripcion_id}/informe", headers=headers, timeout=30).json()


# Paso 2: método, ruta con un id ficticio (9999999, nunca va a existir),
# y el body mínimo si el método lo exige. El role-check de
# requerir_staff_clinico corre en el Depends, ANTES de tocar la BD o de
# validar el body a fondo — verificado manualmente: PATCH/POST con body {}
# contra un usuario 'colegio' ya da 403, nunca 422 ni 404. Por eso no hace
# falta crear baremos/preguntas/etc. reales para esta prueba.
_ENDPOINTS_STAFF_CLINICO_PASO2 = [
    ("POST", "/educativo/baremos/", {}),
    ("PATCH", "/educativo/baremos/9999999", {}),
    ("DELETE", "/educativo/baremos/9999999", None),
    ("POST", "/educativo/dimensiones/", {}),
    ("PATCH", "/educativo/dimensiones/9999999", {}),
    ("DELETE", "/educativo/dimensiones/9999999", None),
    ("POST", "/educativo/opciones/", {}),
    ("PATCH", "/educativo/opciones/9999999", {}),
    ("DELETE", "/educativo/opciones/9999999", None),
    ("POST", "/educativo/preguntas/", {}),
    ("PATCH", "/educativo/preguntas/9999999", {}),
    ("DELETE", "/educativo/preguntas/9999999", None),
    ("POST", "/educativo/pruebas/", {}),
    ("PATCH", "/educativo/pruebas/9999999", {}),
    ("POST", "/educativo/pruebas/carga-masiva/", {}),
    ("POST", "/educativo/pruebas/9999999/reglas", {}),
    ("PATCH", "/educativo/reglas/9999999", {}),
    ("POST", "/educativo/campanas/", {}),
    ("POST", "/educativo/campanas/9999999/matricular-roster", []),
    ("GET", "/educativo/campanas/9999999/links", None),
    ("GET", "/educativo/casos-pendientes", None),
    ("POST", "/educativo/seguimiento/", {}),
    ("GET", "/educativo/seguimiento/9999999", None),
    ("PATCH", "/educativo/seguimiento/9999999", {}),
    ("DELETE", "/educativo/seguimiento/9999999", None),
]

# Mismos endpoints, para confirmar que admin/psicologo SÍ pasan el chequeo de
# rol (pueden seguir adelante hasta chocar con un 404/422 de negocio, nunca
# con el 403 de rol) — regresión negativa.
_MUESTRA_PARA_ADMIN_PSICOLOGO = [
    ("PATCH", "/educativo/baremos/9999999", {}),
    ("GET", "/educativo/casos-pendientes", None),
    ("GET", "/educativo/seguimiento/9999999", None),
    ("GET", "/educativo/campanas/9999999/links", None),
]

_CATALOGO_ABIERTO_A_COLEGIO = [
    ("GET", "/educativo/pruebas/", None),
    ("GET", "/educativo/campanas/", None),
]


def _llamar(metodo, ruta, body, headers):
    url = f"{API}{ruta}"
    if metodo == "GET":
        return requests.get(url, headers=headers, timeout=30)
    if metodo == "DELETE":
        return requests.delete(url, headers=headers, timeout=30)
    if metodo == "POST":
        return requests.post(url, json=body, headers=headers, timeout=30)
    if metodo == "PATCH":
        return requests.patch(url, json=body, headers=headers, timeout=30)
    raise ValueError(metodo)


def main(url_base):
    global API, MOTOR
    API = url_base.rstrip("/") + "/api"
    MOTOR = create_engine(os.getenv("DATABASE_URL"))
    sello = datetime.utcnow().strftime("%H%M%S")

    print("\n0 · Montaje")
    r = login(os.getenv("SUPERADMIN_EMAIL", "admin@aurorasoft.com"), os.getenv("SUPERADMIN_PASSWORD"))
    h_super = {"Authorization": f"Bearer {exigir(r, 200, 'login superadmin')['access_token']}"}

    inst_id = exigir(requests.post(f"{API}/instituciones/", json={
        "nombre": f"{PREFIJO} Colegio {sello}"}, headers=h_super, timeout=30), 201, "Institución")["id"]

    tokens = {}
    for rol in ("admin", "psicologo", "colegio"):
        correo = f"{PREFIJO.lower()}.{rol}.{sello}@ejemplo.com"
        exigir(pedir(requests.post, f"{API}/auth/registro", f"registro {rol}", json={
            "email": correo, "password": "ClaveDePrueba123", "nombre_completo": f"{PREFIJO} {rol}",
            "rol": rol, "institucion_id": inst_id}, headers=h_super), 200, f"Registro {rol}")
        tokens[rol] = exigir(login(correo, "ClaveDePrueba123"), 200, f"Login {rol}")["access_token"]
    h_admin = {"Authorization": f"Bearer {tokens['admin']}"}
    h_psicologo = {"Authorization": f"Bearer {tokens['psicologo']}"}
    h_colegio = {"Authorization": f"Bearer {tokens['colegio']}"}

    print(f"\n1 · 'colegio' recibe 403 en los {len(_ENDPOINTS_STAFF_CLINICO_PASO2)} endpoints cerrados en el paso 2")
    for metodo, ruta, body in _ENDPOINTS_STAFF_CLINICO_PASO2:
        r = _llamar(metodo, ruta, body, h_colegio)
        comprobar(f"{metodo} {ruta}: colegio -> 403", r.status_code == 403, f"status={r.status_code}: {r.text[:120]}")

    print("\n2 · admin/psicologo SIGUEN pasando el chequeo de rol (no reciben el 403 de requerir_staff_clinico)")
    for metodo, ruta, body in _MUESTRA_PARA_ADMIN_PSICOLOGO:
        for rol, headers in (("admin", h_admin), ("psicologo", h_psicologo)):
            r = _llamar(metodo, ruta, body, headers)
            comprobar(f"{metodo} {ruta}: {rol} -> distinto de 403 (sigue adelante en el negocio)",
                      r.status_code != 403, f"status={r.status_code}")

    print("\n3 · El catálogo de lectura sigue abierto a 'colegio' (no se cerró de más)")
    for metodo, ruta, body in _CATALOGO_ABIERTO_A_COLEGIO:
        r = _llamar(metodo, ruta, body, h_colegio)
        comprobar(f"{metodo} {ruta}: colegio -> 200 (sigue pudiendo leer el catálogo)",
                  r.status_code == 200, f"status={r.status_code}: {r.text[:120]}")

    # ==================================================================
    # Paso 3 — informes: colegio solo si informe_liberado=True. Se monta un
    # caso real y COMPLETO (campaña -> matrícula -> examen -> finalizar) con
    # puntaje alto a propósito: PHQ-9 dispara alerta y el informe queda
    # RETENIDO (informe_liberado=False) justo después de finalizar, que es
    # exactamente el estado que se necesita probar primero. Luego se libera
    # con el endpoint real (PATCH liberar-informe, ya usado por el staff en
    # producción) para probar la transición a True sin tocar la BD a mano.
    # ==================================================================
    print("\n4 · Informes — 'colegio' solo ve un informe puntual si informe_liberado=True")
    phq9_id = exigir(requests.get(f"{API}/educativo/pruebas/", headers=h_admin, timeout=30), 200, "Catálogo pruebas")
    phq9_id = next(p["id"] for p in phq9_id if p["nombre_tecnico"] == "PHQ-9")

    persona = exigir(requests.post(f"{API}/admin/personas/", json={
        "institucion_id": inst_id, "nombres": PREFIJO, "apellidos": f"Alumno {sello}",
    }, headers=h_admin, timeout=30), 200, "Crear persona")
    exigir(requests.post(f"{API}/admin/personas/{persona['id']}/consentimiento", json={
        "persona_id": persona["id"], "institucion_id": inst_id,
        "otorgado_por_nombre": "Apoderado de prueba", "tipo_otorgante": "apoderado",
    }, headers=h_admin, timeout=30), 200, "Consentimiento")

    campana = exigir(requests.post(f"{API}/educativo/campanas/", json={
        "prueba_id": phq9_id, "nombre_grupo": f"{PREFIJO} Campaña {sello}", "institucion_id": inst_id,
    }, headers=h_admin, timeout=30), 200, "Crear campaña")
    matricula = exigir(requests.post(f"{API}/educativo/campanas/{campana['id']}/matricular-roster",
                                     json=[persona["id"]], headers=h_admin, timeout=30), 200, "Matricular")
    token = matricula["generados"][0]["token"]

    exigir(requests.post(f"{API}/educativo/comenzar-test/{token}", json={
        "nombre_completo": PREFIJO, "edad": 15, "sexo": "Femenino", "poblacion": "Secundaria",
        "acepta_terminos": True,
    }, timeout=30), 200, "Comenzar test")
    preguntas = exigir(requests.get(f"{API}/educativo/examen/{token}/preguntas", timeout=30), 200, "Preguntas")
    detalle_staff = exigir(requests.get(f"{API}/educativo/pruebas/{phq9_id}/completa", headers=h_admin, timeout=30),
                           200, "Detalle staff PHQ-9")
    preg_por_orden = {p["orden"]: p for p in detalle_staff["dimensiones"][0]["preguntas"]}

    def _opcion_id_publica(pregunta_publica_id, valor_deseado):
        orden = next(o["orden"] for o in preg_por_orden.values() if o["id"] == pregunta_publica_id)
        return next(o["id"] for o in preg_por_orden[orden]["opciones"] if o["valor_puntuacion"] == valor_deseado)

    # Puntaje alto (3 en cada uno de los 9 ítems) -> dispara alerta por baremo
    # -> requiere_revision -> informe_liberado queda False al finalizar.
    respuestas = [{"pregunta_id": p["id"], "opcion_id": _opcion_id_publica(p["id"], 3)} for p in preguntas]
    exigir(requests.post(f"{API}/educativo/examen/{token}/guardar", json={"respuestas": respuestas}, timeout=30),
           200, "Guardar respuestas")
    exigir(requests.post(f"{API}/educativo/examen/{token}/finalizar", timeout=60), 200, "Finalizar")
    inscripcion_id = sql_uno("select id from edu_inscripciones where token=:t", t=token)[0]
    liberado_tras_finalizar = sql_uno(
        "select informe_liberado from edu_inscripciones where id=:i", i=inscripcion_id)[0]
    comprobar("[montaje] Puntaje alto -> informe_liberado=False justo al finalizar (caso a probar primero)",
              liberado_tras_finalizar is False, f"informe_liberado={liberado_tras_finalizar}")

    # El informe se genera en 2º plano (BackgroundTasks) DESPUÉS de que
    # /finalizar ya respondió 200 — hay que esperarlo (con headers de staff,
    # que sí pueden ver la metadata aunque no esté liberado) antes de que
    # las pruebas de abajo signifiquen algo: si se corren demasiado pronto,
    # "colegio no ve nada en el historial" podría ser cierto para TODOS los
    # roles simplemente porque el informe todavía no existía.
    meta_montaje = esperar_informe(h_admin, inscripcion_id)
    comprobar("[montaje] El informe se generó en 2º plano antes de seguir", meta_montaje.get("existe") is True)

    print("\n4a · Con informe_liberado=False: colegio -> 403 en los tres endpoints de informe")
    r = requests.get(f"{API}/educativo/inscripciones/{inscripcion_id}/informe-pdf", headers=h_colegio, timeout=30)
    comprobar("GET .../informe-pdf: colegio -> 403 (retenido)", r.status_code == 403, f"status={r.status_code}")
    r = requests.get(f"{API}/educativo/inscripciones/{inscripcion_id}/informe", headers=h_colegio, timeout=30)
    comprobar("GET .../informe (metadata): colegio -> 403 (retenido)", r.status_code == 403, f"status={r.status_code}")
    r = requests.get(f"{API}/educativo/personas/{persona['id']}/informes", headers=h_colegio, timeout=30)
    historial_retenido = exigir(r, 200, "Historial (colegio, retenido)")
    comprobar("GET .../informes (historial): colegio SÍ recibe 200, pero la fila retenida NO aparece en la lista",
              historial_retenido == [], f"historial={historial_retenido}")

    print("\n4b · admin/psicologo SIGUEN viendo el informe retenido (no son colegio)")
    r = requests.get(f"{API}/educativo/inscripciones/{inscripcion_id}/informe", headers=h_admin, timeout=30)
    comprobar("GET .../informe: admin -> 200 pese a informe_liberado=False", r.status_code == 200,
              f"status={r.status_code}")
    r = requests.get(f"{API}/educativo/personas/{persona['id']}/informes", headers=h_psicologo, timeout=30)
    historial_psicologo = exigir(r, 200, "Historial (psicologo, retenido)")
    comprobar("GET .../informes: psicologo SÍ ve la fila retenida en su historial",
              len(historial_psicologo) == 1, f"historial={historial_psicologo}")

    exigir(requests.patch(f"{API}/educativo/inscripciones/{inscripcion_id}/liberar-informe", json={},
                          headers=h_admin, timeout=30), 200, "Liberar informe")

    print("\n4c · Tras liberar (informe_liberado=True): colegio ya SÍ ve el informe")
    r = requests.get(f"{API}/educativo/inscripciones/{inscripcion_id}/informe-pdf", headers=h_colegio, timeout=30)
    comprobar("GET .../informe-pdf: colegio -> 200 (ya liberado)", r.status_code == 200, f"status={r.status_code}")
    r = requests.get(f"{API}/educativo/inscripciones/{inscripcion_id}/informe", headers=h_colegio, timeout=30)
    comprobar("GET .../informe (metadata): colegio -> 200 (ya liberado)", r.status_code == 200, f"status={r.status_code}")
    r = requests.get(f"{API}/educativo/personas/{persona['id']}/informes", headers=h_colegio, timeout=30)
    historial_liberado = exigir(r, 200, "Historial (colegio, liberado)")
    comprobar("GET .../informes: colegio ahora SÍ ve la fila ya liberada en su historial",
              len(historial_liberado) == 1, f"historial={historial_liberado}")

    # ==================================================================
    # Paso 4 — ficha-completa: reutiliza la MISMA persona (ya tiene un
    # informe real con es_alerta=True) y le agrega una nota de seguimiento
    # real, para que 'colegio' tenga contenido clínico genuino que NO debe
    # ver, en vez de probar contra listas vacías por casualidad.
    # ==================================================================
    print("\n5 · ficha-completa — 'colegio' recibe historial_evaluaciones/sesiones_seguimiento vacíos")
    marca_confidencial = f"{PREFIJO}-NOTA-CONFIDENCIAL-{sello}"
    exigir(requests.post(f"{API}/educativo/seguimiento/", json={
        "persona_id": persona["id"], "observaciones": marca_confidencial, "accion_tomada": "Derivación",
    }, headers=h_admin, timeout=30), 200, "Crear sesión de seguimiento")

    ficha_colegio = exigir(requests.get(f"{API}/admin/personas/{persona['id']}/ficha-completa",
                                        headers=h_colegio, timeout=30), 200, "Ficha completa (colegio)")
    comprobar("ficha-completa (colegio): historial_evaluaciones vacío pese a tener un informe real con alerta",
              ficha_colegio["historial_evaluaciones"] == [], f"historial_evaluaciones={ficha_colegio['historial_evaluaciones']}")
    comprobar("ficha-completa (colegio): sesiones_seguimiento vacío, la nota confidencial NUNCA llega en el JSON",
              ficha_colegio["sesiones_seguimiento"] == [] and marca_confidencial not in json.dumps(ficha_colegio),
              f"sesiones_seguimiento={ficha_colegio['sesiones_seguimiento']}")
    comprobar("ficha-completa (colegio): datos no-clínicos (datos_personales) siguen presentes",
              ficha_colegio["datos_personales"]["id"] == persona["id"])

    ficha_admin = exigir(requests.get(f"{API}/admin/personas/{persona['id']}/ficha-completa",
                                      headers=h_admin, timeout=30), 200, "Ficha completa (admin)")
    comprobar("ficha-completa (admin): SÍ ve el historial de evaluaciones (con 'nivel' real)",
              len(ficha_admin["historial_evaluaciones"]) == 1
              and ficha_admin["historial_evaluaciones"][0]["resultados"][0]["nivel"],
              f"historial_evaluaciones={ficha_admin['historial_evaluaciones']}")
    comprobar("ficha-completa (admin): SÍ ve la nota de seguimiento confidencial",
              marca_confidencial in json.dumps(ficha_admin))

    # ==================================================================
    # Paso 5 — dashboards: mismo campaña/persona con un diagnóstico real
    # (nivel='Severo', el tramo 20-27 de PHQ-9 — ver seeds/seed_phq9.py).
    # El dashboard no depende de informe_liberado (no lo
    # chequea en absoluto) — la redacción por rol es la única protección.
    # ==================================================================
    print("\n6 · Dashboards — 'colegio' ve participación, nunca el diagnóstico por alumno")
    dash_colegio = exigir(requests.get(f"{API}/educativo/campanas/{campana['id']}/dashboard",
                                       headers=h_colegio, timeout=30), 200, "Dashboard campaña (colegio)")
    comprobar("Dashboard campaña (colegio): resultados=[] para el único alumno",
              dash_colegio["alumnos"][0]["resultados"] == [], f"alumnos={dash_colegio['alumnos']}")
    comprobar("Dashboard campaña (colegio): nombre/estado y los conteos agregados se conservan",
              PREFIJO in dash_colegio["alumnos"][0]["nombre_completo"]
              and dash_colegio["total_finalizados"] == 1
              and dash_colegio["distribucion_por_dimension"] != {})

    dash_admin = exigir(requests.get(f"{API}/educativo/campanas/{campana['id']}/dashboard",
                                     headers=h_admin, timeout=30), 200, "Dashboard campaña (admin)")
    comprobar("Dashboard campaña (admin): SÍ ve el 'nivel' real del alumno",
              dash_admin["alumnos"][0]["resultados"][0]["nivel"] == "Severo",
              f"alumnos={dash_admin['alumnos']}")

    excel_colegio = requests.get(f"{API}/educativo/campanas/{campana['id']}/exportar-excel",
                                 headers=h_colegio, timeout=30)
    comprobar("Excel de campaña (colegio): 200", excel_colegio.status_code == 200, f"status={excel_colegio.status_code}")
    filas_excel_colegio = list(load_workbook(io.BytesIO(excel_colegio.content)).active.iter_rows(
        min_row=2, values_only=True))
    comprobar("Excel de campaña (colegio): fila con '-' en vez del nivel real (mismo placeholder que 'sin baremo')",
              len(filas_excel_colegio) == 1 and filas_excel_colegio[0][1:] == ("Finalizado", "-", "-", "-")
              and PREFIJO in filas_excel_colegio[0][0],
              f"filas={filas_excel_colegio}")

    excel_admin = requests.get(f"{API}/educativo/campanas/{campana['id']}/exportar-excel",
                               headers=h_admin, timeout=30)
    comprobar("Excel de campaña (admin): 200", excel_admin.status_code == 200, f"status={excel_admin.status_code}")
    filas_excel_admin = list(load_workbook(io.BytesIO(excel_admin.content)).active.iter_rows(
        min_row=2, values_only=True))
    comprobar("Excel de campaña (admin): SÍ trae el nivel real ('Severo') en la fila",
              any(fila[-1] == "Severo" for fila in filas_excel_admin), f"filas={filas_excel_admin}")

    proceso_colegio = exigir(requests.get(f"{API}/educativo/procesos/dashboard", params={"institucion_id": inst_id},
                                          headers=h_colegio, timeout=30), 200, "Dashboard institucional (colegio)")
    comprobar("Dashboard institucional (colegio): resultados=[] en cada prueba de cada alumno",
              all(p["resultados"] == [] for a in proceso_colegio["alumnos"] for p in a["pruebas"]),
              f"alumnos={proceso_colegio['alumnos']}")

    proceso_admin = exigir(requests.get(f"{API}/educativo/procesos/dashboard", params={"institucion_id": inst_id},
                                        headers=h_admin, timeout=30), 200, "Dashboard institucional (admin)")
    comprobar("Dashboard institucional (admin): SÍ ve el 'nivel' real",
              any(r["nivel"] == "Severo" for a in proceso_admin["alumnos"] for p in a["pruebas"] for r in p["resultados"]),
              f"alumnos={proceso_admin['alumnos']}")

    total = len(_resultados)
    fallos = [n for n, ok, _ in _resultados if not ok]
    print("\n" + "=" * 72)
    print(f"RESULTADO E2E PERMISOS COLEGIO: {total - len(fallos)}/{total} comprobaciones superadas")
    for n in fallos:
        print(f"  - {n}")
    print("=" * 72)
    return 1 if fallos else 0


def limpiar():
    engine = create_engine(os.getenv("DATABASE_URL"))
    with engine.begin() as c:
        c.execute(text(f"""
            delete from edu_informe_versiones where informe_id in (
                select i.id from edu_informes i
                join edu_inscripciones ins on ins.id = i.inscripcion_id
                join edu_evaluaciones ev on ev.id = ins.evaluacion_id
                where ev.nombre_grupo like '{PREFIJO}%')"""))
        c.execute(text(f"""
            delete from edu_informes where inscripcion_id in (
                select ins.id from edu_inscripciones ins
                join edu_evaluaciones ev on ev.id = ins.evaluacion_id
                where ev.nombre_grupo like '{PREFIJO}%')"""))
        c.execute(text(f"delete from edu_diagnosticos_alumno where inscripcion_id in "
                       f"(select ins.id from edu_inscripciones ins join edu_evaluaciones ev on ev.id = ins.evaluacion_id "
                       f"where ev.nombre_grupo like '{PREFIJO}%')"))
        c.execute(text(f"delete from edu_resultados_dimensiones where inscripcion_id in "
                       f"(select ins.id from edu_inscripciones ins join edu_evaluaciones ev on ev.id = ins.evaluacion_id "
                       f"where ev.nombre_grupo like '{PREFIJO}%')"))
        c.execute(text(f"delete from edu_respuestas_alumno where inscripcion_id in "
                       f"(select ins.id from edu_inscripciones ins join edu_evaluaciones ev on ev.id = ins.evaluacion_id "
                       f"where ev.nombre_grupo like '{PREFIJO}%')"))
        c.execute(text(f"""
            delete from edu_sesiones_seguimiento where persona_id in (
                select id from personas where institucion_id in (
                    select id from instituciones where nombre like '{PREFIJO}%'))"""))
        c.execute(text(f"delete from edu_inscripciones where evaluacion_id in "
                       f"(select id from edu_evaluaciones where nombre_grupo like '{PREFIJO}%')"))
        c.execute(text(f"delete from edu_evaluaciones where nombre_grupo like '{PREFIJO}%'"))
        c.execute(text(f"""
            delete from consentimientos_informados where persona_id in (
                select id from personas where institucion_id in (
                    select id from instituciones where nombre like '{PREFIJO}%'))"""))
        c.execute(text(f"delete from personas where institucion_id in "
                       f"(select id from instituciones where nombre like '{PREFIJO}%')"))
        c.execute(text(f"delete from auth_usuarios where email like '{PREFIJO.lower()}.%'"))
        c.execute(text(f"delete from instituciones where nombre like '{PREFIJO}%'"))
    print(f"Datos de prueba {PREFIJO} eliminados.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--limpiar", action="store_true")
    parser.add_argument("--url", default="http://127.0.0.1:8010")
    args = parser.parse_args()
    if args.limpiar:
        limpiar()
        sys.exit(0)
    sys.exit(main(args.url))
